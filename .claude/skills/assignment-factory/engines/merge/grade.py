# -*- coding: utf-8 -*-
"""
다채널 주문 통합(취합/통합) 과제 — 채점 스크립트

사용:
  python eval/grade.py \
    --part1 data/consolidated_template.xlsx \
    --answer-dir eval/answer_key

자기채점(정답키 자기검증, 100점 기대):
  python eval/grade.py \
    --part1 eval/answer_key/consolidated_answer.xlsx \
    --answer-dir eval/answer_key

배점: Part1 통합표 90 (셀 Exact 60 + 행집합 Set-F1 30) / 무결성 10.
  채점 = Set(주문번호) + Exact(정규화 칸별). extra=중복미제거/유령, missing=누락.
  정답키는 build_dataset.py 가 생성한 eval/answer_key/consolidated_answer.xlsx 를 쓴다.
출력: 콘솔 표 + result.json (영역별 점수 / 셀오답·FP·FN 목록).
"""
import argparse, json, re
from pathlib import Path
import openpyxl

KEY = "주문번호"


# ---- 정규화 헬퍼 -----------------------------------------------------------
def norm_str(v):
    """문자 셀: 모든 공백 제거(앞뒤 trim 함정 + 학생 표기 관용 흡수)."""
    if v is None:
        return ""
    return re.sub(r"\s+", "", str(v)).strip()


def norm_int(v):
    """수치 셀: 빈칸 보존(0과 구분), 콤마·공백·.0 제거."""
    if v is None or str(v).strip() == "":
        return ""
    s = re.sub(r"[,\s]", "", str(v))
    try:
        return str(int(float(s)))
    except ValueError:
        return str(v).strip()


def norm_date(v):
    """날짜 셀: 3형식 → YYYY-MM-DD. 빈칸 보존."""
    s = "" if v is None else str(v).strip()
    if s == "":
        return ""
    if " " in s:
        s = s.split()[0]
    if "T" in s:                       # 엑셀이 datetime으로 저장한 경우
        s = s.split("T")[0]
    if "-" in s:
        p = s.split("-")
        if len(p) == 3:
            return f"{p[0]}-{p[1].zfill(2)}-{p[2].zfill(2)}"
        return s
    if "/" in s:
        a, b, c = s.split("/")
        if len(a) == 4:
            return f"{a}-{b.zfill(2)}-{c.zfill(2)}"
        return f"{c}-{a.zfill(2)}-{b.zfill(2)}"
    return s


def norm_channel(v):
    """출처채널: 토큰 분해→정렬→재결합(순서·공백 무관 흡수)."""
    s = norm_str(v)
    if not s:
        return ""
    toks = [t for t in s.split("|") if t]
    return "|".join(sorted(toks))


STD_FIELDS = {
    "고객명": norm_str, "상품명": norm_str, "수량": norm_int,
    "금액_원": norm_int, "주문일": norm_date, "출처채널": norm_channel,
}


def load_sheet(path, sheet):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet in wb.sheetnames else wb[wb.sheetnames[-1]]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    head = [norm_str(h) for h in rows[0]]
    out = []
    for r in rows[1:]:
        if all(c is None or str(c).strip() == "" for c in r):
            continue
        out.append({head[i]: r[i] for i in range(len(head))})
    return out


def rowkey(r):
    return norm_str(r.get(KEY))


# ---- Part1: 통합표 (Set(주문번호) + Exact(정규화 셀)) ------------------------
def grade_part1(sub, ans):
    A = {rowkey(r): r for r in ans}
    # 제출은 주문번호 중복(=dedup 실패) 가능 → 키별 행 목록
    S_list = {}
    for r in sub:
        k = rowkey(r)
        if k:
            S_list.setdefault(k, []).append(r)

    a_keys, s_keys = set(A), set(S_list)
    tp, fp, fn = a_keys & s_keys, s_keys - a_keys, a_keys - s_keys
    # 행집합 정밀도는 '제출한 데이터 행 전체'를 분모로 — 중복 미제거·유령 행이 그대로 FP가 된다
    # (extra=중복미제거/유령 → 정밀도 하락). 재현율은 정답 키 기준.
    sub_total = sum(len(lst) for lst in S_list.values())
    if not a_keys and sub_total == 0:
        prec = rec = f1 = 1.0
    else:
        prec = len(tp) / sub_total if sub_total else 0.0
        rec  = len(tp) / len(a_keys) if a_keys else 0.0
        f1   = (2 * prec * rec / (prec + rec)) if (prec + rec) else 0.0

    # 셀 Exact: 정답 키마다 (중복 제출이면 첫 행) 6칸 대조
    per_field = {f: [0, 0] for f in STD_FIELDS}
    detail = []
    for k, arow in A.items():
        srow = (S_list.get(k) or [{}])[0]
        for f, fn_norm in STD_FIELDS.items():
            exp = fn_norm(arow.get(f)); got = fn_norm(srow.get(f))
            per_field[f][1] += 1
            if exp == got:
                per_field[f][0] += 1
            elif len(detail) < 60:
                detail.append({"주문번호": k, "필드": f, "정답": exp, "제출": got})
    cells_ok = sum(v[0] for v in per_field.values())
    cells_tot = sum(v[1] for v in per_field.values())

    # 무결성: 유령 주문번호(미존재 키) + 중복 미제거(같은 주문번호 2행+)
    ghost = sorted(fp)
    dup = sorted(k for k, lst in S_list.items() if len(lst) > 1)
    return {
        "정확셀": cells_ok, "전체셀": cells_tot,
        "셀정확도": round(cells_ok / cells_tot, 4) if cells_tot else 0,
        "정답행수": len(a_keys), "제출행수(고유)": len(s_keys), "제출행수(원본)": len(sub),
        "행집합_정밀도": round(prec, 4), "행집합_재현율": round(rec, 4), "행집합_F1": round(f1, 4),
        "유령행수": len(ghost), "중복미제거수": len(dup),
        "유령행": ghost[:50], "중복미제거": dup[:50],
        "누락_FN수": len(fn), "FN샘플": sorted(fn)[:20],
        "필드별": {f: round(v[0] / v[1], 3) if v[1] else 0 for f, v in per_field.items()},
        "오답상세": detail,
    }


# ---- 배점 -----------------------------------------------------------------
# Part1 셀 Exact 60 / 행집합 Set-F1 30 / 무결성 10.
def score(p1):
    s_cell = p1["셀정확도"] * 60
    s_row = p1["행집합_F1"] * 30
    viol = p1["유령행수"] + p1["중복미제거수"]
    s_int = max(0, 10 - viol * 0.5)
    return {
        "Part1_셀Exact(60)": round(s_cell, 1),
        "Part1_행집합F1(30)": round(s_row, 1),
        "무결성(10)": round(s_int, 1),
        "총점(100)": round(s_cell + s_row + s_int, 1),
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--part1", required=True)
    ap.add_argument("--answer-dir", required=True)
    args = ap.parse_args()

    adir = Path(args.answer_dir)
    sub = load_sheet(args.part1, "통합주문")
    ans = load_sheet(adir / "consolidated_answer.xlsx", "통합주문")

    r1 = grade_part1(sub, ans)
    sc = score(r1)

    print("=" * 60); print(" 채점 결과 (다채널 주문 통합)")
    print("=" * 60)
    print(f"[Part1] 셀정확도 {r1['셀정확도']*100:.1f}% ({r1['정확셀']}/{r1['전체셀']}) | "
          f"행집합 F1 {r1['행집합_F1']*100:.1f}% (정밀도 {r1['행집합_정밀도']*100:.1f} / 재현율 {r1['행집합_재현율']*100:.1f})")
    print(f"        정답행 {r1['정답행수']} / 제출(고유) {r1['제출행수(고유)']} / 제출(원본) {r1['제출행수(원본)']} | "
          f"누락 {r1['누락_FN수']}")
    print(f"        필드별: {r1['필드별']}")
    if r1["유령행수"]:     print(f"        ⚠ 유령 주문번호(미존재): {r1['유령행수']}건")
    if r1["중복미제거수"]: print(f"        ⚠ 중복 미제거(같은 주문번호 다중 행): {r1['중복미제거수']}건")
    print("-" * 60)
    for k, v in sc.items():
        print(f"  {k:<20} {v}")
    print("=" * 60)

    out = {"part1": r1, "score": sc}
    Path("result.json").write_text(json.dumps(out, ensure_ascii=False, indent=2))
    print("→ result.json 저장")


if __name__ == "__main__":
    main()
