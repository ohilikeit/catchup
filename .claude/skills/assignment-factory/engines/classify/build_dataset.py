# -*- coding: utf-8 -*-
"""
VOC(고객의 소리) 유형·감성 2축 분류 과제 — 데이터셋/정답 생성기
(단일 입력 = items.csv)

[아키타입] 분류 (다클래스 macro-F1 + 혼동행렬 채점). "텍스트 → 범주 라벨".
  판별 기준이 "이게 뭐냐(정답 모호)"라 결정성과 충돌 → **모호함을 규칙으로 결정화**한다:
  자유 판단이 아니라 **우선순위 규칙북**(키워드 → 라벨, 위에서부터 첫 매치)으로 모든 텍스트가
  유일 라벨로 떨어진다. 정답은 생성된 text에 규칙북을 **재적용**해 도출(시나리오 의도 라벨을
  직접 베끼지 않는다 — 규칙 재적용으로 도출해야 결정성/독립검증 성립).

[입력]
  eval/answer_key/items.csv   ★ 시나리오 명세(단일 진실 소스). 헤더(정확히):
    item_id,템플릿,유형,감성,충돌주입,충돌상대,경계케이스
    · item_id   : 고유키(중복 금지).
    · 템플릿     : 문장 조립 템플릿 선택 힌트(T1~T4 또는 auto). 채점·도출과 무관(가독성만).
    · 유형       : 최종 의도 문의유형 라벨 ∈ {환불취소,배송문제,제품불량,사용법문의,칭찬감사,기타}.
                   (충돌 해소 후의 최종 라벨 — 우선순위 높은 쪽을 적는다.)
    · 감성       : 최종 의도 감성 라벨 ∈ {긍정,부정,중립}. (충돌 해소 후 — 부정>긍정이라 충돌이면 부정.)
    · 충돌주입   : none | 유형충돌 | 감성충돌.
        - 유형충돌 : '충돌상대'(유형보다 우선순위 낮은 다른 유형)의 키워드도 함께 주입.
                     규칙(우선순위 첫 매치)이 '유형'을 골라야 한다.
        - 감성충돌 : 부정어 + 긍정어를 함께 주입. 규칙(부정>긍정)이 '부정'을 골라야 한다(감성=부정 강제).
    · 충돌상대   : 충돌주입=유형충돌일 때만 의미 — 함께 주입할 하위 우선순위 유형. 그 외 '없음'.
    · 경계케이스 : none | 판별불가 | 약한신호.
        - 판별불가 : 어떤 키워드도 안 걸리는 텍스트 → 유형=기타·감성=중립(무리한 추측 금지).
        - 약한신호 : 유형 키워드만 있고 감성 신호가 없음 → 감성=중립(유형에서 감성을 추론하지 말 것).

[출력]
  data/voc_데이터.xlsx                       학생 입력(item_id, text). 라벨 없음.
  data/classification_template.xlsx          제출 양식(item_id, 문의유형, 감성 — 빈칸 + '안내')
  eval/answer_key/classification_answer.xlsx 정답(item_id, 문의유형, 감성 — 규칙 재적용 도출)

[규칙북 (rulebook — 안내시트/INPUT_GUIDE에 명시, 정답의 유일 근거)]
  · 문의유형(우선순위 순, 위에서부터 첫 매치): 1.환불취소 2.배송문제 3.제품불량 4.사용법문의 5.칭찬감사 6.기타
    한 텍스트에 여러 유형 키워드가 동시에 걸리면 **번호 작은(우선순위 높은) 라벨**. 아무것도 안 걸리면 기타.
  · 감성: 부정어가 하나라도 있으면 부정. 없고 긍정어가 있으면 긍정. 둘 다 없으면 중립. (부정 > 긍정)

[고정 기준일/시드]
  ASSIGN_DATE = 2026-06-28 (참고 상수, 채점 미사용). datetime.now() 금지.
  random.seed(20260628). items.csv만 바꿔 재실행하면 텍스트·양식·정답이 항상 정합 재생성.

[정답 도출 철학] 시나리오 의도 라벨(유형/감성 컬럼)을 직접 베껴 정답을 만들지 않는다 —
  생성된 text에 규칙북을 *다시 적용*해 라벨을 도출한다(derive_labels). 의도 라벨은 데이터 생성과
  자기검증(cross_check)에만 쓴다 → 결정성·독립검증 성립.
"""
import csv
import sys
import random
import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

# ----------------------------------------------------------------------------
# 경로 / 기준일 / 시드 (결정성 3종의 ①②)
# ----------------------------------------------------------------------------
ROOT = Path(__file__).resolve().parents[2]          # .../3.voc_classification
DATA = ROOT / "data"
ANS  = ROOT / "eval" / "answer_key"
CSV_PATH = ANS / "items.csv"
DATA.mkdir(parents=True, exist_ok=True)
ANS.mkdir(parents=True, exist_ok=True)

ASSIGN_DATE = datetime.date(2026, 6, 28)             # 고정 기준일(참고용)
random.seed(20260628)

# ── 규칙북(rulebook) ────────────────────────────────────────────────────────
# 문의유형: 우선순위 순(위에서부터 첫 매치). '기타'는 키워드 없음(fallback).
TYPE_ORDER = ["환불취소", "배송문제", "제품불량", "사용법문의", "칭찬감사", "기타"]
TYPE_KEYWORDS = {
    "환불취소":   ["환불", "취소", "반품", "돈 돌려", "결제 취소"],
    "배송문제":   ["배송", "택배", "도착", "안 와", "발송"],
    "제품불량":   ["불량", "고장", "깨져", "작동이 안", "하자", "망가"],
    "사용법문의": ["어떻게 쓰", "사용법", "사용 방법", "설정", "연결하는 방법", "방법 문의", "방법 알려"],
    "칭찬감사":   ["친절", "칭찬", "수고", "고맙", "고마웠"],
    # 기타: 키워드 없음
}
# 감성: 부정 > 긍정 > 중립(없음).
SENTI_NEG = ["화가", "실망", "최악", "불편", "짜증", "별로", "엉망"]
SENTI_POS = ["좋아", "만족", "최고", "훌륭", "마음에", "좋네", "좋습니다", "감동"]

SENTI_LABELS = ["긍정", "부정", "중립"]

# ── 자연어 생성 어구(phrase) — 각 어구는 자기 유형/감성 키워드만 포함 ─────────
TYPE_PHRASES = {
    "환불취소":   ["환불해 주세요", "주문을 취소하고 싶어요", "결제 취소 부탁드립니다",
                  "반품하고 환불받고 싶습니다", "돈 돌려주세요"],
    "배송문제":   ["배송이 언제 오나요", "택배가 아직 도착을 안 했어요", "배송이 너무 지연되네요",
                  "주문한 게 아직도 안 와요", "발송이 됐는지 궁금해요"],
    "제품불량":   ["제품이 불량이에요", "받자마자 고장났어요", "화면이 깨져서 왔어요",
                  "작동이 안 됩니다", "하자가 있네요"],
    "사용법문의": ["이거 어떻게 쓰나요", "사용법을 모르겠어요", "설정하는 방법 알려주세요",
                  "연결하는 방법이 궁금해요", "사용 방법 문의드려요"],
    "칭찬감사":   ["응대가 정말 친절하셨어요", "직원분을 칭찬하고 싶어요", "수고 많으셨습니다",
                  "고맙다는 말 전하고 싶어요", "빠른 응대 고마웠어요"],
}
SENTI_PHRASES = {
    "긍정": ["정말 좋아요", "아주 만족합니다", "서비스 최고예요", "응대가 훌륭하네요", "마음에 쏙 들어요"],
    "부정": ["너무 화가 나요", "정말 실망했어요", "최악이네요", "너무 불편합니다", "짜증이 나요"],
}
# 판별불가/기타용 — 어떤 키워드도 포함하지 않는 중립 filler
NEUTRAL_FILLER = ["문의사항이 있어 연락드립니다", "확인 한번 부탁드려요", "안내 좀 받고 싶어요",
                  "담당자분과 통화하고 싶습니다", "궁금한 점이 있어서요"]
TEMPLATES = ["{body}.", "안녕하세요. {body}.", "{body} 확인 부탁드립니다.", "문의드립니다. {body}."]


def _assert_vocab_disjoint():
    """유형 키워드 ∩ 감성 키워드 = ∅ 이어야 derive(유형/감성)가 독립 도출된다."""
    typ = [k for ks in TYPE_KEYWORDS.values() for k in ks]
    sen = SENTI_NEG + SENTI_POS
    for t in typ:
        for s in sen:
            if t in s or s in t:
                sys.exit(f"[중단] 어휘 충돌(유형↔감성): '{t}' ↔ '{s}'. 키워드를 분리하세요.")


# ----------------------------------------------------------------------------
# 1) 입력 파싱 헬퍼
# ----------------------------------------------------------------------------
def s(v):
    return (v or "").strip()


# ----------------------------------------------------------------------------
# 2) items.csv 로드 (fail-loud 검증)
# ----------------------------------------------------------------------------
def load_items():
    if not CSV_PATH.exists():
        sys.exit(
            f"[중단] 입력 파일이 없습니다: {CSV_PATH}\n"
            f"  INPUT_GUIDE.md를 참고해 items.csv를 먼저 채우세요(헤더 고정)."
        )
    rows = []
    seen = set()
    with CSV_PATH.open(encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        for i, r in enumerate(reader):
            iid = s(r.get("item_id"))
            if not iid:
                continue   # 빈 행 무시
            if iid in seen:
                sys.exit(f"[중단] 행 {i+2}: item_id 중복 — {iid}. 유일해야 합니다.")
            seen.add(iid)

            typ   = s(r.get("유형"))
            sen   = s(r.get("감성"))
            conf  = s(r.get("충돌주입")) or "none"
            rival = s(r.get("충돌상대")) or "없음"
            edge  = s(r.get("경계케이스")) or "none"

            if typ not in TYPE_ORDER:
                sys.exit(f"[중단] {iid}: 유형 '{typ}' 허용 외 {TYPE_ORDER}.")
            if sen not in SENTI_LABELS:
                sys.exit(f"[중단] {iid}: 감성 '{sen}' 허용 외 {SENTI_LABELS}.")
            if conf not in ("none", "유형충돌", "감성충돌"):
                sys.exit(f"[중단] {iid}: 충돌주입 '{conf}' 허용 외(none|유형충돌|감성충돌).")
            if edge not in ("none", "판별불가", "약한신호"):
                sys.exit(f"[중단] {iid}: 경계케이스 '{edge}' 허용 외(none|판별불가|약한신호).")

            # ── 입력 모순 fail-loud ─────────────────────────────────────────
            if conf == "유형충돌":
                if typ == "기타":
                    sys.exit(f"[중단] {iid}: 유형충돌인데 유형=기타 — 기타는 키워드가 없어 충돌 해소 불가.")
                if rival not in TYPE_ORDER or rival == "기타":
                    sys.exit(f"[중단] {iid}: 유형충돌의 충돌상대 '{rival}' 가 유효 유형이 아닙니다.")
                if TYPE_ORDER.index(rival) <= TYPE_ORDER.index(typ):
                    sys.exit(f"[중단] {iid}: 충돌상대 '{rival}'(우선순위 {TYPE_ORDER.index(rival)+1})가 "
                             f"유형 '{typ}'(우선순위 {TYPE_ORDER.index(typ)+1})보다 우선순위가 높거나 같음 — "
                             f"규칙이 유형을 못 고릅니다. 충돌상대는 유형보다 우선순위가 낮아야 합니다.")
            if conf == "감성충돌" and sen != "부정":
                sys.exit(f"[중단] {iid}: 감성충돌이면 규칙(부정>긍정)상 최종 감성은 '부정'이어야 합니다(감성='{sen}').")
            if edge == "판별불가":
                if typ != "기타" or sen != "중립":
                    sys.exit(f"[중단] {iid}: 판별불가는 유형=기타·감성=중립이어야 합니다(유형='{typ}', 감성='{sen}').")
                if conf != "none":
                    sys.exit(f"[중단] {iid}: 판별불가인데 충돌주입='{conf}' — 키워드가 없어 충돌 불가.")
            if edge == "약한신호":
                if sen != "중립":
                    sys.exit(f"[중단] {iid}: 약한신호는 감성 신호가 없어 감성=중립이어야 합니다(감성='{sen}').")
                if typ == "기타":
                    sys.exit(f"[중단] {iid}: 약한신호는 유형 키워드가 있어야 합니다(유형=기타 불가).")
                if conf != "none":
                    sys.exit(f"[중단] {iid}: 약한신호인데 충돌주입='{conf}' — 단일 약한신호와 모순.")

            rows.append(dict(item_id=iid, 템플릿=s(r.get("템플릿")) or "auto",
                             유형=typ, 감성=sen, 충돌주입=conf, 충돌상대=rival,
                             경계케이스=edge, _seq=len(rows)))
    if not rows:
        sys.exit(f"[중단] {CSV_PATH} 에 항목 행이 없습니다.")
    return rows


# ----------------------------------------------------------------------------
# 3) 텍스트 합성 — items.csv → 자연어 문장(text)
# ----------------------------------------------------------------------------
def _pick(lst, seq):
    return lst[seq % len(lst)]


def build_text(c):
    """items.csv 한 행 → 자연어 VOC 문장. 의도 라벨이 규칙으로 재도출되도록 키워드를 심는다."""
    seq = c["_seq"]
    frags = []

    if c["경계케이스"] == "판별불가":
        # 어떤 키워드도 없는 중립 텍스트
        frags.append(_pick(NEUTRAL_FILLER, seq))
    else:
        # ── 유형 어구 ───────────────────────────────────────────────
        if c["유형"] == "기타":
            # 기타(비판별불가): 유형 키워드 없이 감성/필러만
            frags.append(_pick(NEUTRAL_FILLER, seq + 2))
        else:
            frags.append(_pick(TYPE_PHRASES[c["유형"]], seq))
            if c["충돌주입"] == "유형충돌":
                # 하위 우선순위 유형 키워드도 주입(규칙이 상위=유형을 골라야 함)
                frags.append(_pick(TYPE_PHRASES[c["충돌상대"]], seq + 1))

        # ── 감성 어구 ───────────────────────────────────────────────
        if c["감성"] != "중립":
            frags.append(_pick(SENTI_PHRASES[c["감성"]], seq))
        if c["충돌주입"] == "감성충돌":
            # 긍정어도 주입(규칙 부정>긍정이 부정을 골라야 함)
            frags.append(_pick(SENTI_PHRASES["긍정"], seq + 1))

    body = ". ".join(frags)
    return _pick(TEMPLATES, seq).format(body=body)


# ----------------------------------------------------------------------------
# 4) 정답 도출 — text에 규칙북을 *다시 적용*해 라벨 도출(의도 라벨 직접 안 읽음)
# ----------------------------------------------------------------------------
def derive_labels(text):
    """규칙북 재적용: (문의유형 우선순위 첫 매치, 감성 부정>긍정>중립)."""
    t = text
    dtype = "기타"
    for label in TYPE_ORDER:                # 우선순위 순(위에서부터 첫 매치)
        if label == "기타":
            continue
        if any(kw in t for kw in TYPE_KEYWORDS[label]):
            dtype = label
            break
    if any(kw in t for kw in SENTI_NEG):
        dsenti = "부정"
    elif any(kw in t for kw in SENTI_POS):
        dsenti = "긍정"
    else:
        dsenti = "중립"
    return dtype, dsenti


def expected_from_flags(c):
    """items.csv 의도 라벨(자기검증/fail-loud 전용 — 정답 도출엔 미사용)."""
    return c["유형"], c["감성"]


def cross_check(items, texts):
    """규칙도출(derive) == 의도라벨(expected) 를 항목마다 대조. 불일치=생성기 버그 → 즉시 중단.
       이것이 '의도 라벨이 규칙대로 풀리는지' 보장하는 자기검증이다(fail-loud 핵심)."""
    for c, text in zip(items, texts):
        got = derive_labels(text)
        exp = expected_from_flags(c)
        if got != exp:
            sys.exit(f"[중단] 자기검증 실패 {c['item_id']}: 규칙도출={got} ≠ 의도라벨={exp}\n"
                     f"       text='{text}'\n"
                     f"       (어휘 누수/우선순위/충돌 설정을 점검하세요.)")


# ----------------------------------------------------------------------------
# 5) 엑셀 산출
# ----------------------------------------------------------------------------
HDR = Font(bold=True, color="FFFFFF")
HDRFILL = PatternFill("solid", fgColor="161616")
GUIDE = PatternFill("solid", fgColor="F4F4F4")

VOC_COLS    = ["item_id", "text"]
SUBMIT_COLS = ["item_id", "문의유형", "감성"]


def style_header(ws, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(1, c)
        cell.font = HDR
        cell.fill = HDRFILL
        cell.alignment = Alignment(vertical="center")


def write_voc(path, items, texts):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "VOC"
    ws.append(VOC_COLS)
    style_header(ws, len(VOC_COLS))
    for c, text in zip(items, texts):
        ws.append([c["item_id"], text])
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 80
    ws.freeze_panes = "A2"
    wb.save(path)


def build_submission(path, items, answers, answer):
    """제출 양식(빈) 또는 정답. answer=False면 라벨 빈칸, True면 도출 라벨 채움."""
    wb = openpyxl.Workbook()
    g = wb.active
    g.title = "안내"
    guide = [
        ["항목", "설명"],
        ["행 단위", "VOC 한 건 = 한 줄. data/voc_데이터.xlsx 의 모든 item_id 가 한 줄씩(item_id 채워져 있음)"],
        ["입력", "data/voc_데이터.xlsx 의 text(고객 문장)를 읽고 두 칸(문의유형·감성)을 채운다"],
        ["채점 칸", "문의유형(6클래스), 감성(3클래스) — 두 축 각각 다클래스 macro-F1"],
        ["문의유형 허용값", "환불취소 · 배송문제 · 제품불량 · 사용법문의 · 칭찬감사 · 기타 (이 6개만)"],
        ["감성 허용값", "긍정 · 부정 · 중립 (이 3개만)"],
        ["── 문의유형 규칙(우선순위) ──", "위에서부터 첫 매치. 여러 키워드가 동시에 걸리면 번호가 작은(위) 라벨을 고른다"],
        ["1. 환불취소", "환불 · 취소 · 반품 · 돈 돌려 · 결제 취소 중 하나라도 있으면"],
        ["2. 배송문제", "배송 · 택배 · 도착 · 안 와 · 발송 중 하나라도 있으면"],
        ["3. 제품불량", "불량 · 고장 · 깨져 · 작동이 안 · 하자 · 망가 중 하나라도 있으면"],
        ["4. 사용법문의", "어떻게 쓰 · 사용법 · 사용 방법 · 설정 · 연결하는 방법 · 방법 문의/알려 중 하나라도 있으면"],
        ["5. 칭찬감사", "친절 · 칭찬 · 수고 · 고맙/고마웠 중 하나라도 있으면"],
        ["6. 기타", "위 어디에도 안 걸리면 기타"],
        ["── 감성 규칙 ──", "부정어 우선. 부정어가 하나라도 있으면 부정, 없고 긍정어가 있으면 긍정, 둘 다 없으면 중립"],
        ["부정어", "화가 · 실망 · 최악 · 불편 · 짜증 · 별로 · 엉망"],
        ["긍정어", "좋아/좋네/좋습니다 · 만족 · 최고 · 훌륭 · 마음에 · 감동"],
        ["★우선순위", "한 문장에 환불·배송이 같이 있으면 환불취소(위 번호). 부정·긍정이 같이 있으면 부정"],
        ["★판별불가 = 기타/중립", "규칙 키워드가 하나도 안 걸리면 무리하게 추측하지 말고 유형=기타·감성=중립"],
        ["★유형≠감성", "유형(무엇을 묻나)과 감성(긍/부/중)은 별개 축. 칭찬이라고 무조건 긍정 아님 — 감성어만 본다"],
        ["행 순서", "채점에 영향 없음(item_id 로 대조)"],
    ]
    for r in guide:
        g.append(r)
    g.cell(1, 1).font = Font(bold=True)
    g.cell(1, 2).font = Font(bold=True)
    for row in g.iter_rows(min_row=1, max_row=g.max_row, max_col=2):
        row[0].fill = GUIDE
    g.column_dimensions["A"].width = 26
    g.column_dimensions["B"].width = 82

    sht = wb.create_sheet("분류")
    sht.append(SUBMIT_COLS)
    style_header(sht, len(SUBMIT_COLS))
    for c in items:
        if answer:
            dt, ds = answers[c["item_id"]]
            sht.append([c["item_id"], dt, ds])
        else:
            sht.append([c["item_id"], "", ""])
    sht.column_dimensions["A"].width = 12
    sht.freeze_panes = "A2"
    wb.save(path)


# ----------------------------------------------------------------------------
# main
# ----------------------------------------------------------------------------
def main():
    _assert_vocab_disjoint()
    items = load_items()
    print(f"[1/4] items.csv 로드: 항목 {len(items)}건")

    texts = [build_text(c) for c in items]
    cross_check(items, texts)   # fail-loud: 규칙도출 == 의도라벨
    print(f"[2/4] 텍스트 합성 + 자기검증 통과: {len(texts)}문장")

    # 정답 도출: 각 텍스트에 규칙 재적용
    answers = {}
    type_count = {t: 0 for t in TYPE_ORDER}
    senti_count = {s_: 0 for s_ in SENTI_LABELS}
    conflict = sum(1 for c in items if c["충돌주입"] != "none")
    edge = sum(1 for c in items if c["경계케이스"] != "none")
    undet = sum(1 for c in items if c["경계케이스"] == "판별불가")
    for c, text in zip(items, texts):
        dt, ds = derive_labels(text)
        answers[c["item_id"]] = (dt, ds)
        type_count[dt] += 1
        senti_count[ds] += 1
    print(f"[3/4] 정답 도출(규칙 재적용): 유형분포={type_count} / 감성분포={senti_count}")

    write_voc(DATA / "voc_데이터.xlsx", items, texts)
    build_submission(DATA / "classification_template.xlsx", items, answers, answer=False)
    build_submission(ANS  / "classification_answer.xlsx",   items, answers, answer=True)
    print("[4/4] 엑셀 산출 완료: voc_데이터 / 제출양식 / 정답키")

    print("=" * 60)
    print("생성 요약")
    print("  항목 수:", len(items))
    print("  유형 분포:", type_count)
    print("  감성 분포:", senti_count)
    print("  충돌 주입:", conflict, " / 경계케이스:", edge, "(판별불가", undet, ")")
    print("=" * 60)


if __name__ == "__main__":
    main()
