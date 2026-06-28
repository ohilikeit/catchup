# -*- coding: utf-8 -*-
"""
3종 대사 검증/감사 과제 — 채점 스크립트

사용:
  python eval/grade.py \
    --part1 data/reconcile_table_template.xlsx \
    --part2 data/exception_template.xlsx \
    --answer-dir eval/answer_key

자기채점(정답키 자기검증, 100점 기대):
  python eval/grade.py \
    --part1 eval/answer_key/reconcile_table_answer.xlsx \
    --part2 eval/answer_key/exception_answer.xlsx \
    --answer-dir eval/answer_key

배점: Part1 대사정리표 셀 Exact 40 / Part2 예외목록 Set-F1 50 / 무결성 10.
정답키는 build_dataset.py 가 생성한 eval/answer_key/*.xlsx 를 사용한다.
출력: 콘솔 표 + result.json (영역별 점수 / FP·FN 목록).
"""
import argparse, json, re, datetime
from pathlib import Path
import openpyxl

VIOLATIONS = {"수량불일치", "단가불일치", "금액오류", "입고누락", "계산서누락"}


# ---- 정규화 헬퍼 -----------------------------------------------------------
def norm_str(v):
    if v is None:
        return ""
    return re.sub(r"\s+", "", str(v)).strip()


def norm_int(v):
    """수치 셀: 빈칸 보존(0과 구분), 콤마·공백·.0 제거."""
    if v is None or str(v).strip() == "":
        return ""
    s = re.sub(r"[,\s]", "", str(v))
    try:
        return str(int(float(s)))
    except ValueError:
        return str(v).strip()


def norm_vio(v):
    """위반유형 문자열 정규화(공백 제거 + 동치 표기 흡수)."""
    s = norm_str(v)
    alias = {
        "수량불일치": "수량불일치", "수량오류": "수량불일치", "수량불일": "수량불일치",
        "단가불일치": "단가불일치", "단가오류": "단가불일치",
        "금액오류": "금액오류", "금액불일치": "금액오류",
        "입고누락": "입고누락", "입고없음": "입고누락",
        "계산서누락": "계산서누락", "세금계산서누락": "계산서누락", "계산서없음": "계산서누락",
    }
    return alias.get(s, s)


def load_sheet(path, sheet):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet in wb.sheetnames else wb[wb.sheetnames[-1]]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    head = [norm_str(h) for h in rows[0]]
    out = []
    for r in rows[1:]:
        if all(c is None or str(c).strip() == "" for c in r):
            continue
        out.append({head[i]: r[i] for i in range(len(head))})
    return out


# ---- Part 1: 대사정리표 (셀 단위 Exact) ------------------------------------
RECON_FIELDS = {
    "발주수량": norm_int, "발주단가": norm_int, "입고수량": norm_int,
    "청구수량": norm_int, "청구단가": norm_int, "청구금액": norm_int,
}


def reckey(r):
    return (norm_str(r.get("발주번호")), norm_str(r.get("품목코드")))


def grade_part1(sub, ans):
    A = {reckey(r): r for r in ans}
    S = {reckey(r): r for r in sub}
    per_field = {f: [0, 0] for f in RECON_FIELDS}   # [correct, total]
    detail = []
    for key, arow in A.items():
        srow = S.get(key, {})
        for f, fn in RECON_FIELDS.items():
            exp = fn(arow.get(f)); got = fn(srow.get(f))
            ok = (exp == got)
            per_field[f][1] += 1
            if ok:
                per_field[f][0] += 1
            elif len(detail) < 60:
                detail.append({"거래": "/".join(key), "필드": f, "정답": exp, "제출": got})
    cells_ok = sum(v[0] for v in per_field.values())
    cells_tot = sum(v[1] for v in per_field.values())
    return {
        "정확셀": cells_ok, "전체셀": cells_tot,
        "셀정확도": round(cells_ok / cells_tot, 4) if cells_tot else 0,
        "거래총수": len(A),
        "필드별": {f: round(v[0] / v[1], 3) if v[1] else 0 for f, v in per_field.items()},
        "오답상세": detail,
    }


# ---- Part 2: 예외목록 (Set-F1) ---------------------------------------------
def excset(rows):
    s = set()
    for r in rows:
        po = norm_str(r.get("발주번호")); it = norm_str(r.get("품목코드"))
        vt = norm_vio(r.get("위반유형"))
        if po.startswith("예시"):
            continue
        if po and it and vt:
            s.add((po, it, vt))
    return s


def grade_part2(sub, ans, valid_keys):
    A = excset(ans); S = excset(sub)
    tp = A & S; fp = S - A; fn = A - S
    if not A and not S:
        prec = rec = f1 = 1.0
    else:
        prec = len(tp) / len(S) if S else 0.0
        rec  = len(tp) / len(A) if A else 0.0
        f1   = (2 * prec * rec / (prec + rec)) if (prec + rec) else 0.0
    # 무결성: 유령 거래(존재하지 않는 발주번호+품목코드) / 허용 외 위반유형
    ghost = sorted({(po, it, vt) for (po, it, vt) in S if (po, it) not in valid_keys})
    bad_type = sorted({(po, it, vt) for (po, it, vt) in S if vt not in VIOLATIONS})
    return {
        "정답위반수": len(A), "제출위반수": len(S),
        "정밀도": round(prec, 4), "재현율": round(rec, 4), "F1": round(f1, 4),
        "오탐_FP수": len(fp), "누락_FN수": len(fn),
        "유령거래": ghost[:50], "허용외위반유형": bad_type[:50],
        "FP샘플": sorted(fp)[:20], "FN샘플": sorted(fn)[:20],
    }


# ---- 배점 -----------------------------------------------------------------
# Part1 셀정확도 40 / Part2 Set-F1 50(검증이 핵심) / 무결성 10.
def score(p1, p2):
    s1 = p1["셀정확도"] * 40
    s2 = p2["F1"] * 50
    viol = len(p2["유령거래"]) + len(p2["허용외위반유형"])
    s3 = max(0, 10 - viol * 0.5)
    return {"Part1_대사(40)": round(s1, 1), "Part2_F1(50)": round(s2, 1),
            "Part2_무결성(10)": round(s3, 1), "총점(100)": round(s1 + s2 + s3, 1)}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--part1", required=True)
    ap.add_argument("--part2", required=True)
    ap.add_argument("--answer-dir", required=True)
    args = ap.parse_args()

    adir = Path(args.answer_dir)
    p1_sub = load_sheet(args.part1, "대사정리")
    p1_ans = load_sheet(adir / "reconcile_table_answer.xlsx", "대사정리")
    p2_sub = load_sheet(args.part2, "예외목록")
    p2_ans = load_sheet(adir / "exception_answer.xlsx", "예외목록")

    valid_keys = {reckey(r) for r in p1_ans}

    r1 = grade_part1(p1_sub, p1_ans)
    r2 = grade_part2(p2_sub, p2_ans, valid_keys)
    sc = score(r1, r2)

    print("=" * 60); print(" 채점 결과 (3종 대사 검증/감사)")
    print("=" * 60)
    print(f"[Part1] 셀정확도 {r1['셀정확도']*100:.1f}% ({r1['정확셀']}/{r1['전체셀']}) | "
          f"거래 {r1['거래총수']}건")
    print(f"        필드별: {r1['필드별']}")
    print(f"[Part2] F1 {r2['F1']*100:.1f}% (정밀도 {r2['정밀도']*100:.1f} / 재현율 {r2['재현율']*100:.1f}) | "
          f"FP {r2['오탐_FP수']} / FN {r2['누락_FN수']} | 정답위반 {r2['정답위반수']} / 제출 {r2['제출위반수']}")
    if r2["유령거래"]:      print(f"        ⚠ 유령 거래(미존재 키): {len(r2['유령거래'])}건")
    if r2["허용외위반유형"]: print(f"        ⚠ 허용 외 위반유형: {len(r2['허용외위반유형'])}건")
    print("-" * 60)
    for k, v in sc.items():
        print(f"  {k:<20} {v}")
    print("=" * 60)

    out = {"part1": r1, "part2": r2, "score": sc}
    Path("result.json").write_text(json.dumps(out, ensure_ascii=False, indent=2))
    print("→ result.json 저장")


if __name__ == "__main__":
    main()
