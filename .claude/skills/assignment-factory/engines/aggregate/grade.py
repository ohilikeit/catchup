# -*- coding: utf-8 -*-
"""
이메일 캠페인 성과 집계 과제 — 채점 스크립트 (집계/분석 아키타입, Computed ±ε)

사용:
  python eval/grade.py \
    --part1 data/campaign_metrics_template.xlsx \
    --part2 data/summary_template.xlsx \
    --answer-dir eval/answer_key

자기채점(정답키 자기검증, 100점 기대):
  python eval/grade.py \
    --part1 eval/answer_key/campaign_metrics_answer.xlsx \
    --part2 eval/answer_key/summary_answer.xlsx \
    --answer-dir eval/answer_key

배점: Part1 캠페인별 집계표 55 / Part2 전체 요약 35 / 무결성 10.
채점 프리미티브:
  · norm_int  — 합계·건수(발송수/오픈수/클릭수/전환수/매출합계, 총*): 정수 Exact(빈칸 보존).
  · norm_float(tol=0.1) — 율·평균(오픈율/클릭율/전환율/ARPU, 전체전환율/전체ARPU):
      |정답-제출| < 0.1 이면 정답. 반올림 경계(X.X5)를 반대 방향으로 굴리면(정답과 0.1 차)
      오답 처리되어 트랩이 동작한다(0.1 미만만 허용, 부동소수 잡음은 4자리 반올림으로 제거).
출력: 콘솔 표 + result.json.
"""
import argparse, json, re
from pathlib import Path
import openpyxl

# ---- 채점 필드 정의 --------------------------------------------------------
P1_INT_FIELDS   = ["발송수", "오픈수", "클릭수", "전환수", "매출합계"]
P1_FLOAT_FIELDS = ["오픈율", "클릭율", "전환율", "ARPU"]
SUM_INT_FIELDS   = ["총발송", "총오픈", "총클릭", "총전환", "총매출"]
SUM_FLOAT_FIELDS = ["전체전환율", "전체ARPU"]
FLOAT_TOL = 0.1


# ---- 정규화 / 비교 프리미티브 ----------------------------------------------
def norm_str(v):
    if v is None:
        return ""
    return re.sub(r"\s+", "", str(v)).strip()


def norm_int(v):
    """합계·건수 셀: 빈칸 보존(0과 구분), 콤마·공백·.0 제거 → 정수 문자열."""
    if v is None or str(v).strip() == "":
        return ""
    s = re.sub(r"[,\s원]", "", str(v))
    try:
        return str(int(float(s)))
    except ValueError:
        return str(v).strip()


def norm_float(v):
    """율·평균 셀: 콤마·%·공백 제거 후 float. 빈칸/파싱불가 → None."""
    if v is None or str(v).strip() == "":
        return None
    s = re.sub(r"[,%\s]", "", str(v))
    try:
        return float(s)
    except ValueError:
        return None


def int_eq(answer, sub):
    return norm_int(answer) == norm_int(sub)


def float_eq(answer, sub, tol=FLOAT_TOL):
    """Computed ±ε. 정답은 항상 숫자. 제출이 비었거나 |차| >= tol 이면 오답.
       4자리 반올림으로 부동소수 잡음 제거 → 경계(0.1 차)는 엄격히 오답."""
    a = norm_float(answer)
    s = norm_float(sub)
    if a is None:
        a = 0.0
    if s is None:
        return False
    return round(abs(a - s), 4) < tol


def load_sheet(path, sheet):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet in wb.sheetnames else wb[wb.sheetnames[-1]]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    head = [norm_str(h) for h in rows[0]]
    out = []
    for r in rows[1:]:
        if all(c is None or str(c).strip() == "" for c in r):
            continue
        out.append({head[i]: r[i] for i in range(len(head))})
    return out


# ---- Part 1: 캠페인별 집계표 (키별 셀 Exact + ±ε) --------------------------
def p1key(r):
    return norm_str(r.get("campaign_id"))


def grade_part1(sub, ans):
    A = {p1key(r): r for r in ans}
    S = {p1key(r): r for r in sub if p1key(r)}
    fields = P1_INT_FIELDS + P1_FLOAT_FIELDS
    per_field = {f: [0, 0] for f in fields}     # [correct, total]
    detail = []
    for key, arow in A.items():
        srow = S.get(key, {})
        for f in P1_INT_FIELDS:
            ok = int_eq(arow.get(f), srow.get(f))
            per_field[f][1] += 1
            if ok:
                per_field[f][0] += 1
            elif len(detail) < 60:
                detail.append({"캠페인": key, "필드": f,
                               "정답": norm_int(arow.get(f)), "제출": norm_int(srow.get(f))})
        for f in P1_FLOAT_FIELDS:
            ok = float_eq(arow.get(f), srow.get(f))
            per_field[f][1] += 1
            if ok:
                per_field[f][0] += 1
            elif len(detail) < 60:
                detail.append({"캠페인": key, "필드": f,
                               "정답": norm_float(arow.get(f)), "제출": norm_float(srow.get(f))})
    cells_ok = sum(v[0] for v in per_field.values())
    cells_tot = sum(v[1] for v in per_field.values())
    # 무결성: 유령 캠페인(정답에 없는 campaign_id) + 제출 행의 퍼널 단조성 위반
    valid = set(A.keys())
    ghost = sorted(k for k in S if k not in valid)
    funnel_bad = []
    for key, srow in S.items():
        try:
            send = int(norm_int(srow.get("발송수")) or -1)
            op = int(norm_int(srow.get("오픈수")) or -1)
            ck = int(norm_int(srow.get("클릭수")) or -1)
            cv = int(norm_int(srow.get("전환수")) or -1)
        except ValueError:
            continue
        if -1 in (send, op, ck, cv):
            continue
        if not (cv <= ck <= op <= send):
            funnel_bad.append(key)
    return {
        "정확셀": cells_ok, "전체셀": cells_tot,
        "셀정확도": round(cells_ok / cells_tot, 4) if cells_tot else 0,
        "캠페인총수": len(A),
        "필드별": {f: round(v[0] / v[1], 3) if v[1] else 0 for f, v in per_field.items()},
        "유령캠페인": ghost[:50], "퍼널위반": sorted(funnel_bad)[:50],
        "오답상세": detail,
    }


# ---- Part 2: 전체 요약 1행 (셀 Exact + ±ε) ---------------------------------
def grade_part2(sub, ans):
    arow = ans[0] if ans else {}
    srow = sub[0] if sub else {}
    fields = SUM_INT_FIELDS + SUM_FLOAT_FIELDS
    per_field = {f: 0 for f in fields}
    detail = []
    for f in SUM_INT_FIELDS:
        ok = int_eq(arow.get(f), srow.get(f))
        per_field[f] = 1 if ok else 0
        if not ok:
            detail.append({"필드": f, "정답": norm_int(arow.get(f)), "제출": norm_int(srow.get(f))})
    for f in SUM_FLOAT_FIELDS:
        ok = float_eq(arow.get(f), srow.get(f))
        per_field[f] = 1 if ok else 0
        if not ok:
            detail.append({"필드": f, "정답": norm_float(arow.get(f)), "제출": norm_float(srow.get(f))})
    cells_ok = sum(per_field.values())
    cells_tot = len(fields)
    return {
        "정확셀": cells_ok, "전체셀": cells_tot,
        "셀정확도": round(cells_ok / cells_tot, 4) if cells_tot else 0,
        "필드별": {f: v for f, v in per_field.items()},
        "오답상세": detail,
    }


# ---- 배점 -----------------------------------------------------------------
# Part1 집계표 셀정확도 55 / Part2 요약 셀정확도 35 / 무결성 10.
def score(p1, p2):
    s1 = p1["셀정확도"] * 55
    s2 = p2["셀정확도"] * 35
    viol = len(p1["유령캠페인"]) + len(p1["퍼널위반"])
    s3 = max(0, 10 - viol * 0.5)
    return {"Part1_집계표(55)": round(s1, 1), "Part2_요약(35)": round(s2, 1),
            "무결성(10)": round(s3, 1), "총점(100)": round(s1 + s2 + s3, 1)}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--part1", required=True)
    ap.add_argument("--part2", required=True)
    ap.add_argument("--answer-dir", required=True)
    args = ap.parse_args()

    adir = Path(args.answer_dir)
    p1_sub = load_sheet(args.part1, "집계표")
    p1_ans = load_sheet(adir / "campaign_metrics_answer.xlsx", "집계표")
    p2_sub = load_sheet(args.part2, "전체요약")
    p2_ans = load_sheet(adir / "summary_answer.xlsx", "전체요약")

    r1 = grade_part1(p1_sub, p1_ans)
    r2 = grade_part2(p2_sub, p2_ans)
    sc = score(r1, r2)

    print("=" * 60); print(" 채점 결과 (이메일 캠페인 성과 집계)")
    print("=" * 60)
    print(f"[Part1] 셀정확도 {r1['셀정확도']*100:.1f}% ({r1['정확셀']}/{r1['전체셀']}) | "
          f"캠페인 {r1['캠페인총수']}개")
    print(f"        필드별: {r1['필드별']}")
    if r1["유령캠페인"]: print(f"        ⚠ 유령 캠페인: {len(r1['유령캠페인'])}건")
    if r1["퍼널위반"]:   print(f"        ⚠ 제출 퍼널 단조성 위반: {len(r1['퍼널위반'])}건")
    print(f"[Part2] 셀정확도 {r2['셀정확도']*100:.1f}% ({r2['정확셀']}/{r2['전체셀']})")
    print(f"        필드별: {r2['필드별']}")
    if r2["오답상세"]:  print(f"        오답: {r2['오답상세']}")
    print("-" * 60)
    for k, v in sc.items():
        print(f"  {k:<20} {v}")
    print("=" * 60)

    out = {"part1": r1, "part2": r2, "score": sc}
    Path("result.json").write_text(json.dumps(out, ensure_ascii=False, indent=2))
    print("→ result.json 저장")


if __name__ == "__main__":
    main()
