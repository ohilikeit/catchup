# -*- coding: utf-8 -*-
"""
이커머스 CRM 리텐션 캠페인 타게팅 과제 — 데이터셋/정답 생성기 (단일 입력 = campaigns.csv)

[입력]
  eval/answer_key/campaigns.csv  ★ 사용자가 실제 캠페인 데이터로 채우는 단일 진실 소스.
                                 헤더(정확히):
                                 campaign_id,name,channel,brief_file,region,age_min,age_max,
                                 dormant_days_min,dormant_days_max,total_spend_min,member_grade,
                                 pref_channel,opt_in_required,app_installed_required,matchable
                                 · dormant_days_min : 휴면 N일 이상(고객 휴면일수 >= N 통과). 비면 무관.
                                 · dormant_days_max : 최근 N일 이내(고객 휴면일수 <= N 통과). 비면 무관.
                                   ★ 휴면일수 = (ASSIGN_DATE − 고객 최근주문일).days (날짜 파생 축)
                                 · total_spend_min  : 누적구매액 하한(원). 고객 누적구매액 >= N 통과(같으면 통과). 비면 무관.
                                 · member_grade     : 회원등급 요건(일반|실버|골드|VIP). 비면 무관.
                                 · pref_channel     : 선호채널 요건(이메일|앱푸시|SMS). 비면 무관.
                                 · opt_in_required  : Y면 마케팅수신동의=Y 고객만. 비면 무관.
                                 · app_installed_required : Y면 앱설치여부=Y 고객만. 비면 무관.
                                 · matchable        : 'Y'(기본)|'N'. N=고객 명단으로 대상을 확인 불가한
                                   캠페인 → 아무에게도 추천하지 않음(매칭 0건, 의도된 결함).
  data/briefs/*                  실제 캠페인 브리프(이미지/PDF) — 운영자가 직접 배치.
                                 ── 이 스크립트는 이 폴더를 절대 만들거나 지우거나 건드리지 않는다.

[출력]
  data/customers.xlsx                  고객 DB (약 800명, 캠페인 axes 기반 경계 함정 자동 포함)
  data/campaign_table_template.xlsx    Part1 제출 양식(식별칸만 채운 빈 표 + '안내' 시트)
  data/targeting_template.xlsx         Part2 제출 양식(빈 표 + '안내' + 예시행)
  data/campaign_table_given.xlsx       ★ 분리형 Part2 입력(정답 캠페인표, matchable!=N만, 올바른 조건)
  eval/answer_key/campaign_table_answer.xlsx   Part1 정답
  eval/answer_key/targeting_answer.xlsx        Part2 정답 ((고객ID,캠페인ID) 매칭 전부)

[고정 기준일]
  ASSIGN_DATE = 2026-06-28 (신입 CRM 마케터가 과제를 받은 날). datetime.now() 절대 금지.
  ★ ASSIGN_DATE는 (1) '만 나이' 계산 (2) '휴면일수' = (ASSIGN_DATE − 최근주문일).days 에만 쓴다.
    매칭 대상 캠페인 = matchable!='N' 인 모든 캠페인.

campaigns.csv만 바꿔서 재실행하면 고객 DB·양식·정답이 항상 정합하게 재생성된다.
"""
import csv
import sys
import random
import datetime
from datetime import timedelta
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

# ----------------------------------------------------------------------------
# 경로 / 기준일 / 시드
# ----------------------------------------------------------------------------
ROOT = Path(__file__).resolve().parents[2]          # .../2.crm_targeting
DATA = ROOT / "data"
ANS  = ROOT / "eval" / "answer_key"
CSV_PATH = ANS / "campaigns.csv"
# 주의: data/briefs/ 의 실제 캠페인 브리프는 운영자가 넣는다 → 코드는 손대지 않는다.
DATA.mkdir(parents=True, exist_ok=True)
ANS.mkdir(parents=True, exist_ok=True)

# ── 과제 부여 기준일(고정) ─────────────────────────────────────────────────
# 이 상수 하나만 고정돼 있으면 만 나이·휴면일수·매칭 정답이 언제 재생성해도 동일하다.
ASSIGN_DATE = datetime.date(2026, 6, 28)                  # ← 고정 기준일(만 나이·휴면일수용)
random.seed(20260628)

# 서울 25 자치구 (지역 매칭의 기준 집합)
SEOUL25 = ["종로구","중구","용산구","성동구","광진구","동대문구","중랑구","성북구",
           "강북구","도봉구","노원구","은평구","서대문구","마포구","양천구","강서구",
           "구로구","금천구","영등포구","동작구","관악구","서초구","강남구","송파구","강동구"]

# 시도별 자치구 풀 (고객 거주지 합성·비매칭 주소 생성에 사용)
SIDO_GU = {
    "서울특별시": SEOUL25,
    "경기도": ["고양시 덕양구", "성남시 분당구", "수원시 영통구", "부천시 원미구"],
    "인천광역시": ["부평구", "남동구"],
}
# 서울 자치구별 동 (표기 변형 함정 생성용)
DONG = {g: [g[:-1] + "1동", g[:-1] + "2동", "중앙동"] for g in SEOUL25}

SURNAME = list("김이박최정강조윤장임한오서신권황안송류전홍")
GIVEN = ["민준","서연","도윤","지우","예준","하은","시우","지민","주원","서아","건우","수빈",
         "현우","지호","유진","준서","채원","지훈","다은","승현","나윤","우진","서윤","동현"]

# ── 회원등급/선호채널/수신동의/앱설치 풀 (각 판정 축 ↔ 고객 컬럼) ──────────
GRADE_ALLOWED = ("일반", "실버", "골드", "VIP")
PREF_ALLOWED  = ("이메일", "앱푸시", "SMS")
GRADE_POOL = ["일반", "일반", "일반", "실버", "골드", "VIP"]   # 일반 비중 ↑
PREF_POOL  = ["이메일", "앱푸시", "SMS"]
OPTIN_POOL = ["Y", "Y", "Y", "N"]    # 수신동의 Y 비중 ↑
APP_POOL   = ["Y", "Y", "N"]
SPEND_POOL = [0, 50_000, 150_000, 350_000, 600_000, 950_000,
              1_500_000, 2_800_000, 5_000_000]
JOIN_POOL  = ["2022-03-14", "2022-08-21", "2023-01-09", "2023-05-30", "2023-11-02"]


# ----------------------------------------------------------------------------
# 1) 입력 파싱 헬퍼
# ----------------------------------------------------------------------------
def int_or_none(v):
    """빈칸/None → None, 그 외 정수로(콤마·공백 제거)."""
    if v is None:
        return None
    s = str(v).replace(",", "").strip()
    if s == "":
        return None
    return int(float(s))


# ── 지역 파싱(일반화): 시도/자치구를 주소 문자열에서 뽑아낸다 ───────────────
# 비서울 시도 키워드 → 정식명. (서울 자치구명 substring보다 먼저 검사해야
#  '부산광역시 중구'·'경기도 강서구' 등이 서울로 오판정되지 않는다.)
NONSEOUL_SIDO = {
    "경기": "경기도", "인천": "인천광역시", "부산": "부산광역시", "대구": "대구광역시",
    "광주": "광주광역시", "대전": "대전광역시", "울산": "울산광역시", "세종": "세종특별자치시",
    "강원": "강원특별자치도", "충청북": "충청북도", "충북": "충청북도",
    "충청남": "충청남도", "충남": "충청남도", "전라북": "전북특별자치도", "전북": "전북특별자치도",
    "전라남": "전라남도", "전남": "전라남도", "경상북": "경상북도", "경북": "경상북도",
    "경상남": "경상남도", "경남": "경상남도", "제주": "제주특별자치도",
}


def detect_sido(addr):
    """주소(또는 캠페인 region) 문자열에서 시도를 판별한다."""
    if not addr:
        return None
    for key, full in NONSEOUL_SIDO.items():   # 명시적 비서울 시도 우선
        if key in addr:
            return full
    if "서울" in addr or any(g in addr for g in SEOUL25):
        return "서울특별시"
    for tok in addr.split():                  # 일반 폴백: 시도형 접미사 토큰
        if tok.endswith(("특별자치시", "특별자치도", "특별시", "광역시", "도")):
            return tok
    return None


def detect_gu(addr):
    """주소(또는 캠페인 district)에서 자치구 토큰을 뽑는다."""
    if not addr:
        return None
    for g in SEOUL25:
        if g in addr:
            return g
    for tok in addr.split():
        if tok.endswith("구"):
            return tok
    return None


def parse_campaign_region(region_str):
    """캠페인 region을 3형식으로 일반화: 전국 / <시도> / <시도> <자치구>."""
    s = (region_str or "").strip()
    if s == "" or s == "전국":
        return dict(nation=True, sido=None, district=None)
    toks = s.split()
    sido = detect_sido(s) or toks[0]
    district = " ".join(toks[1:]) if len(toks) > 1 else None
    return dict(nation=False, sido=sido, district=district)


def region_match(member_addr, region_parsed):
    """고객 거주지가 캠페인 지역 요건을 만족하는지."""
    if region_parsed["nation"]:
        return True
    if detect_sido(member_addr) != region_parsed["sido"]:
        return False
    if region_parsed["district"] is None:
        return True
    member_gu = detect_gu(member_addr)
    policy_gu = detect_gu(region_parsed["district"]) or region_parsed["district"]
    return member_gu == policy_gu


# ----------------------------------------------------------------------------
# 2) campaigns.csv 로드
# ----------------------------------------------------------------------------
def load_campaigns():
    if not CSV_PATH.exists():
        sys.exit(
            f"[중단] 입력 파일이 없습니다: {CSV_PATH}\n"
            f"  INPUT_GUIDE.md를 참고해 campaigns.csv를 먼저 채우세요(헤더 고정)."
        )
    rows = []
    with CSV_PATH.open(encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        for r in reader:
            cid = (r.get("campaign_id") or "").strip()
            if not cid:
                continue   # 빈 행 무시
            region_raw = (r.get("region") or "").strip()
            if not region_raw:   # 빈 region이 조용히 '전국'이 되어 대량 오발송하는 것 방지(fail loud)
                sys.exit(f"[중단] {cid}: region이 비어 있습니다. '전국'·'서울특별시'·'서울특별시 관악구' "
                         f"처럼 명시하세요(빈칸 금지).")
            grade = (r.get("member_grade") or "").strip()
            if grade and grade not in GRADE_ALLOWED:
                sys.exit(f"[중단] {cid}: member_grade '{grade}'는 허용되지 않음({'|'.join(GRADE_ALLOWED)}).")
            pref = (r.get("pref_channel") or "").strip()
            if pref and pref not in PREF_ALLOWED:
                sys.exit(f"[중단] {cid}: pref_channel '{pref}'는 허용되지 않음({'|'.join(PREF_ALLOWED)}).")
            matchable = ((r.get("matchable") or "").strip().upper() or "Y")
            if matchable not in ("Y", "N"):
                sys.exit(f"[중단] {cid}: matchable '{matchable}'는 허용되지 않음(Y|N).")
            opt_in = ((r.get("opt_in_required") or "").strip().upper() == "Y")
            app_req = ((r.get("app_installed_required") or "").strip().upper() == "Y")
            axis_vals = [
                int_or_none(r.get("age_min")), int_or_none(r.get("age_max")),
                int_or_none(r.get("dormant_days_min")), int_or_none(r.get("dormant_days_max")),
                int_or_none(r.get("total_spend_min")),
            ]
            if matchable == "N" and (any(v is not None for v in axis_vals)
                                     or grade or pref or opt_in or app_req):
                sys.exit(f"[중단] {cid}: matchable=N(판별불가) 캠페인엔 자격 축(나이/휴면/누적구매액/"
                         f"등급/채널/수신동의/앱설치)을 비워 두세요 — 정답표의 그 칸은 빈칸이어야 하는데 "
                         f"값이 채워져 자기모순.")
            rows.append(dict(
                id=cid,
                name=(r.get("name") or "").strip(),
                channel=(r.get("channel") or "").strip(),
                brief=(r.get("brief_file") or "").strip(),   # 참조용(생성엔 미사용)
                region_raw=region_raw,
                region=parse_campaign_region(region_raw),
                age_min=int_or_none(r.get("age_min")),
                age_max=int_or_none(r.get("age_max")),
                dormant_days_min=int_or_none(r.get("dormant_days_min")),
                dormant_days_max=int_or_none(r.get("dormant_days_max")),
                total_spend_min=int_or_none(r.get("total_spend_min")),
                member_grade=grade,        # 비면 무관
                pref_channel=pref,         # 비면 무관
                opt_in_required=opt_in,    # bool
                app_installed_required=app_req,  # bool
                matchable=matchable,       # Y|N (N=판별불가 → 추천 0건)
            ))
    if not rows:
        sys.exit(f"[중단] {CSV_PATH} 에 캠페인 행이 없습니다.")
    return rows


# ----------------------------------------------------------------------------
# 3) 고객 합성 — 캠페인 axes를 읽어 generic 경계 케이스 생성(하드코딩 금지)
# ----------------------------------------------------------------------------
def manage(birth):
    """기준일(ASSIGN_DATE) 기준 만 나이. 생일 안 지났으면 -1."""
    a = ASSIGN_DATE.year - birth.year
    if (ASSIGN_DATE.month, ASSIGN_DATE.day) < (birth.month, birth.day):
        a -= 1
    return a


def birth_for_age(age):
    """기준일에 정확히 만 `age`세가 되도록 생년월일 생성(생일은 이미 지난 1~5월)."""
    return datetime.date(ASSIGN_DATE.year - age,
                         random.randint(1, 5), random.randint(1, 28))


def birth_unpassed(manage_target):
    """기준일에 '생일이 아직 안 지난'(7~12월생) 고객으로, **만 나이 = manage_target**.
       단순 출생연도 빼기(2026 - 출생연도)로 풀면 manage_target+1 이 나와 틀린다 — 만 나이 함정."""
    return datetime.date(ASSIGN_DATE.year - 1 - manage_target,
                         random.randint(7, 12), random.randint(1, 28))


def birth_random(age):
    """일반 고객용: 출생월을 1~12월 무작위로(생일 경과/미경과 섞임)."""
    return datetime.date(ASSIGN_DATE.year - age,
                         random.randint(1, 12), random.randint(1, 28))


def order_for_dormant(d):
    """휴면일수 d → 최근주문일 = ASSIGN_DATE − d일. (날짜 파생 축의 역산)"""
    return ASSIGN_DATE - timedelta(days=int(d))


def seoul_addr(gu):
    """서울 자치구 주소를 표기 변형(서울특별시/서울/시 생략 등)으로 생성 — 정규화 함정."""
    dong = random.choice(DONG.get(gu, ["중앙동"]))
    return random.choice([
        f"서울특별시 {gu} {dong}",
        f"서울 {gu} {dong}",
        f"서울시 {gu}",
        f"{gu} {dong}",          # 시 표기 생략(자치구명만)
    ])


def make_address(sido, district):
    """(시도, 자치구) → 고객 주소 문자열."""
    if sido == "서울특별시":
        gu = detect_gu(district) or district
        return seoul_addr(gu)
    return f"{sido} {district} 중앙동"


def matching_address(p):
    """캠페인 p의 지역 요건을 만족하는 고객 주소."""
    reg = p["region"]
    if reg["nation"]:
        return seoul_addr(random.choice(SEOUL25))
    sido = reg["sido"]
    district = reg["district"]
    if district is None:                       # 시도만 지정 → 그 시도의 임의 자치구
        pool = SIDO_GU.get(sido, [])
        district = random.choice(pool) if pool else ""
    return make_address(sido, district)


def nonmatching_address(p):
    """캠페인 p의 지역 요건을 만족하지 않는 고객 주소(전국 캠페인은 진짜 불일치가 없어 다른 시도로 다양화)."""
    reg = p["region"]
    if reg["nation"]:
        return "경기도 고양시 덕양구 화정동"     # 전국엔 불일치 없음 → 지리적 다양성용
    sido = reg["sido"]
    district = reg["district"]
    if district is not None:                   # 같은 시도, 다른 자치구
        gu = detect_gu(district)
        if sido == "서울특별시":
            other = random.choice([g for g in SEOUL25 if g != gu])
            return seoul_addr(other)
        pool = [d for d in SIDO_GU.get(sido, []) if detect_gu(d) != gu]
        return make_address(sido, random.choice(pool)) if pool else make_address(sido, district)
    # 시도만 지정 → 다른 시도
    others = [s for s in SIDO_GU if s != sido]
    osido = random.choice(others) if others else "경기도"
    return make_address(osido, random.choice(SIDO_GU[osido]))


def passing_age(p):
    """캠페인 p의 연령 요건을 확실히 만족하는 나이 하나."""
    lo, hi = p["age_min"], p["age_max"]
    if lo is not None and hi is not None:
        return (lo + hi) // 2
    if lo is not None:
        return lo + 3
    if hi is not None:
        return max(hi - 3, 19)
    return 35


def passing_dormant(p):
    """캠페인 p의 휴면일수 요건을 확실히 만족하는 휴면일수 하나."""
    lo, hi = p["dormant_days_min"], p["dormant_days_max"]
    if lo is not None and hi is not None:
        return (lo + hi) // 2
    if lo is not None:
        return lo + 10
    if hi is not None:
        return max(hi - 2, 0)
    return 30


def passing_spend(p):
    """캠페인 p의 누적구매액 하한을 만족하는 누적구매액(없으면 안전 기본값)."""
    return p["total_spend_min"] if p["total_spend_min"] is not None else 300_000


def passing_grade(p):
    """캠페인 p의 회원등급 요건을 만족하는 회원등급."""
    return p["member_grade"] if p["member_grade"] else "일반"


def passing_pref(p):
    """캠페인 p의 선호채널 요건을 만족하는 선호채널."""
    return p["pref_channel"] if p["pref_channel"] else "이메일"


_members = []
_mid = [1]


def add_member(birth, addr, order_date, spend, grade, pref, optin, app):
    """고객 1명 추가. 자동 ID(U0001…).
       birth=생년월일(date), order_date=최근주문일(date), spend=누적구매액_원,
       grade=회원등급, pref=선호채널, optin=마케팅수신동의(Y|N), app=앱설치여부(Y|N)."""
    n = _mid[0]
    rec = dict(
        고객ID=f"U{n:04d}",
        이름=random.choice(SURNAME) + random.choice(GIVEN),
        생년월일=birth.isoformat(),
        성별=random.choice(["남", "여"]),
        거주지=addr,
        최근주문일=order_date.isoformat(),
        가입일=random.choice(JOIN_POOL),
        누적구매액_원=spend,
        회원등급=grade,
        선호채널=pref,
        마케팅수신동의=optin,
        앱설치여부=app,
        이메일=f"user{n:04d}@example.com",
    )
    _members.append(rec)
    _mid[0] += 1


def plant_boundaries(p):
    """캠페인 p의 각 axis(나이·휴면일수·지역·누적구매액·회원등급·선호채널·수신동의·앱설치)에
       경계 고객을 generic하게 심는다. 경계 외 다른 axis는 모두 '통과' 값으로 둔다."""
    base_addr = matching_address(p)
    base_age = passing_age(p)
    base_dormant = passing_dormant(p)
    base_spend = passing_spend(p)
    base_grade = passing_grade(p)
    base_pref = passing_pref(p)

    def m(age=None, birth=None, addr=None, dormant=None, order_date=None,
          spend=None, grade=None, pref=None, optin="Y", app="Y"):
        if birth is None:
            birth = birth_for_age(base_age if age is None else age)
        if order_date is None:
            order_date = order_for_dormant(base_dormant if dormant is None else dormant)
        add_member(birth, addr or base_addr, order_date,
                   base_spend if spend is None else spend,
                   grade or base_grade, pref or base_pref, optin, app)

    # (A) 나이 경계: age_min-1 / age_min / age_max / age_max+1
    age_targets = []
    if p["age_min"] is not None:
        age_targets += [p["age_min"] - 1, p["age_min"]]
    if p["age_max"] is not None:
        age_targets += [p["age_max"], p["age_max"] + 1]
    for ta in age_targets:
        if ta < 0:
            continue
        m(age=ta)

    # (A') 생일 미경과(7~12월생) 경계 — '단순 출생연도 빼기'로 풀면 틀리는 만 나이 함정
    if p["age_max"] is not None:
        m(birth=birth_unpassed(p["age_max"]))             # 만 나이=max(정답=포함), 단순빼기=max+1(FN)
    if p["age_min"] is not None and p["age_min"] - 1 >= 0:
        m(birth=birth_unpassed(p["age_min"] - 1))         # 만 나이=min-1(정답=제외), 단순빼기=min(FP)

    # (B) 휴면일수 경계: dormant_min-1/min/min+1 (89/90/91) · dormant_max-1/max/max+1
    if p["dormant_days_min"] is not None:
        for d in (p["dormant_days_min"] - 1, p["dormant_days_min"], p["dormant_days_min"] + 1):
            if d < 0:
                continue
            m(dormant=d)
    if p["dormant_days_max"] is not None:
        for d in (max(p["dormant_days_max"] - 1, 0), p["dormant_days_max"], p["dormant_days_max"] + 1):
            m(dormant=d)

    # (C) 지역: 일치 1 + 불일치 1 (다른 axis는 통과)
    m(addr=base_addr)
    m(addr=nonmatching_address(p))

    # (D) 누적구매액 하한 경계: min-1(탈락) / min(통과, >=)
    if p["total_spend_min"] is not None:
        for sp in (max(p["total_spend_min"] - 1, 0), p["total_spend_min"]):
            m(spend=sp)

    # (E) 회원등급 경계: 일치(통과)·불일치(탈락)
    if p["member_grade"]:
        miss = "일반" if p["member_grade"] != "일반" else "VIP"
        m(grade=p["member_grade"])
        m(grade=miss)

    # (F) 선호채널 경계: 일치(통과)·불일치(탈락)
    if p["pref_channel"]:
        miss = "이메일" if p["pref_channel"] != "이메일" else "SMS"
        m(pref=p["pref_channel"])
        m(pref=miss)

    # (G) 수신동의 경계: 정책이 요구하면 Y(통과)·N(탈락)
    if p["opt_in_required"]:
        m(optin="Y")
        m(optin="N")

    # (H) 앱설치 경계: 정책이 요구하면 Y(통과)·N(탈락)
    if p["app_installed_required"]:
        m(app="Y")
        m(app="N")


def synthesize_members(campaigns, target_count=800):
    """캠페인 경계 함정 + 일반 무작위 고객으로 고객 DB를 구성한다."""
    _members.clear()          # 같은 프로세스에서 재호출해도 누적되지 않게 초기화(결정성)
    _mid[0] = 1
    random.seed(20260628)     # 재호출 시에도 동일 난수 시퀀스
    for p in campaigns:
        if p["matchable"] == "N":
            continue   # 판별불가 캠페인 → 어차피 0매칭 → 경계 고객을 심지 않음
        plant_boundaries(p)

    # 일반 무작위 고객으로 약 target_count명까지 채운다.
    all_addr_makers = [lambda: seoul_addr(random.choice(SEOUL25))]
    for sido, gus in SIDO_GU.items():
        if sido == "서울특별시":
            continue
        all_addr_makers.append(lambda s=sido, g=gus: make_address(s, random.choice(g)))
    while len(_members) < target_count:
        a = random.randint(19, 64)
        dormant = random.randint(0, 400)
        add_member(
            birth_random(a),
            random.choice(all_addr_makers)(),
            order_for_dormant(dormant),
            random.choice(SPEND_POOL),
            random.choice(GRADE_POOL),
            random.choice(PREF_POOL),
            random.choice(OPTIN_POOL),
            random.choice(APP_POOL),
        )
    return _members


# ----------------------------------------------------------------------------
# 4) 매칭 정답 계산 (matchable!='N' 캠페인 대상, 모든 축 동시 통과)
# ----------------------------------------------------------------------------
def member_passes(m, p):
    # 판별불가 캠페인: 고객 명단으로 대상을 확인할 수 없음 → 어떤 고객도 통과 못 함(추천 0건).
    if p["matchable"] == "N":
        return False
    a = manage(datetime.date.fromisoformat(m["생년월일"]))
    if p["age_min"] is not None and a < p["age_min"]:
        return False
    if p["age_max"] is not None and a > p["age_max"]:
        return False
    if not region_match(m["거주지"], p["region"]):
        return False
    dormant = (ASSIGN_DATE - datetime.date.fromisoformat(m["최근주문일"])).days
    if p["dormant_days_min"] is not None and dormant < p["dormant_days_min"]:
        return False
    if p["dormant_days_max"] is not None and dormant > p["dormant_days_max"]:
        return False
    if p["total_spend_min"] is not None and (m["누적구매액_원"] or 0) < p["total_spend_min"]:
        return False
    if p["member_grade"] and m["회원등급"] != p["member_grade"]:
        return False
    if p["pref_channel"] and m["선호채널"] != p["pref_channel"]:
        return False
    if p["opt_in_required"] and m["마케팅수신동의"] != "Y":
        return False
    if p["app_installed_required"] and m["앱설치여부"] != "Y":
        return False
    return True


# ----------------------------------------------------------------------------
# 5) 엑셀 산출 (고객DB / 제출양식 / 정답키)
# ----------------------------------------------------------------------------
HDR = Font(bold=True, color="FFFFFF")
HDRFILL = PatternFill("solid", fgColor="161616")
GUIDE = PatternFill("solid", fgColor="F4F4F4")

MEMBER_COLS = ["고객ID", "이름", "생년월일", "성별", "거주지", "최근주문일", "가입일",
               "누적구매액_원", "회원등급", "선호채널", "마케팅수신동의", "앱설치여부", "이메일"]
P1COLS = ["캠페인ID", "캠페인명", "채널", "거주지요건", "연령_min", "연령_max",
          "휴면일수_min", "휴면일수_max", "누적구매액하한_원", "회원등급요건",
          "선호채널요건", "수신동의요건", "앱설치요건"]


def style_header(ws, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(1, c)
        cell.font = HDR
        cell.fill = HDRFILL
        cell.alignment = Alignment(vertical="center")


def write_members(members):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "고객"
    ws.append(MEMBER_COLS)
    style_header(ws, len(MEMBER_COLS))
    for m in members:
        ws.append([m[c] for c in MEMBER_COLS])
    ws.freeze_panes = "A2"
    wb.save(DATA / "customers.xlsx")


def p1_row(p, answer):
    """Part1 행. answer=False면 식별칸(ID/명/채널)만, True면 정답 채움.
       자격 10칸(거주/연령min·max/휴면min·max/누적하한/등급/선호채널/수신동의/앱설치)만."""
    if not answer:
        return [p["id"], p["name"], p["channel"]] + [""] * 10
    return [
        p["id"], p["name"], p["channel"],
        p["region_raw"],
        p["age_min"] if p["age_min"] is not None else "",
        p["age_max"] if p["age_max"] is not None else "",
        p["dormant_days_min"] if p["dormant_days_min"] is not None else "",
        p["dormant_days_max"] if p["dormant_days_max"] is not None else "",
        p["total_spend_min"] if p["total_spend_min"] is not None else "",
        p["member_grade"] if p["member_grade"] else "",
        p["pref_channel"] if p["pref_channel"] else "",
        "필요" if p["opt_in_required"] else "",
        "필요" if p["app_installed_required"] else "",
    ]


def build_p1(path, campaigns, answer):
    wb = openpyxl.Workbook()
    g = wb.active
    g.title = "안내"
    guide = [
        ["항목", "설명"],
        ["행 단위", "캠페인 1건 (data/briefs/ 의 브리프 전부)"],
        ["입력", "data/briefs/ 의 실제 캠페인 브리프(이미지·PDF)"],
        ["채점 칸", "거주지요건, 연령_min/max, 휴면일수_min/max, 누적구매액하한_원, 회원등급요건, "
                  "선호채널요건, 수신동의요건, 앱설치요건"],
        ["거주지요건", "허용값: '전국' 또는 '서울특별시' 또는 '서울특별시 관악구'(시도/자치구). 브리프 그대로. 예: 전국"],
        ["연령_min/max", "형식: 만 나이 정수. 브리프에 하한/상한 없으면 빈칸. 예: 25 / 54"],
        ["휴면일수_min", "형식: 정수(일). '휴면 N일 이상'이면 N. 휴면일수 = 기준일(2026-06-28) − 최근주문일. "
                       "없으면 빈칸. 예: 90"],
        ["휴면일수_max", "형식: 정수(일). '최근 N일 이내 구매'면 N. 없으면 빈칸. 예: 14"],
        ["누적구매액하한_원", "형식: 정수(원). '누적구매액 N원 이상'이면 N(고객 누적구매액 ≥ N 통과, 같으면 통과). "
                          "없으면 빈칸 (0과 빈칸은 다름). 예: 1000000"],
        ["회원등급요건", "허용값: '일반'·'실버'·'골드'·'VIP'. 무관/언급 없으면 빈칸. 예: VIP"],
        ["선호채널요건", "허용값: '이메일'·'앱푸시'·'SMS'. 무관/언급 없으면 빈칸. 예: 앱푸시"],
        ["수신동의요건", "허용값: '필요' 또는 빈칸. 마케팅 수신동의가 필수면 '필요'. 예: 필요"],
        ["앱설치요건", "허용값: '필요' 또는 빈칸. 앱 설치가 필수면 '필요'. 예: 필요"],
        ["빈칸 규칙", "빈칸 = 제약 없음(무관). 0과 빈칸은 다르다 — 무관 칸에 0이나 '무관' 글자를 넣지 말고 비운다"],
        ["★확인 불가 규칙", "브리프에 안 나오거나 고객 명단(customers.xlsx 컬럼)으로 잴 수 없는 자격 칸은 비운다(억지 값=오답). "
                          "명단으로 대상을 끝까지 확인할 수 없는 캠페인은 Part2에서도 아무에게도 추천하지 않는다 — 확실할 때만"],
        ["캠페인명·채널", "참조용(채점 대상 아님). 이미 채워져 있음"],
    ]
    for r in guide:
        g.append(r)
    g.cell(1, 1).font = Font(bold=True)
    g.cell(1, 2).font = Font(bold=True)
    for row in g.iter_rows(min_row=1, max_row=g.max_row, max_col=2):
        row[0].fill = GUIDE
    g.column_dimensions["A"].width = 16
    g.column_dimensions["B"].width = 80

    s = wb.create_sheet("캠페인정리")
    s.append(P1COLS)
    style_header(s, len(P1COLS))
    for p in campaigns:
        s.append(p1_row(p, answer))
    s.freeze_panes = "A2"
    wb.save(path)


def build_p2(path, matches, answer):
    wb = openpyxl.Workbook()
    g = wb.active
    g.title = "안내"
    guide = [
        ["항목", "설명"],
        ["행 단위", "(고객 × 추천캠페인) 1조합"],
        ["입력", "data/customers.xlsx + data/campaign_table_given.xlsx(주어진 정답 캠페인표). "
               "★ 2단계는 1단계 결과가 아니라 제공된 정답 캠페인표를 쓴다 — 1단계를 틀려도 영향 없다"],
        ["고객ID", "customers.xlsx의 고객 식별자(U0001 …)"],
        ["추천캠페인ID", "자격(대상 조건)이 맞는 캠페인의 캠페인ID"],
        ["기록 규칙", "한 고객이 한 캠페인의 대상 조건을 '모두' 통과하면 그 (고객,캠페인)을 한 줄. "
                    "어떤 캠페인에도 안 맞는 고객은 적지 않음"],
        ["── 대상 판정 규칙 ──", "아래를 모두 통과해야 추천. 캠페인표 칸이 비어 있으면(무관) 그 조건은 보지 않는다"],
        ["나이", "고객 만 나이가 캠페인 연령(연령_min~연령_max, 양끝 포함) 안. "
                "만 나이 = 기준일(2026-06-28) 기준이며, 생일이 아직 안 지났으면 한 살 뺀다(연도만 빼면 틀림)"],
        ["거주지", "거주지요건이 '전국'=모든 고객 / 시도(예 서울특별시)=그 시도 거주자 / "
                  "시도+자치구=그 자치구 거주자. 고객 주소 표기가 제각각이어도 시·도와 자치구로 판단"],
        ["휴면일수", "휴면일수 = 기준일(2026-06-28) − 고객 최근주문일(일수). "
                   "휴면일수_min이 있으면 고객 휴면일수 ≥ min(휴면 오래된 고객) / "
                   "휴면일수_max가 있으면 ≤ max(최근 구매 고객) / 빈칸=모두"],
        ["누적구매액", "누적구매액하한_원이 있으면 고객 누적구매액_원이 그 이상(같으면 통과) / 빈칸=모두"],
        ["회원등급", "회원등급요건이 있으면(예 VIP) 고객 회원등급이 그 값과 같아야 통과 / 빈칸=모두"],
        ["선호채널", "선호채널요건이 있으면(예 앱푸시) 고객 선호채널이 그 값과 같아야 통과 / 빈칸=모두"],
        ["수신동의", "수신동의요건이 '필요'면 고객 마케팅수신동의가 'Y' / 빈칸=모두"],
        ["앱설치", "앱설치요건이 '필요'면 고객 앱설치여부가 'Y' / 빈칸=모두"],
        ["★확인 불가 캠페인", "고객 명단(customers.xlsx의 컬럼)만으로는 대상을 끝까지 확인할 수 없는 캠페인이 섞여 있을 수 있다. "
                          "그런 캠페인은 아무에게도 추천하지 않는다 — 확실할 때만 추천하고, 무리하면 오발송(FP)이다"],
        ["행 순서", "채점에 영향 없음. 중복 행은 1개로 간주"],
        ["예시", "아래 '타게팅' 시트의 예시 행 참고(실제 답 아님, 채점 시 무시)"],
    ]
    for r in guide:
        g.append(r)
    g.cell(1, 1).font = Font(bold=True)
    g.cell(1, 2).font = Font(bold=True)
    for row in g.iter_rows(min_row=1, max_row=g.max_row, max_col=2):
        row[0].fill = GUIDE
    g.column_dimensions["A"].width = 18
    g.column_dimensions["B"].width = 80

    s = wb.create_sheet("타게팅")
    s.append(["고객ID", "추천캠페인ID"])
    style_header(s, 2)
    if answer:
        for mid_, pid in sorted(set(matches)):
            s.append([mid_, pid])
    else:
        s.append(["예시U0001", "C01"])
    s.freeze_panes = "A2"
    wb.save(path)


# ----------------------------------------------------------------------------
# main
# ----------------------------------------------------------------------------
def main():
    campaigns = load_campaigns()
    print(f"[1/4] campaigns.csv 로드: 캠페인 {len(campaigns)}건")

    members = synthesize_members(campaigns)
    print(f"[2/4] 고객 {len(members)}명 합성(캠페인 axes 기반 경계 함정 포함)")

    # 매칭 대상 = matchable!='N' 인 모든 캠페인(판별불가 캠페인은 추천 0건).
    targets = [p for p in campaigns if p["matchable"] != "N"]
    matches = []
    for m in members:
        for p in targets:
            if member_passes(m, p):
                matches.append((m["고객ID"], p["id"]))
    print(f"[3/4] 매칭 {len(matches)}건 "
          f"(대상 캠페인 {len(targets)}개, 고객 {len(set(x[0] for x in matches))}명 발송)")

    write_members(members)
    build_p1(DATA / "campaign_table_template.xlsx", campaigns, answer=False)
    build_p1(ANS  / "campaign_table_answer.xlsx",   campaigns, answer=True)
    # ★ 분리형(decouple): Part2 입력용 '정답 캠페인표' — matchable!=N 만, 올바른 조건으로 채움.
    given = [p for p in campaigns if p["matchable"] != "N"]
    build_p1(DATA / "campaign_table_given.xlsx", given, answer=True)
    build_p2(DATA / "targeting_template.xlsx", matches, answer=False)
    build_p2(ANS  / "targeting_answer.xlsx",   matches, answer=True)
    print("[4/4] 엑셀 산출 완료: customers / 제출양식 2 / 분리형입력 1 / 정답키 2")

    # ── 생성 요약 ───────────────────────────────────────────────────────────
    per_policy = {p["id"]: sum(1 for x in matches if x[1] == p["id"]) for p in targets}
    print("=" * 60)
    print("생성 요약")
    print("  대상 캠페인(matchable!=N):", [p["id"] for p in targets])
    print("  캠페인별 매칭 수:", per_policy)
    print("  고객 수:", len(members))
    print("=" * 60)


if __name__ == "__main__":
    main()
