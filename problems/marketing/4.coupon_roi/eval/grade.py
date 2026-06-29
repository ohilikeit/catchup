# -*- coding: utf-8 -*-
"""
쿠폰 성과 집계(ROI·적자) 과제 — 채점 스크립트 (집계/분석 아키타입)

사용:
  python eval/grade.py \
    --part1 데이터/쿠폰별_성과표_작성용.xlsx \
    --part2 데이터/전체요약_작성용.xlsx \
    --answer-dir eval/answer_key

자기채점(정답키 자기검증, 100점 기대):
  python eval/grade.py \
    --part1 eval/answer_key/coupon_roi_answer.xlsx \
    --part2 eval/answer_key/summary_answer.xlsx \
    --answer-dir eval/answer_key

배점: Part1 쿠폰별 집계표 55 / Part2 전체 요약 35 / 무결성 10.
채점 프리미티브:
  · norm_int  — 건수·합계·이익(사용건수/매출합계/할인합계/공헌이익/순이익, 총*·적자쿠폰수):
      정수 Exact(빈칸 보존, 빈칸≠0).
  · norm_float(tol=0.1) — 율(ROI, 전체ROI): |정답−제출| < 0.1 이면 정답(부동소수 잡음은 4자리 반올림 제거).
  · norm_channel/norm_deficit — 문자(주채널/적자여부): 정규화 후 Exact. 적자여부는 '적자'(별칭
      Y·손해·loss 등 흡수)/빈칸(흑자·N 등) 두 부류로 정규화. 주채널은 대소문자·공백·한글 별칭 흡수.
출력: 콘솔 표 + result.json.
"""
import argparse, json, re
from pathlib import Path
import openpyxl

# ---- 채점 필드 정의 --------------------------------------------------------
P1_INT_FIELDS   = ["사용건수", "매출합계", "할인합계", "공헌이익", "순이익"]
P1_FLOAT_FIELDS = ["ROI"]
P1_STR_FIELDS   = ["주채널", "적자여부"]
SUM_INT_FIELDS   = ["총사용건수", "총매출", "총할인", "총공헌이익", "총순이익", "적자쿠폰수"]
SUM_FLOAT_FIELDS = ["전체ROI"]
FLOAT_TOL = 0.1

CHANNELS = ["Organic Search", "Paid Search", "Email", "Direct",
            "Social", "Referral", "Display"]
# 주채널 동치 흡수(영문 채널명 정규화 + 흔한 한글 별칭)
CHANNEL_ALIAS = {
    "organicsearch": "organic search", "유기검색": "organic search", "자연검색": "organic search",
    "paidsearch": "paid search", "유료검색": "paid search", "검색광고": "paid search",
    "email": "email", "이메일": "email",
    "direct": "direct", "직접": "direct",
    "social": "social", "소셜": "social",
    "referral": "referral", "추천": "referral", "리퍼럴": "referral",
    "display": "display", "디스플레이": "display", "배너": "display",
}
DEFICIT_TRUE = {"적자", "deficit", "loss", "y", "손해", "마이너스"}


# ---- 정규화 / 비교 프리미티브 ----------------------------------------------
def norm_str(v):
    if v is None:
        return ""
    return re.sub(r"\s+", "", str(v)).strip()


def norm_int(v):
    """건수·합계·이익 셀: 빈칸 보존(0과 구분), 콤마·공백·원 제거 → 정수 문자열."""
    if v is None or str(v).strip() == "":
        return ""
    s = re.sub(r"[,\s원]", "", str(v))
    try:
        return str(int(float(s)))
    except ValueError:
        return str(v).strip()


def norm_float(v):
    """율 셀: 콤마·%·공백 제거 후 float. 빈칸/파싱불가 → None."""
    if v is None or str(v).strip() == "":
        return None
    s = re.sub(r"[,%\s]", "", str(v))
    try:
        return float(s)
    except ValueError:
        return None


def norm_channel(v):
    """주채널: 소문자·공백제거 후 별칭 흡수 → 정규 채널명(소문자) 또는 ''(빈칸)."""
    if v is None or str(v).strip() == "":
        return ""
    raw = str(v).strip().lower()
    key = re.sub(r"\s+", "", raw)
    if key in CHANNEL_ALIAS:
        return CHANNEL_ALIAS[key]
    return raw  # 알 수 없는 값은 원문(소문자)으로 비교 → 불일치


def norm_deficit(v):
    """적자여부: '적자'류 → '적자', 그 외(빈칸·흑자·정상·N) → ''."""
    if v is None or str(v).strip() == "":
        return ""
    key = re.sub(r"\s+", "", str(v).strip().lower())
    return "적자" if key in DEFICIT_TRUE else ""


def int_eq(a, s):
    return norm_int(a) == norm_int(s)


def float_eq(a, s, tol=FLOAT_TOL):
    """Computed ±ε. 정답키 ROI는 항상 숫자(미사용·할인0도 0.0)이므로 정답이 비면
       생성기 버그 → 오답 처리. 제출이 비었거나 |차|>=tol 이면 오답."""
    av = norm_float(a)
    sv = norm_float(s)
    if av is None or sv is None:
        return False
    return round(abs(av - sv), 4) < tol


def str_eq(field, a, s):
    if field == "주채널":
        return norm_channel(a) == norm_channel(s)
    return norm_deficit(a) == norm_deficit(s)


def load_sheet(path, sheet):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet in wb.sheetnames else wb[wb.sheetnames[-1]]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    head = [norm_str(h) for h in rows[0]]
    out = []
    for r in rows[1:]:
        if all(c is None or str(c).strip() == "" for c in r):
            continue
        out.append({head[i]: r[i] for i in range(len(head))})
    return out


# ---- Part 1: 쿠폰별 집계표 (키별 셀 Exact + ±ε + 문자) ---------------------
def p1key(r):
    return norm_str(r.get("쿠폰코드"))


def grade_part1(sub, ans):
    A = {p1key(r): r for r in ans}
    S = {p1key(r): r for r in sub if p1key(r)}
    fields = P1_INT_FIELDS + P1_FLOAT_FIELDS + P1_STR_FIELDS
    per_field = {f: [0, 0] for f in fields}     # [correct, total]
    detail = []
    for key, arow in A.items():
        srow = S.get(key, {})
        for f in P1_INT_FIELDS:
            ok = int_eq(arow.get(f), srow.get(f))
            per_field[f][1] += 1
            per_field[f][0] += 1 if ok else 0
            if not ok and len(detail) < 80:
                detail.append({"쿠폰": key, "필드": f,
                               "정답": norm_int(arow.get(f)), "제출": norm_int(srow.get(f))})
        for f in P1_FLOAT_FIELDS:
            ok = float_eq(arow.get(f), srow.get(f))
            per_field[f][1] += 1
            per_field[f][0] += 1 if ok else 0
            if not ok and len(detail) < 80:
                detail.append({"쿠폰": key, "필드": f,
                               "정답": norm_float(arow.get(f)), "제출": norm_float(srow.get(f))})
        for f in P1_STR_FIELDS:
            ok = str_eq(f, arow.get(f), srow.get(f))
            per_field[f][1] += 1
            per_field[f][0] += 1 if ok else 0
            if not ok and len(detail) < 80:
                detail.append({"쿠폰": key, "필드": f,
                               "정답": norm_str(arow.get(f)), "제출": norm_str(srow.get(f))})
    cells_ok = sum(v[0] for v in per_field.values())
    cells_tot = sum(v[1] for v in per_field.values())
    # 무결성: 유령 쿠폰(정답에 없는 쿠폰코드) + 제출 행의 음수 금액/건수
    # (순이익은 적자 쿠폰에서 정상적으로 음수이므로 음수 검사에서 제외)
    valid = set(A.keys())
    ghost = sorted(k for k in S if k not in valid)
    NONNEG_FIELDS = ["사용건수", "매출합계", "할인합계", "공헌이익"]
    neg_bad = []
    for key, srow in S.items():
        for f in NONNEG_FIELDS:
            ns = norm_int(srow.get(f))
            try:
                if ns != "" and int(ns) < 0:
                    neg_bad.append(key)
                    break
            except ValueError:
                continue
    return {
        "정확셀": cells_ok, "전체셀": cells_tot,
        "셀정확도": round(cells_ok / cells_tot, 4) if cells_tot else 0,
        "쿠폰총수": len(A),
        "필드별": {f: round(v[0] / v[1], 3) if v[1] else 0 for f, v in per_field.items()},
        "유령쿠폰": ghost[:50], "음수오류": sorted(set(neg_bad))[:50],
        "오답상세": detail,
    }


# ---- Part 2: 전체 요약 1행 (셀 Exact + ±ε) ---------------------------------
def grade_part2(sub, ans):
    arow = ans[0] if ans else {}
    srow = sub[0] if sub else {}
    fields = SUM_INT_FIELDS + SUM_FLOAT_FIELDS
    per_field = {f: 0 for f in fields}
    detail = []
    for f in SUM_INT_FIELDS:
        ok = int_eq(arow.get(f), srow.get(f))
        per_field[f] = 1 if ok else 0
        if not ok:
            detail.append({"필드": f, "정답": norm_int(arow.get(f)), "제출": norm_int(srow.get(f))})
    for f in SUM_FLOAT_FIELDS:
        ok = float_eq(arow.get(f), srow.get(f))
        per_field[f] = 1 if ok else 0
        if not ok:
            detail.append({"필드": f, "정답": norm_float(arow.get(f)), "제출": norm_float(srow.get(f))})
    cells_ok = sum(per_field.values())
    cells_tot = len(fields)
    return {
        "정확셀": cells_ok, "전체셀": cells_tot,
        "셀정확도": round(cells_ok / cells_tot, 4) if cells_tot else 0,
        "필드별": dict(per_field),
        "오답상세": detail,
    }


# ---- 배점 -----------------------------------------------------------------
# Part1 집계표 셀정확도 55 / Part2 요약 셀정확도 35 / 무결성 10.
def score(p1, p2):
    s1 = p1["셀정확도"] * 55
    s2 = p2["셀정확도"] * 35
    viol = len(p1["유령쿠폰"]) + len(p1["음수오류"])
    s3 = max(0, 10 - viol * 0.5)
    return {"Part1_집계표(55)": round(s1, 1), "Part2_요약(35)": round(s2, 1),
            "무결성(10)": round(s3, 1), "총점(100)": round(s1 + s2 + s3, 1)}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--part1", required=True)
    ap.add_argument("--part2", required=True)
    ap.add_argument("--answer-dir", required=True)
    args = ap.parse_args()

    adir = Path(args.answer_dir)
    p1_sub = load_sheet(args.part1, "집계표")
    p1_ans = load_sheet(adir / "coupon_roi_answer.xlsx", "집계표")
    p2_sub = load_sheet(args.part2, "전체요약")
    p2_ans = load_sheet(adir / "summary_answer.xlsx", "전체요약")

    r1 = grade_part1(p1_sub, p1_ans)
    r2 = grade_part2(p2_sub, p2_ans)
    sc = score(r1, r2)

    print("=" * 60); print(" 채점 결과 (쿠폰 성과 집계 · ROI/적자)")
    print("=" * 60)
    print(f"[Part1] 셀정확도 {r1['셀정확도']*100:.1f}% ({r1['정확셀']}/{r1['전체셀']}) | "
          f"쿠폰 {r1['쿠폰총수']}개")
    print(f"        필드별: {r1['필드별']}")
    if r1["유령쿠폰"]: print(f"        ⚠ 유령 쿠폰: {len(r1['유령쿠폰'])}건")
    if r1["음수오류"]: print(f"        ⚠ 음수 금액/건수: {len(r1['음수오류'])}건")
    print(f"[Part2] 셀정확도 {r2['셀정확도']*100:.1f}% ({r2['정확셀']}/{r2['전체셀']})")
    print(f"        필드별: {r2['필드별']}")
    if r2["오답상세"]:  print(f"        오답: {r2['오답상세']}")
    print("-" * 60)
    for k, v in sc.items():
        print(f"  {k:<20} {v}")
    print("=" * 60)

    out = {"part1": r1, "part2": r2, "score": sc}
    Path("result.json").write_text(json.dumps(out, ensure_ascii=False, indent=2))
    print("→ result.json 저장")


if __name__ == "__main__":
    main()
