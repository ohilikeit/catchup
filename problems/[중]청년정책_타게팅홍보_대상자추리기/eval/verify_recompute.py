# -*- coding: utf-8 -*-
"""S10 게이트 1 — 독립 재계산 (Independent Recompute).

build_dataset.py를 import·복붙하지 않고, 학생에게 보이는 규칙 서술만 근거로
(회원, 정책) 자격 판정을 처음부터 재구현해 정답키와 대조한다.

규칙 근거(이것만 보고 구현했음 — build_dataset.py 코드는 보지 않음):
  * 제출물/최종제출.xlsx '안내' 시트 (자격 판정 규칙 9축 + 확인 불가 정책)
  * 제출물/최종제출.xlsx '소득기준표' 시트 (가구원수별 중위소득 % 상한)
  * 문제.md (기준일 2026-06-28, 날짜 안 봄, 판별불가 정책 추천 금지)
  * eval/answer_key/INPUT_GUIDE.md (policies.csv 컬럼 의미·region 3형식·빈칸=무관)

입력(직접 파싱): eval/answer_key/policies.csv + 데이터/회원명단.xlsx
대조 대상:      eval/answer_key/2_추천리스트_정답.xlsx (Part2, (회원ID,정책ID) 집합)
                eval/answer_key/1_정책정리표_정답.xlsx (Part1, 자격 칸 셀 단위)

통과 기준: Part2 extra=0 & missing=0, Part1 불일치 셀 0건, 경계 표본(B1~B8)이
실제 모집단에 존재하고 올바른 방향으로 갈림.

실행: python3 eval/verify_recompute.py   (과제 루트에서)
"""
import csv
import re
import sys
from datetime import date
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent  # 과제 루트
ASSIGN_DATE = date(2026, 6, 28)  # 문제.md: "오늘은 2026년 6월 28일입니다"

# ---------------------------------------------------------------- 입력 로드

POLICY_COLS = [
    "policy_id", "name", "field", "poster_file", "region", "age_min",
    "age_max", "employment", "company_size", "income_max_month",
    "income_median_pct", "income_year_max", "no_house", "enroll_status",
    "marital", "matchable", "apply_start", "apply_end",
    "alt_age_max", "alt_income_median_pct",
]


def int_or_none(v):
    """빈칸=None(무관), 0은 0 그대로 — '0과 빈칸은 다르다'(안내 시트)."""
    if v is None:
        return None
    s = str(v).strip().replace(",", "")
    if s == "":
        return None
    return int(s)


def load_policies():
    path = ROOT / "eval" / "answer_key" / "policies.csv"
    with open(path, newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        assert reader.fieldnames == POLICY_COLS, f"헤더 불일치: {reader.fieldnames}"
        pols = []
        for row in reader:
            p = {k: (row[k] or "").strip() for k in POLICY_COLS}
            for k in ("age_min", "age_max", "income_max_month",
                      "income_median_pct", "income_year_max",
                      "alt_age_max", "alt_income_median_pct"):
                p[k] = int_or_none(p[k])
            pols.append(p)
    return pols


def load_members():
    wb = openpyxl.load_workbook(ROOT / "데이터" / "회원명단.xlsx", data_only=True)
    ws = wb["회원"]
    hdr = [c.value for c in ws[1]]
    members = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        members.append(dict(zip(hdr, row)))
    return members


def load_income_table():
    """제출 양식 '소득기준표' 시트 — 가구원수(1~7)별 기준중위소득(월,원).

    8인 이상: '1인 늘 때마다 (7인-6인)만큼 가산'(시트 말미 각주).
    % 상한 열은 개수·순서가 고정이 아니다(정책마다 다른 %가 쓰일 수 있음 — 예:
    기본 트랙 150%/200% + OR 분기 정책의 대안 트랙 100% 등). 헤더 문구
    '중위소득{N}%_월상한(원)'에서 N을 직접 파싱해 동적으로 읽는다(열 위치 하드코딩 금지).
    """
    wb = openpyxl.load_workbook(ROOT / "제출물" / "최종제출.xlsx", data_only=True)
    ws = wb["소득기준표"]
    hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    pct_col = {}
    for i, h in enumerate(hdr):
        mm = re.match(r"중위소득(\d+)%_월상한", str(h or ""))
        if mm:
            pct_col[i] = int(mm.group(1))
    base, caps = {}, {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not isinstance(row[0], int):
            continue
        n = row[0]
        base[n] = int(row[1])
        for i, pct in pct_col.items():
            caps[(n, pct)] = int(row[i])
    return base, caps


# ------------------------------------------------------- 규칙 재구현 (안내 시트)

SEOUL_GU = [
    "종로구", "중구", "용산구", "성동구", "광진구", "동대문구", "중랑구",
    "성북구", "강북구", "도봉구", "노원구", "은평구", "서대문구", "마포구",
    "양천구", "강서구", "구로구", "금천구", "영등포구", "동작구", "관악구",
    "서초구", "강남구", "송파구", "강동구",
]
# 비서울 시도 키워드 → 정식 명칭 (비서울을 먼저 판정: '분당구' 같은 타지역 구 오인 방지)
SIDO_KEYS = [
    ("인천", "인천광역시"), ("경기", "경기도"), ("부산", "부산광역시"),
    ("대구", "대구광역시"), ("대전", "대전광역시"), ("광주광역시", "광주광역시"),
    ("울산", "울산광역시"), ("세종", "세종특별자치시"), ("강원", "강원특별자치도"),
    ("충청북", "충청북도"), ("충북", "충청북도"), ("충청남", "충청남도"),
    ("충남", "충청남도"), ("전라북", "전북특별자치도"), ("전북", "전북특별자치도"),
    ("전라남", "전라남도"), ("전남", "전라남도"), ("경상북", "경상북도"),
    ("경북", "경상북도"), ("경상남", "경상남도"), ("경남", "경상남도"),
    ("제주", "제주특별자치도"),
]


def parse_addr(addr):
    """회원 주소 → (시도, 서울자치구|None). 표기 4형('서울특별시 관악구 봉천동'/
    '서울 관악구'/'서울시 관악구'/'관악구 신림동')을 모두 흡수 — 안내 시트
    "회원 주소 표기가 제각각이어도 시·도와 자치구로 판단"의 독립 구현."""
    s = str(addr).strip()
    toks = s.split()
    # 1) 비서울 시도 먼저
    for key, full in SIDO_KEYS:
        if toks and toks[0].startswith(key):
            return full, None  # 정책 자치구 요건은 서울만 존재
    # 2) 서울 (명시 or 자치구 단독 표기)
    if toks and toks[0].startswith("서울"):
        gu = next((t for t in toks[1:] if t in SEOUL_GU), None)
        return "서울특별시", gu
    if toks and toks[0] in SEOUL_GU:
        return "서울특별시", toks[0]
    raise ValueError(f"주소 시도 판정 불가(fail-loud): {addr!r}")


def full_age(birth, on=ASSIGN_DATE):
    """만 나이 — 기준일에 생일이 아직 안 지났으면 한 살 뺀다(안내 시트)."""
    if isinstance(birth, date):
        b = birth
    else:
        y, m, d = map(int, str(birth).strip()[:10].split("-"))
        b = date(y, m, d)
    age = on.year - b.year
    if (on.month, on.day) < (b.month, b.day):
        age -= 1
    return age


def median_cap(hh, pct, base, caps):
    """가구원수 hh, 중위소득 pct% 상한(원). 표에 있으면 표값, 없으면
    기준중위소득×%÷100 원단위 반올림(안내 시트 '직접 계산' 규칙)."""
    if (hh, pct) in caps:
        return caps[(hh, pct)]
    if hh in base:
        b = base[hh]
    elif hh > 7:
        b = base[7] + (hh - 7) * (base[7] - base[6])  # 8인 이상 가산 각주
    else:
        raise ValueError(f"가구원수 {hh} 기준중위소득 없음")
    return int((Decimal(b) * Decimal(pct) / Decimal(100)).quantize(
        Decimal("1"), rounding=ROUND_HALF_UP))


def eligible(m, p, base, caps, skip_axis=None):
    """안내 시트 자격 판정 규칙의 독립 재구현. 모든 축 AND.
    skip_axis: 경계 표본 점검용 — 특정 축을 제외한 나머지 축 통과 여부."""
    # ★확인 불가 정책: 아무에게도 추천하지 않는다
    if p["matchable"] == "N":
        return False
    ax = {}
    # 나이 (양끝 포함). 복수 자격 트랙(OR 분기): alt_age_max가 있으면 기본 트랙
    # (age_min~age_max) OR 대안 트랙(나이<=alt_age_max AND 가구중위소득<=alt_income_median_pct%)
    # 으로 판정한다. alt_age_max가 없는 정책은 기존과 완전히 동일(단일 age_max).
    # (규칙 근거: 이 정책의 포스터에 "만 19~39세, 단 기준중위소득 100% 이하 저소득
    # 가구는 만45세까지"처럼 조건부 완화 조항이 서술돼 있다고 가정하고 policies.csv의
    # alt_age_max/alt_income_median_pct 값을 그대로 읽어 재구현했다 — build_dataset.py는
    # import하지 않음.)
    a = full_age(m["생년월일"])
    ok = True
    if p["age_min"] is not None and a < p["age_min"]:
        ok = False
    elif p["alt_age_max"] is not None:
        base_age_ok = p["age_max"] is None or a <= p["age_max"]
        if not base_age_ok:
            if p["alt_income_median_pct"] is None or a > p["alt_age_max"]:
                ok = False
            else:
                cap = median_cap(int(m["가구원수"]), p["alt_income_median_pct"], base, caps)
                if int(m["월소득(원)"]) > cap:
                    ok = False
    elif p["age_max"] is not None and a > p["age_max"]:
        ok = False
    ax["age"] = ok
    # 거주지: 전국 / 시도 / 시도+자치구
    sido, gu = parse_addr(m["거주지"])
    r = p["region"]
    if r == "전국":
        ax["region"] = True
    else:
        rt = r.split()
        if len(rt) == 1:
            ax["region"] = (sido == rt[0])
        elif len(rt) == 2:
            ax["region"] = (sido == rt[0] and gu == rt[1])
        else:
            raise ValueError(f"region 형식 오류(fail-loud): {r!r}")
    # 취업: 미취업={미취업,단기근로} / 재직={재직,단기근로} / 무관·빈칸=모두
    emp = p["employment"]
    st = str(m["취업상태"]).strip()
    if emp in ("", "무관"):
        ax["emp"] = True
    elif emp == "미취업":
        ax["emp"] = st in ("미취업", "단기근로")
    elif emp == "재직":
        ax["emp"] = st in ("재직", "단기근로")
    else:
        raise ValueError(f"employment 허용값 아님(fail-loud): {emp!r}")
    # 기업규모: 값 있으면 정확 일치
    ax["comp"] = (p["company_size"] == "" or
                  str(m["기업규모"]).strip() == p["company_size"])
    # 월소득(본인 상한): 이하(<=), 빈칸=무관
    ax["inc_m"] = (p["income_max_month"] is None or
                   int(m["월소득(원)"]) <= p["income_max_month"])
    # 월소득(중위소득 %): 가구원수별 상한 이하(<=)
    if p["income_median_pct"] is None:
        ax["inc_pct"] = True
    else:
        cap = median_cap(int(m["가구원수"]), p["income_median_pct"], base, caps)
        ax["inc_pct"] = int(m["월소득(원)"]) <= cap
    # 연소득: 이하(<=), 빈칸=무관
    ax["inc_y"] = (p["income_year_max"] is None or
                   int(m["연소득(원)"]) <= p["income_year_max"])
    # 무주택: '필요'(csv Y)면 주택소유여부=아니오
    ax["house"] = (p["no_house"] == "" or
                   str(m["주택소유여부"]).strip() == "아니오")
    # 학력: 재학={재학,휴학} / 졸업={졸업} / 빈칸=모두
    es = p["enroll_status"]
    rs = str(m["재학상태"]).strip()
    if es == "":
        ax["enroll"] = True
    elif es == "재학":
        ax["enroll"] = rs in ("재학", "휴학")
    elif es == "졸업":
        ax["enroll"] = rs == "졸업"
    else:
        raise ValueError(f"enroll_status 허용값 아님(fail-loud): {es!r}")
    # 혼인: 신혼부부={기혼} / 미혼={미혼} / 빈칸=모두
    ma = p["marital"]
    mv = str(m["혼인여부"]).strip()
    if ma == "":
        ax["marital"] = True
    elif ma == "신혼부부":
        ax["marital"] = mv == "기혼"
    elif ma == "미혼":
        ax["marital"] = mv == "미혼"
    else:
        raise ValueError(f"marital 허용값 아님(fail-loud): {ma!r}")
    if skip_axis is not None:
        return all(v for k, v in ax.items() if k != skip_axis)
    return all(ax.values())


# ---------------------------------------------------------------- 대조

def recompute_part2(members, policies, base, caps):
    got = set()
    for p in policies:
        for m in members:
            if eligible(m, p, base, caps):
                got.add((m["회원ID"], p["policy_id"]))
    return got


def load_answer_part2():
    wb = openpyxl.load_workbook(
        ROOT / "eval" / "answer_key" / "2_추천리스트_정답.xlsx", data_only=True)
    ws = wb["타게팅"]
    ans = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        ans.add((str(row[0]).strip(), str(row[1]).strip()))
    return ans


# Part1: csv 자격값 → 정답표 표현으로 사상 (안내 시트의 허용값 서술 근거:
#   취업 '무관이면 빈칸', 무주택 허용값 '필요' 또는 빈칸)
P1_MAP = [
    ("거주지요건", lambda p: p["region"] or None),
    ("연령_최소", lambda p: p["age_min"]),
    ("연령_최대", lambda p: p["age_max"]),
    ("취업요건", lambda p: None if p["employment"] in ("", "무관") else p["employment"]),
    ("기업규모요건", lambda p: p["company_size"] or None),
    ("월소득상한(원)", lambda p: p["income_max_month"]),
    ("월소득_중위소득기준(%)", lambda p: p["income_median_pct"]),
    ("연소득상한(원)", lambda p: p["income_year_max"]),
    ("무주택요건", lambda p: "필요" if p["no_house"] == "Y" else None),
    ("학력요건", lambda p: p["enroll_status"] or None),
    ("혼인요건", lambda p: p["marital"] or None),
]


def compare_part1(policies):
    wb = openpyxl.load_workbook(
        ROOT / "eval" / "answer_key" / "1_정책정리표_정답.xlsx", data_only=True)
    ws = wb["정책정리"]
    hdr = [c.value for c in ws[1]]
    rows = {str(r[0]).strip(): dict(zip(hdr, r))
            for r in ws.iter_rows(min_row=2, values_only=True) if r[0]}
    mism = []
    for p in policies:
        pid = p["policy_id"]
        if pid not in rows:
            mism.append((pid, "<행 없음>", None, None))
            continue
        for col, fn in P1_MAP:
            exp = fn(p)
            gotv = rows[pid].get(col)
            gv = None if gotv in (None, "") else (
                int(str(gotv).replace(",", "")) if isinstance(exp, int) and
                str(gotv).replace(",", "").lstrip("-").isdigit() else
                str(gotv).strip())
            if isinstance(exp, int):
                same = (gv == exp)
            else:
                same = ((exp or None) == (gv or None))
            if not same:
                mism.append((pid, col, exp, gotv))
    return mism, rows


# ------------------------------------------------------------ 경계 표본 점검

def boundary_report(members, policies, base, caps, match):
    """independent_recompute.md '점검할 경계 표본 목록' — 각 경계가 실제로
    존재하고 올바른 방향으로 갈렸는지 출력. 반환: (검사수, 실패목록)."""
    checks, fails = 0, []
    bym = {m["회원ID"]: m for m in members}

    def note(name, cond, detail):
        nonlocal checks
        checks += 1
        mark = "OK " if cond else "FAIL"
        print(f"  [{mark}] {name}: {detail}")
        if not cond:
            fails.append(name)

    for p in policies:
        if p["matchable"] == "N":
            continue
        pid = p["policy_id"]
        # 나이 ±1 (다른 축 통과자 중에서)
        if p["age_min"] is not None or p["age_max"] is not None:
            pool = [m for m in members if eligible(m, p, base, caps, skip_axis="age")]
            buckets = {}
            for m in pool:
                buckets.setdefault(full_age(m["생년월일"]), []).append(m)
            # alt_age_max가 있는 정책은 "max+1" 경계에서 대안 트랙(소득에 따라 통과/탈락이
            # 갈림)이 개입해 단일 이진 기대값("전원 제외")이 성립하지 않는다 — 이 경계는
            # 아래 B9 블록에서 소득을 통제한 전용 표본으로 별도 검증한다.
            age_labels = [
                ("min-1", None if p["age_min"] is None else p["age_min"] - 1, False),
                ("min", p["age_min"], True),
                ("max", p["age_max"], True),
            ]
            if p["alt_age_max"] is None:
                age_labels.append(
                    ("max+1", None if p["age_max"] is None else p["age_max"] + 1, False))
            for label, tgt, should in age_labels:
                if tgt is None:
                    continue
                ms = buckets.get(tgt, [])
                if not ms:
                    note(f"B1 {pid} 나이 {label}", False, f"만나이={tgt} 표본 없음")
                    continue
                ok = all(((m["회원ID"], pid) in match) == should for m in ms)
                note(f"B1 {pid} 나이 {label}", ok,
                     f"만나이={tgt} {len(ms)}명 → {'포함' if should else '제외'}")
                # B2: 생일 미경과(7~12월생) 표본
                if should is True and tgt == p["age_max"]:
                    un = [m for m in ms
                          if int(str(m["생년월일"])[5:7]) >= 7]
                    note(f"B2 {pid} 만나이 미경과(max 포함)", bool(un) and all(
                        (m["회원ID"], pid) in match for m in un),
                        f"7~12월생 만나이={tgt} {len(un)}명 포함" if un
                        else "표본 없음")
                if should is False and p["age_min"] is not None and tgt == p["age_min"] - 1:
                    un = [m for m in ms if int(str(m["생년월일"])[5:7]) >= 7]
                    note(f"B2 {pid} 만나이 미경과(min-1 제외)", bool(un) and all(
                        (m["회원ID"], pid) not in match for m in un),
                        f"7~12월생 만나이={tgt} {len(un)}명 제외" if un
                        else "표본 없음")
            # B9: 복수 자격 트랙(OR 분기) 경계 — alt_age_max가 있는 정책 전용(generic).
            #   기본상한/기본상한+1은 소득을 통제(고소득)해 대안 트랙 개입 없이 순수
            #   기본 트랙 판정만 관찰하고, 대안 트랙 경계는 기본상한 초과~alt_age_max
            #   이내 나이에서 가구원수 1·3의 중위소득 상한/상한+1로 구제/최종탈락을 본다.
            if p["alt_age_max"] is not None and p["age_max"] is not None:
                alt_pct = p["alt_income_median_pct"]
                base_max = p["age_max"]
                cap1 = median_cap(1, alt_pct, base, caps) if alt_pct is not None else None
                hi_at_base = [mm for mm in buckets.get(base_max, [])
                              if cap1 is not None and int(mm["월소득(원)"]) > cap1]
                note(f"B9 {pid} 기본상한({base_max}) 고소득자 포함(기본트랙)",
                     bool(hi_at_base) and all(
                         (mm["회원ID"], pid) in match for mm in hi_at_base),
                     f"{len(hi_at_base)}명 나이={base_max} 고소득 → 포함" if hi_at_base
                     else "표본 없음")
                hi_over = [mm for mm in buckets.get(base_max + 1, [])
                           if cap1 is not None and int(mm["월소득(원)"]) > cap1]
                note(f"B9 {pid} 기본상한+1({base_max+1}) 고소득자 제외(대안도 탈락)",
                     bool(hi_over) and all(
                         (mm["회원ID"], pid) not in match for mm in hi_over),
                     f"{len(hi_over)}명 나이={base_max+1} 고소득 → 제외" if hi_over
                     else "표본 없음")
                if alt_pct is not None:
                    mid_ages = [a for a in buckets if base_max < a <= p["alt_age_max"]]
                    for hh in (1, 3):
                        capv = median_cap(hh, alt_pct, base, caps)
                        rescued = [mm for a in mid_ages for mm in buckets[a]
                                   if int(mm["가구원수"]) == hh and int(mm["월소득(원)"]) == capv]
                        rejected = [mm for a in mid_ages for mm in buckets[a]
                                    if int(mm["가구원수"]) == hh
                                    and int(mm["월소득(원)"]) == capv + 1]
                        note(f"B9 {pid} 대안트랙 구제(가구{hh}인, 상한={capv})",
                             bool(rescued) and all(
                                 (mm["회원ID"], pid) in match for mm in rescued),
                             f"{len(rescued)}명 기본상한 초과·대안상한 이하 → 포함" if rescued
                             else "표본 없음")
                        note(f"B9 {pid} 대안트랙 소득초과 최종탈락(가구{hh}인)",
                             bool(rejected) and all(
                                 (mm["회원ID"], pid) not in match for mm in rejected),
                             f"{len(rejected)}명 대안상한 초과 → 제외" if rejected
                             else "표본 없음")
        # B3 소득 경계
        for axis, mcol, capv in (
                ("월소득상한", "월소득(원)", p["income_max_month"]),
                ("연소득상한", "연소득(원)", p["income_year_max"])):
            if capv is None:
                continue
            pool = [m for m in members
                    if eligible(m, p, base, caps,
                                skip_axis="inc_m" if axis == "월소득상한" else "inc_y")]
            at = [m for m in pool if int(m[mcol]) == capv]
            over = [m for m in pool if int(m[mcol]) == capv + 1]
            note(f"B3 {pid} {axis}={capv} (통과)", bool(at) and all(
                (m["회원ID"], pid) in match for m in at),
                f"{len(at)}명 == 상한 → 포함" if at else "표본 없음")
            note(f"B3 {pid} {axis}+1 (탈락)", bool(over) and all(
                (m["회원ID"], pid) not in match for m in over),
                f"{len(over)}명 == 상한+1 → 제외" if over else "표본 없음")
        if p["income_median_pct"] is not None:
            pool = [m for m in members if eligible(m, p, base, caps, skip_axis="inc_pct")]
            at = [m for m in pool if int(m["월소득(원)"]) ==
                  median_cap(int(m["가구원수"]), p["income_median_pct"], base, caps)]
            over = [m for m in pool if int(m["월소득(원)"]) ==
                    median_cap(int(m["가구원수"]), p["income_median_pct"], base, caps) + 1]
            note(f"B3 {pid} 중위{p['income_median_pct']}% 상한 (통과)",
                 bool(at) and all((m["회원ID"], pid) in match for m in at),
                 f"{len(at)}명 == 가구원수별 상한 → 포함" if at else "표본 없음")
            note(f"B3 {pid} 중위{p['income_median_pct']}% 상한+1 (탈락)",
                 bool(over) and all((m["회원ID"], pid) not in match for m in over),
                 f"{len(over)}명 == 상한+1 → 제외" if over else "표본 없음")
        # B8 취업 예외(단기근로)
        if p["employment"] in ("미취업", "재직"):
            pool = [m for m in members if eligible(m, p, base, caps)
                    and str(m["취업상태"]).strip() == "단기근로"]
            note(f"B8 {pid} 단기근로 통과({p['employment']} 요건)", bool(pool),
                 f"{len(pool)}명 매칭")
        # 무주택/학력/혼인 탈락 표본
        for axis, cond, want in (
                ("house", p["no_house"] == "Y", "주택소유=예 제외"),
                ("enroll", p["enroll_status"] != "", f"학력≠{p['enroll_status']} 제외"),
                ("marital", p["marital"] != "", f"혼인≠요건 제외")):
            if not cond:
                continue
            pool = [m for m in members
                    if eligible(m, p, base, caps, skip_axis=axis)
                    and not eligible(m, p, base, caps)]
            note(f"B5/B6류 {pid} {want}", bool(pool) and all(
                (m["회원ID"], pid) not in match for m in pool),
                f"{len(pool)}명 그 축으로만 탈락")

    # B4 지역 표기 4형 (자치구 한정 정책)
    for p in policies:
        if p["matchable"] == "N" or len(p["region"].split()) != 2:
            continue
        pid = p["policy_id"]
        gu = p["region"].split()[1]
        forms = {}
        for m in members:
            a = str(m["거주지"])
            if gu in a:
                t0 = a.split()[0]
                key = ("서울특별시" if t0.startswith("서울특별시") else
                       "서울시" if t0.startswith("서울시") else
                       "서울" if t0.startswith("서울") else "자치구단독")
                forms.setdefault(key, []).append(m)
        note(f"B4 {pid} {gu} 표기형 다양성", len(forms) >= 4,
             f"표기형 {sorted(forms)} ({ {k: len(v) for k, v in sorted(forms.items())} })")
        same = [m for k, v in forms.items() for m in v
                if eligible(m, p, base, caps, skip_axis="region")]
        ok = all(eligible(m, p, base, caps) for m in same)
        note(f"B4 {pid} 표기 불문 동일 판정", ok,
             f"{len(same)}명 (타 축 통과자) 모두 지역 통과")
        # 같은 시도 다른 자치구 탈락
        other = [m for m in members
                 if eligible(m, p, base, caps, skip_axis="region")
                 and parse_addr(m["거주지"])[0] == "서울특별시"
                 and parse_addr(m["거주지"])[1] != gu]
        note(f"B4 {pid} 타 자치구 탈락", bool(other) and all(
            (m["회원ID"], pid) not in match for m in other),
            f"서울 타구 {len(other)}명 제외")

    # 전국 정책: 비서울 시도 회원 통과 존재
    for p in policies:
        if p["matchable"] == "N" or p["region"] != "전국":
            continue
        pid = p["policy_id"]
        non_seoul = [m for m in members if eligible(m, p, base, caps)
                     and parse_addr(m["거주지"])[0] != "서울특별시"]
        sidos = sorted({parse_addr(m["거주지"])[0] for m in non_seoul})
        note(f"전국 {pid} 비서울 통과", bool(non_seoul),
             f"{len(non_seoul)}명 ({', '.join(sidos)})")

    # B6 판별불가 0건
    for p in policies:
        if p["matchable"] != "N":
            continue
        pid = p["policy_id"]
        n = sum(1 for mid, x in match if x == pid)
        note(f"B6 {pid} matchable=N 매칭 0건", n == 0, f"재계산 {n}건")

    # 빈칸=무관: 소득 축이 전부 빈 정책에서 고소득자도 통과하는 표본
    for p in policies:
        if p["matchable"] == "N":
            continue
        if (p["income_max_month"] is None and p["income_median_pct"] is None
                and p["income_year_max"] is None):
            rich = [m for m in members if eligible(m, p, base, caps)
                    and int(m["월소득(원)"]) >= 4_000_000]
            note(f"B5 {p['policy_id']} 소득 빈칸=무관", bool(rich),
                 f"월소득≥400만 {len(rich)}명 통과 (소득 무관 확인)")
    return checks, fails


def main():
    policies = load_policies()
    members = load_members()
    base, caps = load_income_table()
    print(f"입력: 정책 {len(policies)}건, 회원 {len(members)}명, "
          f"기준일 {ASSIGN_DATE} (문제.md)")

    # 소득기준표 자기일관성: 표의 %상한 == 기준중위소득×%÷100 반올림
    for (hh, pct), v in caps.items():
        calc = int((Decimal(base[hh]) * pct / 100).quantize(
            Decimal("1"), rounding=ROUND_HALF_UP))
        assert calc == v, f"소득기준표 불일치: {hh}인 {pct}% 표={v} 계산={calc}"
    print("소득기준표 시트 ↔ '기준중위소득×%÷100 원단위 반올림' 일치 확인")

    match = recompute_part2(members, policies, base, caps)
    ans = load_answer_part2()
    extra = match - ans
    missing = ans - match
    print(f"\n== Part2 (회원ID,정책ID) 집합 대조 ==")
    print(f"재계산 {len(match)}건 vs 정답키 {len(ans)}건")
    print(f"extra(재계산에만 있음)  : {len(extra)}건 {sorted(extra)[:10]}")
    print(f"missing(정답키에만 있음): {len(missing)}건 {sorted(missing)[:10]}")

    mism, _ = compare_part1(policies)
    print(f"\n== Part1 정책정리표 셀 대조 ==")
    print(f"불일치 셀: {len(mism)}건")
    for row in mism[:20]:
        print("  ", row)

    print(f"\n== 경계 표본 점검 (independent_recompute.md B1~B8) ==")
    checks, bfails = boundary_report(members, policies, base, caps, match)
    print(f"\n경계 점검 {checks}건 중 실패 {len(bfails)}건: {bfails}")

    ok = not extra and not missing and not mism and not bfails
    print(f"\n===== 독립 재계산 결과: {'PASS' if ok else 'FAIL'} =====")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
