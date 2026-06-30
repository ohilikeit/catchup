# -*- coding: utf-8 -*-
"""
계약 기한 관리 — 계약서에서 기한·금액을 추출해 계약대장을 채우고, 만료·갱신 대상을 추리는
추출/파싱 + 검증/감사 과제 — 데이터셋/정답 생성기 (단일 입력 = contracts.csv)

[아키타입] 검증/감사 (Set-F1 채점) + 추출/파싱(셀 Exact). "비정형 계약서 → 구조화 추출 →
  규칙 적용 → 대상 집합". 정답이 곧 규칙.
  Part1: 제각각 양식의 계약서 PDF에서 기한·금액·자동갱신 조항을 읽어 계약대장(정형표)을 채운다.
  Part2: 본인이 Part1에서 채운 계약대장에 '갱신 규칙'을 적용해 만료/갱신 대상만 골라낸다. 정상은
    올리지 않는다. (통합형 — 추출이 틀리면 그 계약의 만료 판정도 함께 틀린다. Part2 정답키는 늘
    '진짜 정답 계약대장' 기준이므로, 끝까지 정확히 맞췄는지를 본다.)

[입력]
  eval/answer_key/contracts.csv   ★ 시나리오 명세(단일 진실 소스). 헤더(정확히):
    case_id,contract_id,계약명,거래상대방,계약유형,계약기간_개월,잔여일수,계약금액,
    자동갱신,갱신통지일,금액표기변형,날짜표기변형,종료일미명시,금액미기재
    · 계약 키 = contract_id (유일).
    · 잔여일수(int) : 종료일 = ASSIGN_DATE + 잔여일수 (음수 = 이미 만료). ← 날짜 결정성의 축.
    · 계약기간_개월(int) : 시작일 = 종료일 - 계약기간_개월 (말일 보정).
    · 계약금액(int, 빈칸 가능) : 금액미기재=Y면 비운다.
    · 자동갱신(Y/N) : Y면 '만료 N일 전 통지 없으면 자동 갱신' 조항. N이면 통지일 무의미(0).
    · 갱신통지일(int) : 자동갱신=Y일 때만 의미. 통지마감일 = 종료일 - 갱신통지일.
    · 금액표기변형/날짜표기변형(Y/N) : 계약서에서 금액·날짜를 까다로운 표기로 렌더(FP 함정,
      정규화하면 같은 값 → 위반 아님).
    · 종료일미명시(Y/N) : Y면 계약서에 종료일을 안 쓰고 '시작일 + 기간(개월)'만 → 학생이 계산.
    · 금액미기재(Y/N) : Y면 계약서에 금액 조항 없음 → 계약대장 금액칸은 빈칸이 정답.

[출력]
  데이터/1_계약서/<contract_id>.pdf          계약서 PDF (양식 4종 순환 + 표기변형) — 학생 입력
  데이터/1_계약대장_제출용.xlsx              Part1 제출 양식(식별칸=계약번호만, 빈 표 + '안내')
                                            ★Part2 입력 = 학생이 직접 채운 이 파일(별도 참고용 없음)
  데이터/2_만료대상_제출용.xlsx              Part2 제출 양식(빈 표 + '안내' + 예시행)
  eval/answer_key/1_계약대장_정답.xlsx       Part1 정답(계약대장)
  eval/answer_key/2_만료대상_정답.xlsx       Part2 정답(만료·갱신 대상표 — 대상 전부)

[갱신 규칙 (rulebook — 안내시트에 명시)]
  기준일 ASSIGN_DATE = 2026-06-28. 잔여일수 = (종료일 - 기준일).
  · 만료경과 : 종료일 < 기준일 (잔여 < 0).
  · 만료임박 : 0 <= 잔여 <= 60 (오늘 포함 60일 이내 만료).
  · 자동연장주의 : 자동갱신=Y 이고 (아직 만료 안 됨) 이고 통지마감일까지 잔여 <= 30
        (통지마감잔여 = 잔여 - 갱신통지일). 통지기한이 임박/경과해 손 안 쓰면 자동 연장됨.
  · 한 계약이 복수 대상 가능(예: 만료임박 + 자동연장주의 → 대상 행 2개). 어디에도 안 들면 안 올림.

[고정 기준일]
  ASSIGN_DATE = 2026-06-28 (종료일·시작일 합성 기준). datetime.now() 금지.
  모든 합성(양식·날짜표기·금액표기 선택)은 행 순서 _seq 로 결정 — 무작위 없음.
  contracts.csv만 바꿔 재실행하면 PDF·양식·정답이 항상 정합 재생성.

[정답 도출 철학] 시나리오 플래그를 직접 읽어 정답을 만들지 않는다 — 합성한 계약(종료일·자동갱신·
  통지일)에 갱신 규칙을 *다시 적용*해 대상을 도출한다(derive_targets). cross_check 는 종료일
  산술(end_date = 기준일 + 잔여일수)이 깨지지 않았는지 보는 라운드트립 자기검증이다.
  ★ 생성기·채점기 공통 버그를 잡는 진짜 독립검증은 eval/verify_recompute.py(이 생성기를
  import하지 않는 제3 구현)가 수행한다 — 발행 게이트는 그 통과를 필수로 한다.
"""
import csv
import sys
import datetime
import calendar
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
)

# ----------------------------------------------------------------------------
# 경로 / 기준일 / 시드 (결정성 3종의 ①②)
# ----------------------------------------------------------------------------
ROOT = Path(__file__).resolve().parents[2]          # .../3.contract_expiry_tracking
DATA = ROOT / "데이터"
CONTRACTS_DIR = DATA / "1_계약서"
ANS = ROOT / "eval" / "answer_key"
CSV_PATH = ANS / "contracts.csv"
DATA.mkdir(parents=True, exist_ok=True)
CONTRACTS_DIR.mkdir(parents=True, exist_ok=True)
ANS.mkdir(parents=True, exist_ok=True)

ASSIGN_DATE = datetime.date(2026, 6, 28)             # 고정 기준일 (무작위 미사용)

# ── 갱신 규칙(rulebook) 상수 ────────────────────────────────────────────────
SOON_DAYS = 60      # 만료임박: 잔여 0~60일
NOTICE_SOON = 30    # 자동연장주의: 통지마감까지 잔여 <= 30일

CONTRACT_TYPES = {"임대", "용역", "구독", "유지보수", "라이선스"}

# 대상유형 (Set-F1 채점 라벨) — 출력 정렬 기준 순서
TARGETS = ["만료경과", "만료임박", "자동연장주의"]
REASON = {
    "만료경과": "종료일이 기준일을 이미 지남(만료됨)",
    "만료임박": "종료일까지 60일 이내(곧 만료)",
    "자동연장주의": "자동갱신 계약인데 해지 통지 마감이 임박/경과(방치 시 자동 연장)",
}

# ── 한글 폰트 등록 (계약서 PDF용) ───────────────────────────────────────────
# 환경별 폰트 위치가 달라도 동작하도록 후보 경로를 탐색(없으면 명확한 안내로 중단).
_FONT_CANDIDATES = {
    # 등록명: (serif 후보들, sans 후보들)
    "serif": ["NanumMyeongjo.ttf", "NotoSerifCJK-Regular.ttc", "NanumGothic.ttf"],
    "serif_b": ["NanumMyeongjoBold.ttf", "NotoSerifCJK-Bold.ttc", "NanumGothicBold.ttf"],
    "sans": ["NanumGothic.ttf", "NotoSansCJK-Regular.ttc", "NanumBarunGothic.ttf"],
    "sans_b": ["NanumGothicBold.ttf", "NotoSansCJK-Bold.ttc", "NanumBarunGothicBold.ttf"],
}
_FONT_DIRS = [
    Path("/usr/share/fonts/truetype/nanum"),
    Path("/usr/share/fonts/opentype/noto"),
    Path("/usr/share/fonts"),
    Path("/Library/Fonts"),
    Path.home() / ".fonts",
]


def _find_font(names):
    for nm in names:
        for d in _FONT_DIRS:
            p = d / nm
            if p.exists():
                return p
        # 하위 디렉터리까지 탐색(느슨한 폴백)
        for d in _FONT_DIRS:
            if d.exists():
                hit = next(iter(d.rglob(nm)), None)
                if hit:
                    return hit
    return None


def _register_fonts():
    reg = {}
    for key, cands in _FONT_CANDIDATES.items():
        p = _find_font(cands)
        if p is None:
            sys.exit(f"[중단] 한글 폰트를 찾지 못했습니다({key}): {cands}.\n"
                     f"  나눔 또는 Noto CJK 폰트를 설치하거나 _FONT_DIRS 에 경로를 추가하세요.")
        reg[key] = p
    pdfmetrics.registerFont(TTFont("Nanum", str(reg["serif"])))
    pdfmetrics.registerFont(TTFont("NanumB", str(reg["serif_b"])))
    pdfmetrics.registerFont(TTFont("NanumG", str(reg["sans"])))
    pdfmetrics.registerFont(TTFont("NanumGB", str(reg["sans_b"])))


_register_fonts()


# ----------------------------------------------------------------------------
# 1) 입력 파싱 헬퍼
# ----------------------------------------------------------------------------
def int_or_none(v):
    """빈칸/None → None, 그 외 정수로(콤마·공백 제거)."""
    if v is None:
        return None
    s = str(v).replace(",", "").strip()
    if s == "":
        return None
    return int(float(s))


def yn(v, default="N"):
    """Y/N 정규화. 빈칸 → default."""
    s = (v or "").strip().upper()
    if s == "":
        return default
    return s


def add_months(d, months):
    """날짜 d에 months(음수 가능)를 더한다. 말일 보정(clamp)."""
    m = d.month - 1 + months
    y = d.year + m // 12
    mm = m % 12 + 1
    day = min(d.day, calendar.monthrange(y, mm)[1])
    return datetime.date(y, mm, day)


# ----------------------------------------------------------------------------
# 2) contracts.csv 로드 (fail-loud 검증)
# ----------------------------------------------------------------------------
def load_contracts():
    if not CSV_PATH.exists():
        sys.exit(
            f"[중단] 입력 파일이 없습니다: {CSV_PATH}\n"
            f"  INPUT_GUIDE.md를 참고해 contracts.csv를 먼저 채우세요(헤더 고정)."
        )
    rows = []
    seen = set()
    with CSV_PATH.open(encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        for i, r in enumerate(reader):
            cid = (r.get("case_id") or "").strip()
            ct = (r.get("contract_id") or "").strip()
            if not cid and not ct:
                continue  # 빈 행 무시
            if not (cid and ct):
                sys.exit(f"[중단] 행 {i+2}: case_id·contract_id는 비울 수 없습니다.")
            if ct in seen:
                sys.exit(f"[중단] {cid}: contract_id 중복 — {ct}. 계약 키는 유일해야 합니다.")
            seen.add(ct)

            name = (r.get("계약명") or "").strip()
            party = (r.get("거래상대방") or "").strip()
            ctype = (r.get("계약유형") or "").strip()
            if not (name and party and ctype):
                sys.exit(f"[중단] {cid}: 계약명·거래상대방·계약유형은 비울 수 없습니다.")
            if ctype not in CONTRACT_TYPES:
                sys.exit(f"[중단] {cid}: 계약유형 '{ctype}' 허용 외 "
                         f"({'|'.join(sorted(CONTRACT_TYPES))}).")

            term = int_or_none(r.get("계약기간_개월"))
            rem = int_or_none(r.get("잔여일수"))
            amount = int_or_none(r.get("계약금액"))
            if term is None or term <= 0:
                sys.exit(f"[중단] {cid}: 계약기간_개월은 양의 정수여야 합니다('{r.get('계약기간_개월')}').")
            if rem is None:
                sys.exit(f"[중단] {cid}: 잔여일수는 정수여야 합니다('{r.get('잔여일수')}').")

            auto = yn(r.get("자동갱신"))
            notice = int_or_none(r.get("갱신통지일")) or 0
            amt_var = yn(r.get("금액표기변형"))
            date_var = yn(r.get("날짜표기변형"))
            no_end = yn(r.get("종료일미명시"))
            no_amt = yn(r.get("금액미기재"))
            for fname, fval in (("자동갱신", auto), ("금액표기변형", amt_var),
                                ("날짜표기변형", date_var), ("종료일미명시", no_end),
                                ("금액미기재", no_amt)):
                if fval not in ("Y", "N"):
                    sys.exit(f"[중단] {cid}: {fname} '{fval}'는 허용되지 않음(Y|N).")

            # ── 입력 모순 fail-loud ─────────────────────────────────────────
            if no_amt == "Y" and amount is not None:
                sys.exit(f"[중단] {cid}: 금액미기재=Y인데 계약금액={amount} (조항이 없으면 금액 무의미). "
                         f"계약금액을 비우세요.")
            if no_amt == "N" and (amount is None or amount <= 0):
                sys.exit(f"[중단] {cid}: 금액미기재=N인데 계약금액이 비었거나 0 이하('{r.get('계약금액')}'). "
                         f"금액을 채우거나 금액미기재=Y로 두세요.")
            if auto == "N" and notice != 0:
                sys.exit(f"[중단] {cid}: 자동갱신=N인데 갱신통지일={notice} (자동갱신이 없으면 통지 무의미). "
                         f"갱신통지일을 0(빈칸)으로 두세요.")
            if auto == "Y" and notice <= 0:
                sys.exit(f"[중단] {cid}: 자동갱신=Y인데 갱신통지일={notice} (자동갱신은 통지기한이 필요). "
                         f"갱신통지일을 양의 정수로 두세요.")
            if no_amt == "Y" and amt_var == "Y":
                sys.exit(f"[중단] {cid}: 금액미기재=Y인데 금액표기변형=Y (금액 조항이 없으면 표기변형이 무의미). "
                         f"금액표기변형을 N으로 두세요.")

            # 종료일미명시=Y는 학생이 시작일+기간으로 종료일을 계산 →
            # 말일 보정(clamp) 모호를 피하려고 종료일의 day가 28 이하여야 한다(결정성 봉쇄 G5).
            end_date = ASSIGN_DATE + datetime.timedelta(days=rem)
            if no_end == "Y" and end_date.day > 28:
                sys.exit(f"[중단] {cid}: 종료일미명시=Y인데 종료일({end_date})의 일(day)이 28을 초과 — "
                         f"시작일+개월 계산 시 말일 보정 모호. 잔여일수를 조정해 종료일 일자를 1~28로 두세요.")

            rows.append(dict(
                case_id=cid, contract_id=ct, 계약명=name, 거래상대방=party, 계약유형=ctype,
                계약기간_개월=term, 잔여일수=rem, 계약금액=amount,
                자동갱신=auto, 갱신통지일=notice,
                금액표기변형=amt_var, 날짜표기변형=date_var,
                종료일미명시=no_end, 금액미기재=no_amt,
                _seq=len(rows),
            ))
    if not rows:
        sys.exit(f"[중단] {CSV_PATH} 에 계약 행이 없습니다.")
    return rows


# ----------------------------------------------------------------------------
# 3) 계약 객체 합성 — contracts.csv → 종료일/시작일 계산
# ----------------------------------------------------------------------------
def build_txns(contracts):
    txns = []
    for c in contracts:
        end_date = ASSIGN_DATE + datetime.timedelta(days=c["잔여일수"])
        start_date = add_months(end_date, -c["계약기간_개월"])
        # 종료일미명시 케이스의 역산 일관성(시작일+기간 == 종료일) 보증
        if c["종료일미명시"] == "Y" and add_months(start_date, c["계약기간_개월"]) != end_date:
            sys.exit(f"[중단] {c['case_id']}: 시작일+{c['계약기간_개월']}개월 ≠ 종료일 "
                     f"(말일 보정 충돌). 종료일미명시 케이스는 종료일 일자를 1~28로.")
        txns.append(dict(
            contract_id=c["contract_id"], 계약명=c["계약명"], 거래상대방=c["거래상대방"],
            계약유형=c["계약유형"], 계약기간_개월=c["계약기간_개월"],
            start_date=start_date, end_date=end_date,
            계약금액=c["계약금액"], 자동갱신=c["자동갱신"], 갱신통지일=c["갱신통지일"],
            금액표기변형=c["금액표기변형"], 날짜표기변형=c["날짜표기변형"],
            종료일미명시=c["종료일미명시"], 금액미기재=c["금액미기재"],
            _flags=c, _seq=c["_seq"],
        ))
    return txns


# ----------------------------------------------------------------------------
# 4) 정답 도출 — 합성한 계약에 갱신 규칙을 *다시 적용*해 대상 계산(플래그 직접 안 읽음)
# ----------------------------------------------------------------------------
def derive_targets(txn):
    """계약(txn)에서 대상유형 집합을 규칙으로 도출. 어디에도 안 들면 빈 집합."""
    t = set()
    rem = (txn["end_date"] - ASSIGN_DATE).days
    if rem < 0:
        t.add("만료경과")
    else:
        if rem <= SOON_DAYS:
            t.add("만료임박")
        if txn["자동갱신"] == "Y":
            notice_rem = rem - txn["갱신통지일"]
            if notice_rem <= NOTICE_SOON:
                t.add("자동연장주의")
    return t


def expected_from_flags(c):
    """contracts.csv 플래그가 의도한 대상 집합(자기검증/fail-loud 전용 — 정답 도출엔 미사용)."""
    t = set()
    rem = c["잔여일수"]
    if rem < 0:
        t.add("만료경과")
    else:
        if rem <= SOON_DAYS:
            t.add("만료임박")
        if c["자동갱신"] == "Y":
            notice_rem = rem - c["갱신통지일"]
            if notice_rem <= NOTICE_SOON:
                t.add("자동연장주의")
    return t


def cross_check(txns):
    """라운드트립 자기검증: 규칙도출(derive, end_date에서 잔여일수 재계산) == 플래그의도(expected,
       잔여일수 노브 직접). 종료일 산술(end_date = 기준일 + 잔여일수)이 어긋나면 여기서 잡힌다.
       ※ 이 둘은 같은 입력에 같은 규칙을 적용하므로 산술 가드 수준의 검증이다 — 생성기·채점기
         공통 버그를 잡는 '진짜' 독립검증은 eval/verify_recompute.py(이 파일 미사용)가 수행한다."""
    for t in txns:
        got = derive_targets(t)
        exp = expected_from_flags(t["_flags"])
        if got != exp:
            sys.exit(f"[중단] 자기검증 실패 {t['_flags']['case_id']} ({t['contract_id']}): "
                     f"규칙도출={sorted(got)} ≠ 플래그의도={sorted(exp)}. 생성기/입력 모순.")


# ----------------------------------------------------------------------------
# 5) 표기 렌더러 (날짜·금액 — 계약서 PDF용 FP 함정)
# ----------------------------------------------------------------------------
def fmt_date(d, style):
    if style == "iso":
        return d.isoformat()                       # 2026-08-07
    if style == "kor":
        return f"{d.year}년 {d.month}월 {d.day}일"   # 2026년 8월 7일
    if style == "dot":
        return f"{d.year}. {d.month:02d}. {d.day:02d}."  # 2026. 08. 07.
    if style == "slash":
        return f"{d.year}/{d.month:02d}/{d.day:02d}"      # 2026/08/07
    return d.isoformat()


DIG = "영일이삼사오육칠팔구"
_SMALL = [(1000, "천"), (100, "백"), (10, "십"), (1, "")]


def _k4(n):
    """0~9999 한글(천/백/십). 일십 → 십 으로 축약."""
    s = ""
    for div, nm in _SMALL:
        q = n // div
        if q:
            s += ("" if (q == 1 and nm) else DIG[q]) + nm
        n %= div
    return s


def korean_amount(n):
    """정수 금액 → 한글(억/만 단위). 예: 36000000 → 삼천육백만."""
    if n == 0:
        return "영"
    out = ""
    for div, nm in ((10 ** 8, "억"), (10 ** 4, "만"), (1, "")):
        q = n // div
        if q:
            out += _k4(q) + nm
        n %= div
    return out


def fmt_amount(n, variant, seq):
    """금액 문자열. 변형=N → '1,280,000원'. 변형=Y → 통화기호/한글병기(숫자 항상 병기)."""
    if variant != "Y":
        return f"{n:,}원"
    forms = [
        f"₩{n:,}",
        f"金 {korean_amount(n)}원정 (₩{n:,})",
        f"일금 {korean_amount(n)}원정 ({n:,}원)",
    ]
    return forms[seq % len(forms)]


def date_style_for(txn, layout):
    """양식 기본 날짜형식 + 날짜표기변형이면 더 까다로운 형식으로 교체."""
    base = {1: "kor", 2: "iso", 3: "kor", 4: "dot"}[layout]
    if txn["날짜표기변형"] == "Y":
        return ["dot", "slash", "kor"][txn["_seq"] % 3]
    return base


# ----------------------------------------------------------------------------
# 6) 계약서 PDF 렌더 (양식 4종 순환)
# ----------------------------------------------------------------------------
def clause_period(txn, dstyle):
    """계약기간 문구. 종료일미명시=Y면 종료일 대신 '시작일 + 기간(개월)'만 노출."""
    sd = fmt_date(txn["start_date"], dstyle)
    if txn["종료일미명시"] == "Y":
        return f"{sd}부터 {txn['계약기간_개월']}개월간"
    ed = fmt_date(txn["end_date"], dstyle)
    return f"{sd}부터 {ed}까지"


def clause_amount(txn):
    if txn["금액미기재"] == "Y":
        return "개별 발주 단가에 따라 별도 산정한다(총액 미기재)."
    return fmt_amount(txn["계약금액"], txn["금액표기변형"], txn["_seq"])


def clause_renew(txn):
    if txn["자동갱신"] == "Y":
        return (f"본 계약은 계약 만료일 {txn['갱신통지일']}일 전까지 어느 일방의 "
                f"서면에 의한 해지 통지가 없는 경우, 동일한 조건으로 1년간 자동으로 갱신된다.")
    return "본 계약은 계약기간 만료로 종료되며, 갱신을 원하는 경우 당사자는 별도의 재계약을 체결한다."


def _styles():
    return {
        "title_m": ParagraphStyle("tm", fontName="NanumB", fontSize=18,
                                  alignment=1, spaceAfter=14, leading=24),
        "title_g": ParagraphStyle("tg", fontName="NanumGB", fontSize=17,
                                  alignment=1, spaceAfter=14, leading=23),
        "body_m": ParagraphStyle("bm", fontName="Nanum", fontSize=10.5,
                                 leading=18, spaceAfter=4),
        "body_g": ParagraphStyle("bg", fontName="NanumG", fontSize=10.5,
                                 leading=18, spaceAfter=4),
        "small": ParagraphStyle("sm", fontName="Nanum", fontSize=9,
                                leading=14, textColor=colors.HexColor("#525252")),
        "cell": ParagraphStyle("cell", fontName="NanumG", fontSize=9.5, leading=14),
        "cellb": ParagraphStyle("cellb", fontName="NanumGB", fontSize=9.5, leading=14),
    }


def _info_table(data, st, col_widths):
    t = Table(data, colWidths=col_widths)
    t.setStyle(TableStyle([
        ("FONT", (0, 0), (-1, -1), "NanumG", 9.5),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#8d8d8d")),
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#f4f4f4")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    return t


def render_layout1(txn, st, dstyle):
    """양식 1 — 표준 조문형(명조). 제1조~제5조 줄글 조문."""
    e = []
    e.append(Paragraph(f"{txn['계약유형']} 계약서", st["title_m"]))
    e.append(Paragraph(f"계약번호: {txn['contract_id']}", st["small"]))
    e.append(Spacer(1, 6))
    e.append(Paragraph(
        f"주식회사 캐치업(이하 “갑”)과 {txn['거래상대방']}(이하 “을”)은 "
        f"「{txn['계약명']}」에 관하여 다음과 같이 계약을 체결한다.", st["body_m"]))
    e.append(Spacer(1, 8))
    e.append(Paragraph(f"<b>제1조 (목적)</b> 본 계약은 {txn['계약명']}에 관한 권리·의무를 정함을 목적으로 한다.", st["body_m"]))
    e.append(Paragraph(f"<b>제2조 (계약기간)</b> 계약기간은 {clause_period(txn, dstyle)}로 한다.", st["body_m"]))
    e.append(Paragraph(f"<b>제3조 (계약금액)</b> 계약금액은 {clause_amount(txn)}으로 정한다.", st["body_m"]))
    e.append(Paragraph(f"<b>제4조 (계약의 갱신)</b> {clause_renew(txn)}", st["body_m"]))
    e.append(Paragraph("<b>제5조 (기타)</b> 본 계약에 정하지 아니한 사항은 관계 법령 및 상관례에 따른다.", st["body_m"]))
    e.append(Spacer(1, 18))
    e.append(Paragraph(f"계약 체결일: {fmt_date(txn['start_date'], dstyle)}", st["body_m"]))
    e.append(Spacer(1, 10))
    e.append(Paragraph("갑: 주식회사 캐치업 (인)", st["body_m"]))
    e.append(Paragraph(f"을: {txn['거래상대방']} (인)", st["body_m"]))
    return e


def render_layout2(txn, st, dstyle):
    """양식 2 — 표 요약 + 갱신 조문(고딕). 핵심 정보를 상단 표로."""
    e = []
    e.append(Paragraph(f"{txn['계약명']}", st["title_g"]))
    e.append(Paragraph("계약 체결 확인서", st["small"]))
    e.append(Spacer(1, 8))
    cp = lambda s: Paragraph(str(s), st["cell"])
    rows = [
        ["계약번호", cp(txn["contract_id"])],
        ["거래상대방", cp(txn["거래상대방"])],
        ["계약유형", cp(txn["계약유형"])],
        ["계약기간", cp(clause_period(txn, dstyle))],
        ["계약금액", cp(clause_amount(txn))],
    ]
    e.append(_info_table(rows, st, [35 * mm, 120 * mm]))
    e.append(Spacer(1, 12))
    e.append(Paragraph("◆ 계약의 갱신", st["body_g"]))
    e.append(Paragraph(clause_renew(txn), st["body_g"]))
    e.append(Spacer(1, 16))
    e.append(Paragraph(f"확인일: {fmt_date(txn['start_date'], dstyle)}  /  주식회사 캐치업", st["body_g"]))
    return e


def render_layout3(txn, st, dstyle):
    """양식 3 — 약정서/레터형(명조). 정보를 줄글 문단 안에 녹임."""
    e = []
    e.append(Paragraph("약 정 서", st["title_m"]))
    e.append(Spacer(1, 6))
    e.append(Paragraph(
        f"주식회사 캐치업과 {txn['거래상대방']}은(는) 「{txn['계약명']}」(계약번호 {txn['contract_id']})에 "
        f"관하여 아래와 같이 약정합니다.", st["body_m"]))
    e.append(Spacer(1, 8))
    e.append(Paragraph(
        f"본 약정의 계약기간은 {clause_period(txn, dstyle)}이며, 그 대가로 지급할 계약금액은 "
        f"{clause_amount(txn)}으로 한다. {clause_renew(txn)} "
        f"양 당사자는 위 내용을 성실히 이행할 것을 확약한다.", st["body_m"]))
    e.append(Spacer(1, 18))
    e.append(Paragraph(f"작성일: {fmt_date(txn['start_date'], dstyle)}", st["body_m"]))
    e.append(Paragraph(f"수신: {txn['거래상대방']} 귀중", st["body_m"]))
    return e


def render_layout4(txn, st, dstyle):
    """양식 4 — 주문서/구독 확인서형(고딕). 전부 표로."""
    e = []
    e.append(Paragraph("SERVICE ORDER / 계약 주문서", st["title_g"]))
    e.append(Spacer(1, 8))
    cp = lambda s: Paragraph(str(s), st["cell"])
    rows = [
        ["항목", "내용"],
        ["계약번호 (No.)", cp(txn["contract_id"])],
        ["계약명 (Item)", cp(txn["계약명"])],
        ["공급자 (Vendor)", cp(txn["거래상대방"])],
        ["분류 (Type)", cp(txn["계약유형"])],
        ["계약기간 (Term)", cp(clause_period(txn, dstyle))],
        ["금액 (Amount)", cp(clause_amount(txn))],
        ["자동갱신 (Auto-Renewal)", cp(clause_renew(txn))],
    ]
    t = Table(rows, colWidths=[42 * mm, 113 * mm])
    t.setStyle(TableStyle([
        ("FONT", (0, 0), (-1, -1), "NanumG", 9.5),
        ("FONT", (0, 0), (-1, 0), "NanumGB", 10),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#8d8d8d")),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#161616")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("BACKGROUND", (0, 1), (0, -1), colors.HexColor("#f4f4f4")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    e.append(t)
    e.append(Spacer(1, 12))
    e.append(Paragraph(f"발행일: {fmt_date(txn['start_date'], dstyle)}  ·  주식회사 캐치업", st["small"]))
    return e


LAYOUTS = {1: render_layout1, 2: render_layout2, 3: render_layout3, 4: render_layout4}


def render_contract_pdf(txn):
    """계약 1건을 PDF로. 양식은 (seq 기준) 4종 순환 — 계약서마다 레이아웃/표기 다양."""
    layout = (txn["_seq"] % 4) + 1
    dstyle = date_style_for(txn, layout)
    st = _styles()
    path = CONTRACTS_DIR / f"{txn['contract_id']}.pdf"
    doc = SimpleDocTemplate(
        str(path), pagesize=A4,
        leftMargin=24 * mm, rightMargin=24 * mm, topMargin=22 * mm, bottomMargin=22 * mm,
        title=f"{txn['계약명']} ({txn['contract_id']})",
    )
    doc.build(LAYOUTS[layout](txn, st, dstyle))


# ----------------------------------------------------------------------------
# 7) 엑셀 산출 (계약대장 / 대상표)
# ----------------------------------------------------------------------------
HDR = Font(bold=True, color="FFFFFF")
HDRFILL = PatternFill("solid", fgColor="161616")
GUIDE = PatternFill("solid", fgColor="F4F4F4")

LEDGER_COLS = ["계약번호", "계약명", "거래상대방", "계약유형", "시작일", "종료일",
               "계약금액", "자동갱신", "갱신통지일"]
TARGET_COLS = ["계약번호", "대상유형", "사유"]


def style_header(ws, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(1, c)
        cell.font = HDR
        cell.fill = HDRFILL
        cell.alignment = Alignment(vertical="center")


def ledger_row(txn, answer):
    """계약대장 한 줄. answer=False면 식별칸(계약번호)만, True면 추출 정답."""
    if not answer:
        return [txn["contract_id"], "", "", "", "", "", "", "", ""]
    return [
        txn["contract_id"], txn["계약명"], txn["거래상대방"], txn["계약유형"],
        txn["start_date"].isoformat(), txn["end_date"].isoformat(),
        "" if txn["금액미기재"] == "Y" else txn["계약금액"],
        txn["자동갱신"],
        txn["갱신통지일"] if txn["자동갱신"] == "Y" else "",
    ]


def build_ledger(path, txns, answer):
    wb = openpyxl.Workbook()
    g = wb.active
    g.title = "안내"
    guide = [
        ["항목", "설명"],
        ["오늘(기준일)", "2026-06-28 (계약 만료까지 남은 날을 셀 때 이 날짜를 기준으로 한다)"],
        ["행 단위", "계약 1건 = 데이터/1_계약서/ 의 계약서 PDF 1개. 계약번호가 미리 채워져 있음"],
        ["입력", "데이터/1_계약서/<계약번호>.pdf — 양식이 제각각인 계약서. 여기서 값을 읽어 옮긴다"],
        ["채점 칸", "계약명·거래상대방·계약유형·시작일·종료일·계약금액·자동갱신·갱신통지일 (8칸, 셀 단위)"],
        ["계약명/거래상대방", "계약서에 적힌 그대로 옮긴다(앞뒤 공백은 무시됨)"],
        ["계약유형", "임대 · 용역 · 구독 · 유지보수 · 라이선스 중 하나"],
        ["시작일/종료일", "YYYY-MM-DD 로 통일해 적는다. 계약서 날짜 표기(2026년 8월 7일, 2026.08.07 등)는 이 형식으로 바꿔 적는다"],
        ["★종료일 미표기", "계약서에 종료일이 없고 '시작일부터 N개월'만 있으면 시작일 + N개월로 종료일을 계산해 적는다"],
        ["계약금액", "숫자만 적는다(예: 1280000). '₩', '원', 콤마, 한글 병기는 무시된다"],
        ["★금액 미기재", "계약서에 계약금액(총액) 조항이 없으면 비운다. 0이 아니라 빈칸이다"],
        ["자동갱신", "자동 갱신 조항이 있으면 Y, 기간 만료로 종료(재계약)면 N"],
        ["갱신통지일", "자동갱신=Y인 계약만, '만료 N일 전까지 통지' 의 N(정수). 자동갱신=N이면 비운다"],
        ["행 순서", "채점에 영향 없음(계약번호로 대조)"],
    ]
    for r in guide:
        g.append(r)
    g.cell(1, 1).font = Font(bold=True)
    g.cell(1, 2).font = Font(bold=True)
    for row in g.iter_rows(min_row=1, max_row=g.max_row, max_col=2):
        row[0].fill = GUIDE
    g.column_dimensions["A"].width = 22
    g.column_dimensions["B"].width = 86

    s = wb.create_sheet("계약대장")
    s.append(LEDGER_COLS)
    style_header(s, len(LEDGER_COLS))
    for t in txns:
        s.append(ledger_row(t, answer))
    s.freeze_panes = "A2"
    wb.save(path)


def build_targets(path, target_rows, answer):
    wb = openpyxl.Workbook()
    g = wb.active
    g.title = "안내"
    guide = [
        ["항목", "설명"],
        ["오늘(기준일)", "2026-06-28 (이 날짜를 기준으로 만료까지 남은 날을 센다)"],
        ["행 단위", "대상 1건 = (계약번호, 대상유형). 어디에도 해당 안 되는 계약은 적지 않는다"],
        ["입력", "데이터/1_계약대장_제출용.xlsx (1단계에서 본인이 채운 계약대장). 거기 적은 종료일·자동갱신·갱신통지일을 보고 고른다"],
        ["채점 칸", "계약번호, 대상유형 (2칸 집합 채점). 사유는 참고용(채점 안 함)"],
        ["대상유형", "허용값: 만료경과 · 만료임박 · 자동연장주의"],
        ["── 판정 규칙 ──", "남은날 = 종료일 − 2026-06-28. 아래를 적용해 해당하는 계약만 올린다(한 계약이 여러 유형이면 행도 여러 개)"],
        ["만료경과", "종료일이 기준일보다 과거이면(남은날 < 0) 위반"],
        ["만료임박", "남은날이 0~60일이면(0일=오늘 만료 포함, 60일 포함) 대상"],
        ["자동연장주의", "자동갱신=Y이고 아직 만료 전이며, '통지마감(=종료일 − 갱신통지일)'까지 남은날이 30일 이하이면 대상"],
        ["★경계", "남은날 60일=만료임박(O), 61일=아님. 통지마감까지 30일=자동연장주의(O), 31일=아님"],
        ["★정상은 올리지 않음", "세 규칙 중 어디에도 안 들면 적지 않는다 — 올리면 오탐(FP)으로 감점"],
        ["★감(感)으로 넣지 않음", "규칙이 정한 기준으로만 판단한다. '급해 보인다'는 느낌으로 추가하지 않는다"],
        ["행 순서", "채점에 영향 없음. 중복 행은 1개로 간주"],
        ["예시", "아래 '대상목록' 시트의 예시 행 참고(실제 답 아님, 채점 시 무시)"],
    ]
    for r in guide:
        g.append(r)
    g.cell(1, 1).font = Font(bold=True)
    g.cell(1, 2).font = Font(bold=True)
    for row in g.iter_rows(min_row=1, max_row=g.max_row, max_col=2):
        row[0].fill = GUIDE
    g.column_dimensions["A"].width = 22
    g.column_dimensions["B"].width = 86

    s = wb.create_sheet("대상목록")
    s.append(TARGET_COLS)
    style_header(s, len(TARGET_COLS))
    if answer:
        for row in target_rows:
            s.append(row)
    else:
        s.append(["예시CT-2025-999", "만료임박", "(예시 — 실제 답 아님, 채점 시 무시)"])
    s.freeze_panes = "A2"
    wb.save(path)


# ----------------------------------------------------------------------------
# main
# ----------------------------------------------------------------------------
def main():
    contracts = load_contracts()
    print(f"[1/4] contracts.csv 로드: 계약 {len(contracts)}건")

    txns = build_txns(contracts)
    cross_check(txns)   # fail-loud: 규칙도출 == 플래그의도
    print(f"[2/4] 계약 합성 + 자기검증 통과")

    # 정답 도출: 각 계약에 갱신 규칙 재적용 → 대상표
    target_rows = []
    tgt_count = {t: 0 for t in TARGETS}
    normal = 0
    for t in txns:
        ts = derive_targets(t)
        if not ts:
            normal += 1
        for tg in sorted(ts, key=TARGETS.index):
            target_rows.append([t["contract_id"], tg, REASON[tg]])
            tgt_count[tg] += 1
    print(f"[3/4] 정답 도출: 대상 {len(target_rows)}건 / 정상 {normal}건  대상유형별={tgt_count}")

    # 계약서 PDF (학생 입력)
    for t in txns:
        render_contract_pdf(t)
    # Part1 계약대장: 빈 양식(학생) + 정답키. (통합형 — Part2는 학생이 채운 위 양식을 입력으로 쓴다)
    build_ledger(DATA / "1_계약대장_제출용.xlsx", txns, answer=False)
    build_ledger(ANS / "1_계약대장_정답.xlsx", txns, answer=True)
    # Part2 대상표: 빈 양식 + 정답
    build_targets(DATA / "2_만료대상_제출용.xlsx", target_rows, answer=False)
    build_targets(ANS / "2_만료대상_정답.xlsx", target_rows, answer=True)
    print(f"[4/4] 산출 완료: 계약서 PDF {len(txns)}개 / 계약대장 양식+정답 / 대상표 양식+정답")

    print("=" * 60)
    print("생성 요약")
    print("  계약 수:", len(txns))
    print("  대상유형별 건수:", tgt_count)
    print("  정상(대상 없음) 계약:", normal)
    print("  대상 행 합계:", len(target_rows))
    print("=" * 60)


if __name__ == "__main__":
    main()
