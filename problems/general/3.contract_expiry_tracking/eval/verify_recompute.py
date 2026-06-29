# -*- coding: utf-8 -*-
"""
독립 재계산 검증 (S10) — build_dataset.py 로직을 '재사용하지 않고'
contracts.csv + 명시 규칙만으로 정답키를 다시 계산해 *_answer.xlsx 와 셀/집합 단위로 대조한다.

목적: 생성기(build_dataset)와 채점기(grade)에 같은 버그가 있으면 자기채점 100점이어도 못 잡는다.
      이 스크립트는 제3의 독립 구현으로 정답키 자체의 정합성을 검증한다(불일치 = 발행 차단).

실행: python eval/verify_recompute.py
"""
import csv
import sys
import datetime
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
CSV_PATH = ROOT / "eval" / "answer_key" / "contracts.csv"
ANS = ROOT / "eval" / "answer_key"

ASSIGN = datetime.date(2026, 6, 28)
SOON, NOTICE = 60, 30


def months_before(end: datetime.date, n: int) -> datetime.date:
    """end 에서 n개월 전 날짜(독립 구현: 연/월 직접 계산 + 말일 보정)."""
    total = (end.year * 12 + (end.month - 1)) - n
    y, m = divmod(total, 12)
    m += 1
    # 말일 보정 — 해당 월의 마지막 날을 직접 판정(calendar 미사용)
    if m == 12:
        nxt = datetime.date(y + 1, 1, 1)
    else:
        nxt = datetime.date(y, m + 1, 1)
    last = (nxt - datetime.timedelta(days=1)).day
    return datetime.date(y, m, min(end.day, last))


def i(v):
    s = (v or "").replace(",", "").strip()
    return int(s) if s != "" else None


def recompute():
    ledger = {}     # contract_id -> dict of 정답 셀
    targets = set()  # (contract_id, 대상유형)
    with CSV_PATH.open(encoding="utf-8-sig", newline="") as f:
        for r in csv.DictReader(f):
            ct = (r["contract_id"] or "").strip()
            if not ct:
                continue
            rem = i(r["잔여일수"])
            term = i(r["계약기간_개월"])
            auto = (r["자동갱신"] or "N").strip().upper()
            notice = i(r["갱신통지일"]) or 0
            no_amt = (r["금액미기재"] or "N").strip().upper()
            amt = i(r["계약금액"])

            end = ASSIGN + datetime.timedelta(days=rem)
            start = months_before(end, term)
            ledger[ct] = {
                "계약명": (r["계약명"] or "").strip(),
                "거래상대방": (r["거래상대방"] or "").strip(),
                "계약유형": (r["계약유형"] or "").strip(),
                "시작일": start.isoformat(),
                "종료일": end.isoformat(),
                "계약금액": "" if no_amt == "Y" else str(amt),
                "자동갱신": auto,
                "갱신통지일": str(notice) if auto == "Y" else "",
            }
            # 만료·갱신 규칙(독립 재구현)
            if rem < 0:
                targets.add((ct, "만료경과"))
            else:
                if rem <= SOON:
                    targets.add((ct, "만료임박"))
                if auto == "Y" and (rem - notice) <= NOTICE:
                    targets.add((ct, "자동연장주의"))
    return ledger, targets


def load_ledger_answer():
    wb = openpyxl.load_workbook(ANS / "contract_ledger_answer.xlsx", data_only=True)
    ws = wb["계약대장"]
    rows = list(ws.iter_rows(values_only=True))
    head = [str(h).strip() for h in rows[0]]
    out = {}
    for r in rows[1:]:
        d = {head[k]: r[k] for k in range(len(head))}
        ct = str(d["계약번호"]).strip()
        norm = {}
        for f in ("계약명", "거래상대방", "계약유형", "시작일", "종료일", "자동갱신"):
            norm[f] = "" if d[f] is None else str(d[f]).strip()
        for f in ("계약금액", "갱신통지일"):
            v = d[f]
            norm[f] = "" if v is None or str(v).strip() == "" else str(int(float(str(v))))
        # 날짜 셀이 datetime 으로 읽힐 수 있어 ISO 앞부분만
        for f in ("시작일", "종료일"):
            norm[f] = norm[f].split(" ")[0].split("T")[0]
        out[ct] = norm
    return out


def load_target_answer():
    wb = openpyxl.load_workbook(ANS / "expiry_targets_answer.xlsx", data_only=True)
    ws = wb["대상목록"]
    s = set()
    for r in list(ws.iter_rows(values_only=True))[1:]:
        if r[0] is None:
            continue
        ct = str(r[0]).strip()
        tg = str(r[1]).strip()
        if ct.startswith("예시"):
            continue
        s.add((ct, tg))
    return s


def main():
    exp_ledger, exp_targets = recompute()
    got_ledger = load_ledger_answer()
    got_targets = load_target_answer()

    errs = []
    # 1) 계약대장 셀 대조
    if set(exp_ledger) != set(got_ledger):
        errs.append(f"계약번호 집합 불일치: 재계산-정답키={set(exp_ledger)-set(got_ledger)} / "
                    f"정답키-재계산={set(got_ledger)-set(exp_ledger)}")
    for ct in sorted(set(exp_ledger) & set(got_ledger)):
        for f in exp_ledger[ct]:
            e, g = exp_ledger[ct][f], got_ledger[ct].get(f, "")
            if e != g:
                errs.append(f"[Part1] {ct}.{f}: 재계산='{e}' ≠ 정답키='{g}'")
    # 2) 대상 집합 대조
    if exp_targets != got_targets:
        errs.append(f"[Part2] 대상 집합 불일치: 재계산-정답키={sorted(exp_targets-got_targets)} / "
                    f"정답키-재계산={sorted(got_targets-exp_targets)}")

    print("=" * 60)
    print(" 독립 재계산 검증 (build_dataset 미사용)")
    print("=" * 60)
    print(f"  계약 {len(exp_ledger)}건 / 재계산 대상 {len(exp_targets)}행 / 정답키 대상 {len(got_targets)}행")
    if errs:
        print(f"  ✗ 불일치 {len(errs)}건 — 발행 차단")
        for e in errs[:40]:
            print("   -", e)
        sys.exit(1)
    print("  ✓ 계약대장 셀·대상 집합 모두 일치 — 정답키 정합")
    print("=" * 60)


if __name__ == "__main__":
    main()
