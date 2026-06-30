# -*- coding: utf-8 -*-
"""
거래처 견적 통합(취합/통합 + 보조 집계) 과제 — 데이터셋/정답 생성기 (단일 입력 = quotes.csv)

[아키타입] 취합/통합 (Set(행)+Exact(정규화 셀) 채점). "행을 보존하며 합치고, 최저가를 고른다".
  거래처 4곳(가람상사·나래유통·다온오피스·라온문구)의 견적 파일은 양식(컬럼명·단가표기·
  단위·품목명 표기)이 제각각이다. 같은 품목이 여러 거래처 파일에 흩어져 있다. 이들을 품목별
  한 줄(거래처 단가를 옆으로 펼친 비교표)로 통합하고, 각 품목의 최저가·최저가거래처를 표시한다.
  한 거래처만 견적한 품목도 보존한다.

[입력]
  eval/answer_key/quotes.csv   ★ canonical 1행=1품목(단일 진실 소스). 헤더(정확히):
    item_id,품목명,품목명_별칭,규격,단위,가람상사_단가,나래유통_단가,다온오피스_단가,라온문구_단가,표기변형,비교가능
    · item_id  = 품목 고유키(품목코드). 거래처마다 품목명을 다르게 불러도 이 코드로 같은 품목을 묶는다.
    · <거래처>_단가 = 그 거래처가 그 품목에 매긴 단가(원). 비우면 그 거래처는 견적 안 함(빈칸 보존 대상).
      비지 않은 단가 칸들이 곧 그 품목을 견적한 거래처 집합(present)이다.
    · 표기변형 = none | 품목명공백 | 품목명별칭
        - 품목명공백 : 모든 거래처 파일에서 품목명 앞뒤에 공백 주입(→ trim 함정).
        - 품목명별칭 : 우선순위 1위(present 중) 거래처만 정식 품목명, 나머지는 '품목명_별칭' 사용
                       (→ 이름이 달라도 품목코드로 같은 품목임을 알아내는 함정). present≥2 필요.
    · 비교가능 = Y | N
        - Y : 모든 견적 거래처가 같은 단위 → 최저가/최저가거래처를 정상 산출.
        - N : 견적 거래처들이 서로 다른 단위로 적음(개/박스 등) → 가격 직접 비교 불가.
              최저가는 비우고 최저가거래처='비교불가'. present≥2 필요(단위가 갈리려면).

[양식이 다른 4소스 — 코드가 canonical을 각 거래처 스키마/형식으로 렌더]
  데이터/가람상사_견적서.xlsx   품목코드, 품명, 규격, 단위, 단가(천단위 콤마 문자열 "12,500")
  데이터/나래유통_견적서.csv    code, name, spec, unit, price(정수)
  데이터/다온오피스_견적서.xlsx 자재코드, 품목명, 사양, 단위, 공급가(원 표기 "12,500원")
  데이터/라온문구_견적서.xlsx   코드, 상품명, 규격, 판매단위, 판매가(정수)

[표준 스키마 (정답 = 견적 비교표, 단일 산출)]
  품목코드, 품목명, 규격, 가람상사_단가, 나래유통_단가, 다온오피스_단가, 라온문구_단가, 최저가, 최저가거래처

[통합 규칙 (rulebook — 안내시트에 명시)]
  · 같은 품목코드 = 한 줄(거래처별 단가를 옆으로). 품목명/규격은 거래처마다 표기가 달라도
    우선순위(가나다: 가람상사>나래유통>다온오피스>라온문구) 1위 거래처 표기를 쓴다.
  · 단가 정규화 : 콤마·'원' 제거 → 정수. 견적 안 한 거래처 칸은 빈칸(0 아님).
  · 최저가 = 그 품목을 견적한 거래처 단가 중 최솟값. 최저가거래처 = 그 거래처(동점이면 가나다 우선).
  · 비교불가(단위 불일치) = 최저가 비움, 최저가거래처='비교불가'.
  · 한 거래처만 견적한 품목도 보존한다.

[고정 시드] random.seed(20260628). 품목명 공백 주입 폭만 시드에 묶인다. datetime.now() 미사용.

[정답 도출 철학] quotes.csv를 직접 베끼지 않는다 — 생성된 4소스를 *다시 파싱*하고 통합 규칙을
  *다시 적용*해(정규화·우선순위·최저가) 비교표를 재구성한다(consolidate). canonical 플래그는
  데이터 생성과 자기검증(cross_check)에만 쓴다 → 결정성·독립검증 성립.
"""
import csv
import sys
import random
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

# ----------------------------------------------------------------------------
# 경로 / 시드 (결정성)
# ----------------------------------------------------------------------------
ROOT = Path(__file__).resolve().parents[2]          # .../5.vendor_quote_consolidation
DATA = ROOT / "데이터"
ANS  = ROOT / "eval" / "answer_key"
CSV_PATH = ANS / "quotes.csv"
DATA.mkdir(parents=True, exist_ok=True)
ANS.mkdir(parents=True, exist_ok=True)

random.seed(20260628)

# ── 통합 규칙(rulebook) 상수 ────────────────────────────────────────────────
# 거래처 우선순위 = 리스트 순서(가나다). 품목명 표기 선택·최저가 동점 tie-break 둘 다에 쓴다.
VENDORS  = ["가람상사", "나래유통", "다온오피스", "라온문구"]
PRIORITY = {v: i for i, v in enumerate(VENDORS)}
PRICE_COLS = [f"{v}_단가" for v in VENDORS]
VARIANTS = {"none", "품목명공백", "품목명별칭"}
COMPARABLE = {"Y", "N"}

KEY = "품목코드"
STD_COLS = ["품목코드", "품목명", "규격"] + PRICE_COLS + ["최저가", "최저가거래처"]

# 거래처별 견적 파일 양식(컬럼명 제각각) + 단가 표기 형식 + 파일 종류
VENDOR_FILES = {
    "가람상사":   dict(파일="가람상사_견적서.xlsx", 시트="견적",   종류="xlsx",
                     cols=["품목코드", "품명", "규격", "단위", "단가"],   fmt="comma"),
    "나래유통":   dict(파일="나래유통_견적서.csv",                      종류="csv",
                     cols=["code", "name", "spec", "unit", "price"],     fmt="int"),
    "다온오피스": dict(파일="다온오피스_견적서.xlsx", 시트="Quote",  종류="xlsx",
                     cols=["자재코드", "품목명", "사양", "단위", "공급가"], fmt="won"),
    "라온문구":   dict(파일="라온문구_견적서.xlsx", 시트="견적서", 종류="xlsx",
                     cols=["코드", "상품명", "규격", "판매단위", "판매가"], fmt="int"),
}


# ----------------------------------------------------------------------------
# 1) 입력/정규화 헬퍼 (학생이 적용할 정규화 = 정답 도출에도 동일 적용)
# ----------------------------------------------------------------------------
def int_or_none(v):
    """빈칸/None → None, 그 외 정수로(콤마·'원'·'₩'·공백 제거)."""
    if v is None:
        return None
    s = str(v).replace(",", "").replace("원", "").replace("₩", "").strip()
    if s == "":
        return None
    return int(float(s))


def trim_text(s):
    """앞뒤 공백 trim(내부 공백은 보존 — 읽기 좋은 정답)."""
    return ("" if s is None else str(s)).strip()


def parse_price(v):
    """견적 단가 셀(콤마/'원'/정수) → 정수. 빈칸 → None."""
    return int_or_none(v)


# ----------------------------------------------------------------------------
# 2) quotes.csv 로드 (fail-loud 검증)
# ----------------------------------------------------------------------------
def load_quotes():
    if not CSV_PATH.exists():
        sys.exit(f"[중단] 입력 파일이 없습니다: {CSV_PATH}\n  INPUT_GUIDE.md 참고해 quotes.csv를 먼저 채우세요(헤더 고정).")
    rows, seen = [], set()
    with CSV_PATH.open(encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        for i, r in enumerate(reader):
            code = (r.get("item_id") or "").strip()
            if not code:
                if any((r.get(k) or "").strip() for k in ("품목명", *PRICE_COLS)):
                    sys.exit(f"[중단] 행 {i+2}: item_id가 비었습니다.")
                continue            # 빈 행 무시
            if code in seen:
                sys.exit(f"[중단] {code}: item_id 중복 — canonical 품목은 유일해야 합니다.")
            seen.add(code)

            name = (r.get("품목명") or "").strip()
            spec = (r.get("규격") or "").strip()
            unit = (r.get("단위") or "").strip()
            alias = (r.get("품목명_별칭") or "").strip()
            if not name or not spec or not unit:
                sys.exit(f"[중단] {code}: 품목명·규격·단위는 비울 수 없습니다.")

            prices = {}
            for v in VENDORS:
                p = int_or_none(r.get(f"{v}_단가"))
                if p is not None and p <= 0:
                    sys.exit(f"[중단] {code}: {v} 단가는 양의 정수여야 합니다('{r.get(f'{v}_단가')}').")
                if p is not None:
                    prices[v] = p
            present = [v for v in VENDORS if v in prices]
            if not present:
                sys.exit(f"[중단] {code}: 최소 한 거래처의 단가가 있어야 합니다(present=0).")

            variant = (r.get("표기변형") or "none").strip() or "none"
            comp = (r.get("비교가능") or "Y").strip() or "Y"
            if variant not in VARIANTS:
                sys.exit(f"[중단] {code}: 표기변형 '{variant}' 허용 안 됨({'|'.join(sorted(VARIANTS))}).")
            if comp not in COMPARABLE:
                sys.exit(f"[중단] {code}: 비교가능 '{comp}' 허용 안 됨(Y|N).")
            # ── 입력 모순 fail-loud ────────────────────────────────────────
            if variant == "품목명별칭":
                if not alias:
                    sys.exit(f"[중단] {code}: 표기변형=품목명별칭인데 품목명_별칭이 비었습니다.")
                if alias == name:
                    sys.exit(f"[중단] {code}: 품목명_별칭이 품목명과 같습니다 — 실질 변형이 없어 별칭 함정이 무력화됩니다.")
                if len(present) < 2:
                    sys.exit(f"[중단] {code}: 품목명별칭은 거래처 2곳 이상이어야 의미가 있습니다(present={len(present)}).")
            elif alias:
                sys.exit(f"[중단] {code}: 품목명_별칭이 채워졌는데 표기변형={variant}입니다 — 별칭이 조용히 무시됩니다(표기변형=품목명별칭으로 두거나 별칭을 비우세요).")
            if comp == "N" and len(present) < 2:
                sys.exit(f"[중단] {code}: 비교가능=N(단위 불일치)은 거래처 2곳 이상이어야 합니다(present={len(present)}).")

            rows.append(dict(
                item_id=code, 품목명=name, 품목명_별칭=alias, 규격=spec, 단위=unit,
                prices=prices, present=present, 표기변형=variant, 비교가능=comp,
                _seq=len(rows),
            ))
    if not rows:
        sys.exit(f"[중단] {CSV_PATH} 에 품목 행이 없습니다.")
    return rows


# ----------------------------------------------------------------------------
# 3) 4소스 렌더 — canonical → 각 거래처 스키마·형식으로. 표기변형·단위불일치 주입.
# ----------------------------------------------------------------------------
def pad_spaces(text):
    """앞뒤 공백 주입(표기변형 함정). trim 정규화로 사라짐."""
    return " " * random.randint(1, 2) + text + " " * random.randint(1, 2)


def alt_unit(u):
    """비교가능=N에서 비우선 거래처가 쓰는 '다른 단위'(canonical과 항상 다름)."""
    return "개" if u != "개" else "박스"


def render_price(vendor, n):
    fmt = VENDOR_FILES[vendor]["fmt"]
    if fmt == "comma":
        return f"{n:,}"          # 가람: "12,500"
    if fmt == "won":
        return f"{n:,}원"        # 다온: "12,500원"
    return n                     # 나래/라온: 정수


def build_sources(items):
    """canonical → 거래처별 견적 행 리스트. 같은 품목명을 거래처마다 다르게 표기하고(공백/별칭),
       비교가능=N이면 비우선 거래처에 다른 단위를 주입한다. 출력은 학생이 보는 셀 그대로."""
    src = {v: [] for v in VENDORS}
    n_pad = n_alias = n_incomp = 0
    for it in items:
        present = it["present"]                 # 우선순위 정렬(VENDORS 순서 = 가나다)
        top = present[0]
        if it["표기변형"] == "품목명공백":
            n_pad += 1
        if it["표기변형"] == "품목명별칭":
            n_alias += 1
        if it["비교가능"] == "N":
            n_incomp += 1
        for v in present:
            # 품목명 표기변형
            disp_name = it["품목명"]
            if it["표기변형"] == "품목명공백":
                disp_name = pad_spaces(it["품목명"])
            elif it["표기변형"] == "품목명별칭" and v != top:
                disp_name = it["품목명_별칭"]
            # 단위(비교가능=N이면 비우선 거래처에 다른 단위)
            disp_unit = it["단위"]
            if it["비교가능"] == "N" and v != top:
                disp_unit = alt_unit(it["단위"])
            # 단가(거래처 형식)
            cell_price = render_price(v, it["prices"][v])
            # 거래처 파일 컬럼 순서: 코드, 이름, 규격, 단위, 단가
            src[v].append([it["item_id"], disp_name, it["규격"], disp_unit, cell_price])
    return src, n_pad, n_alias, n_incomp


# ----------------------------------------------------------------------------
# 4) 정답 도출 — 4소스를 *다시 파싱* + 통합 규칙 *재적용* (quotes.csv 직접 안 읽음)
# ----------------------------------------------------------------------------
def parse_sources(src):
    """거래처별 견적 셀 → 표준 정규화 레코드(품목코드별 거래처 출현). 학생 정규화와 동일 규칙."""
    recs = []
    for v, rows in src.items():
        for r in rows:
            recs.append(dict(
                품목코드=str(r[0]).strip(), vendor=v,
                품목명=trim_text(r[1]), 규격=trim_text(r[2]),
                단위=trim_text(r[3]), 단가=parse_price(r[4]),
            ))
    return recs


def consolidate(recs):
    """품목코드로 그룹 → 1행. 품목명/규격 = 우선순위 1위 거래처 표기. 거래처별 단가 = 옆으로 펼침.
       최저가 = 견적 단가 최솟값(동점이면 가나다 우선). 단위가 갈리면 비교불가."""
    groups = {}
    for rec in recs:
        groups.setdefault(rec["품목코드"], []).append(rec)
    out = {}
    for code, g in groups.items():
        g_sorted = sorted(g, key=lambda x: PRIORITY[x["vendor"]])
        top = g_sorted[0]
        units = {x["단위"] for x in g}
        prices = {x["vendor"]: x["단가"] for x in g}
        row = dict(품목코드=code, 품목명=top["품목명"], 규격=top["규격"])
        for v in VENDORS:
            row[f"{v}_단가"] = prices.get(v)            # 견적 없으면 None(빈칸)
        if len(units) > 1:                              # 단위 불일치 → 비교불가
            row["최저가"] = None
            row["최저가거래처"] = "비교불가"
        else:
            minp = min(prices.values())
            winner = next(v for v in VENDORS if prices.get(v) == minp)   # 가나다 우선
            row["최저가"] = minp
            row["최저가거래처"] = winner
        out[code] = row
    return out


def expected_from_canonical(items):
    """canonical 플래그가 의도한 비교표 행(자기검증 전용)."""
    exp = {}
    for it in items:
        prices = it["prices"]
        row = dict(품목코드=it["item_id"], 품목명=it["품목명"], 규격=it["규격"])
        for v in VENDORS:
            row[f"{v}_단가"] = prices.get(v)
        if it["비교가능"] == "N":
            row["최저가"] = None
            row["최저가거래처"] = "비교불가"
        else:
            minp = min(prices[v] for v in it["present"])
            winner = next(v for v in VENDORS if prices.get(v) == minp)
            row["최저가"] = minp
            row["최저가거래처"] = winner
        exp[it["item_id"]] = row
    return exp


def _cmp_cell(v):
    """비교용 정규화(빈칸/None 통일)."""
    if v is None:
        return ""
    return str(v).replace(" ", "")


def cross_check(consolidated, expected):
    """소스 재구성(consolidate) == 플래그 의도(expected) 자기검증. 불일치=생성기 버그 → 중단."""
    if set(consolidated) != set(expected):
        miss = set(expected) - set(consolidated)
        extra = set(consolidated) - set(expected)
        sys.exit(f"[중단] 자기검증 실패: 품목코드 집합 불일치. 누락={sorted(miss)[:5]} 유령={sorted(extra)[:5]}")
    for code in expected:
        a, b = consolidated[code], expected[code]
        for col in STD_COLS:
            if _cmp_cell(a[col]) != _cmp_cell(b[col]):
                sys.exit(f"[중단] 자기검증 실패 {code}: '{col}' 재구성={a[col]!r} ≠ 의도={b[col]!r}. 생성기/입력 모순.")


# ----------------------------------------------------------------------------
# 5) 엑셀/CSV 산출
# ----------------------------------------------------------------------------
HDR = Font(bold=True, color="FFFFFF")
HDRFILL = PatternFill("solid", fgColor="161616")
GUIDE = PatternFill("solid", fgColor="F4F4F4")


def style_header(ws, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(1, c)
        cell.font = HDR
        cell.fill = HDRFILL
        cell.alignment = Alignment(vertical="center")


def write_vendor_xlsx(path, title, cols, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    ws.append(cols)
    style_header(ws, len(cols))
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    wb.save(path)


def write_vendor_csv(path, cols, rows):
    with Path(path).open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(cols)
        w.writerows(rows)


def std_row(rec):
    """비교표 한 줄(표준 스키마, 정답값). 빈칸(None)은 빈 문자열로(0 아님)."""
    out = [rec["품목코드"], rec["품목명"], rec["규격"]]
    for v in VENDORS:
        val = rec[f"{v}_단가"]
        out.append("" if val is None else val)
    out.append("" if rec["최저가"] is None else rec["최저가"])
    out.append(rec["최저가거래처"])
    return out


def build_comparison(path, ordered_recs, answer):
    """견적 비교표. answer=False → 빈 양식 + '안내'(+예시행). True → 정답 채움."""
    wb = openpyxl.Workbook()
    g = wb.active
    g.title = "안내"
    guide = [
        ["항목", "설명"],
        ["할 일", "거래처 4곳(가람상사·나래유통·다온오피스·라온문구)의 견적 파일을 품목별 비교표 "
                "한 장으로 합친다. 같은 품목은 한 줄에 모으고, 거래처별 단가를 옆으로 나란히 놓는다. "
                "품목마다 가장 싼 거래처를 표시한다. 한 거래처만 견적한 품목도 빠짐없이 넣는다."],
        ["입력 파일", "데이터/가람상사_견적서.xlsx · 데이터/나래유통_견적서.csv · "
                   "데이터/다온오피스_견적서.xlsx · 데이터/라온문구_견적서.xlsx (양식이 넷 다 다름)"],
        ["행 단위", "품목 1개 = 한 줄. 같은 품목코드는 결과 표에 딱 한 번만 나온다."],
        ["표준 컬럼(9칸)", "품목코드 · 품목명 · 규격 · 가람상사_단가 · 나래유통_단가 · "
                       "다온오피스_단가 · 라온문구_단가 · 최저가 · 최저가거래처"],
        ["컬럼명 맞추기", "가람상사: 품목코드/품명/규격/단위/단가 · 나래유통: code/name/spec/unit/price "
                      "· 다온오피스: 자재코드/품목명/사양/단위/공급가 · 라온문구: 코드/상품명/규격/판매단위/판매가 "
                      "→ 모두 위 표준 컬럼으로 맞춘다."],
        ["★같은 품목 묶기", "거래처마다 품목 이름을 조금씩 다르게 적어도(길게/줄여서), "
                        "품목코드가 같으면 같은 품목이다 — 한 줄로 모은다. 이름이 다르다고 갈라 놓으면 감점."],
        ["★품목명·규격 표기", "거래처마다 품목명 표기가 다르면, 거래처 우선순위가 가장 높은 곳의 표기를 쓴다. "
                         "우선순위(가나다순): 가람상사 > 나래유통 > 다온오피스 > 라온문구. 규격도 같은 규칙."],
        ["★거래처별 단가", "각 거래처가 그 품목에 매긴 단가를 해당 거래처 칸에 적는다. "
                       "그 거래처가 견적하지 않은 품목은 그 칸을 비운다(0 아님)."],
        ["단가 정규화", "단가는 정수로. 가람상사의 콤마 '12,500' → 12500, 다온오피스의 '12,500원' → 12500 "
                    "(콤마·'원' 제거)."],
        ["★최저가", "그 품목을 견적한 거래처들의 단가 중 가장 싼 값을 '최저가'에, 그 거래처 이름을 "
                  "'최저가거래처'에 적는다. 가장 싼 값이 두 거래처에서 같으면(동점) 우선순위가 빠른 거래처를 적는다."],
        ["★비교불가(단위 다름)", "한 품목을 견적한 거래처들이 서로 다른 단위(예: 한 곳은 '박스', 다른 곳은 '개')로 "
                           "적었으면 가격을 곧바로 비교할 수 없다. 이때는 '최저가'를 비우고 "
                           "'최저가거래처'에 '비교불가'라고 적는다."],
        ["★규격이 다르면 다른 품목", "이름이 비슷해도 규격(품목코드)이 다르면 다른 품목이다. 함부로 한 줄로 합치지 않는다."],
        ["★빈칸 보존", "견적 안 한 거래처 칸, 비교불가 품목의 최저가 칸은 비운다. 0과 빈칸은 다르다(빈칸을 0으로 채우면 감점)."],
        ["행 순서", "채점에 영향 없음(품목코드로 대조). 같은 품목코드는 1줄로 간주."],
        ["예시", "아래 '견적비교표' 시트의 예시 행 참고(실제 답 아님, 채점 시 무시)."],
    ]
    for r in guide:
        g.append(r)
    g.cell(1, 1).font = Font(bold=True)
    g.cell(1, 2).font = Font(bold=True)
    for row in g.iter_rows(min_row=1, max_row=g.max_row, max_col=2):
        row[0].fill = GUIDE
    g.column_dimensions["A"].width = 22
    g.column_dimensions["B"].width = 92

    s = wb.create_sheet("견적비교표")
    s.append(STD_COLS)
    style_header(s, len(STD_COLS))
    if answer:
        for rec in ordered_recs:
            s.append(std_row(rec))
    else:
        s.append(["예시OF-0001", "예시 품목", "예시 규격", 10000, 9800, "", 9900, 9800, "나래유통"])
    s.freeze_panes = "A2"
    wb.save(path)


# ----------------------------------------------------------------------------
# main
# ----------------------------------------------------------------------------
def main():
    items = load_quotes()
    print(f"[1/4] quotes.csv 로드: canonical 품목 {len(items)}건")

    src, n_pad, n_alias, n_incomp = build_sources(items)
    counts = " / ".join(f"{v} {len(src[v])}" for v in VENDORS)
    print(f"[2/4] 4소스 렌더: {counts} 행 (공백 {n_pad} · 별칭 {n_alias} · 비교불가 {n_incomp})")

    # 정답 도출: 소스 재파싱 → 통합 규칙 재적용
    recs = parse_sources(src)
    consolidated = consolidate(recs)
    expected = expected_from_canonical(items)
    cross_check(consolidated, expected)        # fail-loud 자기검증
    single = sum(1 for it in items if len(it["present"]) == 1)
    multi = len(items) - single
    print(f"[3/4] 비교표 도출: 품목 {len(consolidated)}건 (단일거래처 {single} / 복수거래처 {multi}) — 자기검증 통과")

    ordered = [consolidated[it["item_id"]] for it in items]   # 출력 순서 = canonical 순

    # 4소스 (학생 입력 원본)
    for v in VENDORS:
        meta = VENDOR_FILES[v]
        path = DATA / meta["파일"]
        if meta["종류"] == "csv":
            write_vendor_csv(path, meta["cols"], src[v])
        else:
            write_vendor_xlsx(path, meta["시트"], meta["cols"], src[v])
    # 제출 양식(빈) + 정답
    build_comparison(DATA / "견적비교표_제출용.xlsx", ordered, answer=False)
    build_comparison(ANS  / "견적비교표_정답.xlsx",   ordered, answer=True)
    print("[4/4] 산출 완료: 4소스(거래처 견적서) / 견적비교표_제출용 / 견적비교표_정답")

    n_incomp_rows = sum(1 for r in ordered if r["최저가거래처"] == "비교불가")
    n_blank_cells = sum(1 for r in ordered for v in VENDORS if r[f"{v}_단가"] is None)
    win = {v: sum(1 for r in ordered if r["최저가거래처"] == v) for v in VENDORS}
    print("=" * 64)
    print("생성 요약")
    print("  canonical 품목:", len(items))
    print("  소스 행 합계:", sum(len(src[v]) for v in VENDORS))
    print("  단일거래처 / 복수거래처:", single, "/", multi)
    print("  표기 공백 / 별칭 / 비교불가:", n_pad, "/", n_alias, "/", n_incomp_rows)
    print("  빈 단가 칸(견적 안 함):", n_blank_cells)
    print("  최저가거래처 분포:", win)
    print("=" * 64)


if __name__ == "__main__":
    main()
