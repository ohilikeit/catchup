# -*- coding: utf-8 -*-
"""
3종 대사(발주서·입고증·세금계산서) 검증/감사 과제 — 데이터셋/정답 생성기
(단일 입력 = cases.csv)

[아키타입] 검증/감사 (Set-F1 채점). "데이터 + 규칙 → 위반 판정". 정답이 곧 규칙.
  세 장부(발주/입고/계산서)를 같은 거래끼리 조인하고, 허용오차 규칙을 넘는 어긋남·누락만
  예외로 골라낸다. 정상(위반 없음)은 예외에 올리지 않는다(올리면 FP).

[입력]
  eval/answer_key/cases.csv   ★ 시나리오 명세(단일 진실 소스). 헤더(정확히):
    case_id,발주번호,품목코드,품목명,발주수량,발주단가,입고수량차,청구단가차,
    금액오류,입고누락,계산서누락,품목명표기변형
    · 거래 키 = (발주번호+품목코드). 한 발주번호에 여러 품목(라인아이템) 가능.
    · 입고수량차(int,0=일치) : 입고수량 = 발주수량 + 입고수량차.
    · 청구단가차(int 원,0=일치) : 청구단가 = 발주단가 + 청구단가차.
    · 금액오류(Y/N) : Y면 청구금액을 청구수량×청구단가에서 ±1원 초과 어긋나게 주입.
    · 입고누락/계산서누락(Y/N) : Y면 그 장부에서 행을 제외(누락).
    · 품목명표기변형(Y/N) : Y면 입고(없으면 계산서) 품목명만 다르게 표기. 조인은 품목코드로 →
      표기변형 자체는 위반이 아님(FP 함정).

[출력]
  data/발주_table.xlsx       발주서  (발주번호,품목코드,품목명,발주수량,발주단가)
  data/입고_table.xlsx       입고증  (발주번호,품목코드,품목명,입고수량) ── 입고누락=Y면 행 없음
  data/계산서_table.xlsx     세금계산서(발주번호,품목코드,품목명,청구수량,청구단가,청구금액)
                                                          ── 계산서누락=Y면 행 없음
  data/reconcile_table_template.xlsx   Part1 제출 양식(식별칸만, 빈 표 + '안내')
  data/exception_template.xlsx         Part2 제출 양식(빈 표 + '안내' + 예시행)
  data/reconcile_table_given.xlsx      ★분리형: Part2 입력용 '정답 대사표'(올바른 조인)
  eval/answer_key/reconcile_table_answer.xlsx  Part1 정답(대사정리표)
  eval/answer_key/exception_answer.xlsx        Part2 정답(예외목록 — 위반 전부)

[허용오차 규칙 (rulebook — 안내시트에 명시)]
  · 수량 : 발주수량 = 입고수량 = 청구수량 정확 일치(허용오차 0). 다르면 '수량불일치'.
  · 단가 : |청구단가 - 발주단가| <= 10 정상, > 10 이면 '단가불일치'. (10=정상, 11=위반)
  · 금액 : |청구금액 - 청구수량×청구단가| <= 1 정상, > 1 이면 '금액오류'. (1=정상, 2=위반)
  · 누락 : 입고 행 없음 → '입고누락'. 계산서 행 없음 → '계산서누락'.
  · 한 거래가 복수 위반 가능(예외 행 여러 개). 정상이면 예외에 올리지 않는다.

[고정 기준일/시드]
  ASSIGN_DATE = 2026-06-28 (장부일자 합성용, 채점엔 미사용). datetime.now() 금지.
  random.seed(20260628). cases.csv만 바꿔 재실행하면 3장부·양식·정답이 항상 정합 재생성.

[정답 도출 철학] 시나리오 플래그(금액오류=Y 등)를 직접 읽어 정답을 만들지 않는다 —
  3장부 테이블에 허용오차 규칙을 *다시 적용*해 위반을 도출한다(derive_violations).
  플래그는 데이터 생성과 자기검증(cross-check)에만 쓴다 → 결정성·독립검증 성립.
"""
import csv
import sys
import random
import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

# ----------------------------------------------------------------------------
# 경로 / 기준일 / 시드 (결정성 3종의 ①②)
# ----------------------------------------------------------------------------
ROOT = Path(__file__).resolve().parents[2]          # .../1.three_way_match
DATA = ROOT / "data"
ANS  = ROOT / "eval" / "answer_key"
CSV_PATH = ANS / "cases.csv"
# 주의: data/sources/ 의 운영자 배치 원본 스캔은 코드가 만들거나 지우지 않는다.
DATA.mkdir(parents=True, exist_ok=True)
ANS.mkdir(parents=True, exist_ok=True)

ASSIGN_DATE = datetime.date(2026, 6, 28)             # 고정 기준일(장부일자 합성용)
random.seed(20260628)

# ── 허용오차 규칙(rulebook) 상수 ────────────────────────────────────────────
QTY_TOL   = 0    # 수량: 정확 일치
PRICE_TOL = 10   # 단가: ±10원
AMOUNT_TOL = 1   # 금액: ±1원

# 위반유형 (Set-F1 채점 라벨) — 출력 정렬 기준 순서
VIOLATIONS = ["수량불일치", "단가불일치", "금액오류", "입고누락", "계산서누락"]
REASON = {
    "수량불일치": "발주·입고·청구 수량이 정확히 일치하지 않음",
    "단가불일치": "청구단가가 발주단가와 ±10원을 초과해 차이",
    "금액오류":   "청구금액이 청구수량×청구단가와 ±1원을 초과해 차이",
    "입고누락":   "입고증에 해당 거래 행이 없음",
    "계산서누락": "세금계산서에 해당 거래 행이 없음",
}


# ----------------------------------------------------------------------------
# 1) 입력 파싱 헬퍼
# ----------------------------------------------------------------------------
def int_or_none(v):
    """빈칸/None → None, 그 외 정수로(콤마·공백 제거)."""
    if v is None:
        return None
    s = str(v).replace(",", "").strip()
    if s == "":
        return None
    return int(float(s))


def yn(v, default="N"):
    """Y/N 정규화. 빈칸 → default."""
    s = (v or "").strip().upper()
    if s == "":
        return default
    return s


# ----------------------------------------------------------------------------
# 2) cases.csv 로드 (fail-loud 검증)
# ----------------------------------------------------------------------------
def load_cases():
    if not CSV_PATH.exists():
        sys.exit(
            f"[중단] 입력 파일이 없습니다: {CSV_PATH}\n"
            f"  INPUT_GUIDE.md를 참고해 cases.csv를 먼저 채우세요(헤더 고정)."
        )
    rows = []
    seen = set()
    with CSV_PATH.open(encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        for i, r in enumerate(reader):
            cid = (r.get("case_id") or "").strip()
            po  = (r.get("발주번호") or "").strip()
            code = (r.get("품목코드") or "").strip()
            if not cid and not po:
                continue   # 빈 행 무시
            if not (cid and po and code):
                sys.exit(f"[중단] 행 {i+2}: case_id·발주번호·품목코드는 비울 수 없습니다.")
            key = (po, code)
            if key in seen:
                sys.exit(f"[중단] {cid}: 거래 키(발주번호+품목코드) 중복 — {key}. "
                         f"한 거래는 유일해야 합니다.")
            seen.add(key)

            po_qty   = int_or_none(r.get("발주수량"))
            po_price = int_or_none(r.get("발주단가"))
            if po_qty is None or po_qty <= 0:
                sys.exit(f"[중단] {cid}: 발주수량은 양의 정수여야 합니다('{r.get('발주수량')}').")
            if po_price is None or po_price <= 0:
                sys.exit(f"[중단] {cid}: 발주단가는 양의 정수여야 합니다('{r.get('발주단가')}').")

            qd = int_or_none(r.get("입고수량차")) or 0
            pd = int_or_none(r.get("청구단가차")) or 0
            amt    = yn(r.get("금액오류"))
            grmiss = yn(r.get("입고누락"))
            invmiss = yn(r.get("계산서누락"))
            namevar = yn(r.get("품목명표기변형"))
            for fname, fval in (("금액오류", amt), ("입고누락", grmiss),
                                ("계산서누락", invmiss), ("품목명표기변형", namevar)):
                if fval not in ("Y", "N"):
                    sys.exit(f"[중단] {cid}: {fname} '{fval}'는 허용되지 않음(Y|N).")

            # ── 입력 모순 fail-loud ─────────────────────────────────────────
            if grmiss == "Y" and qd != 0:
                sys.exit(f"[중단] {cid}: 입고누락=Y인데 입고수량차={qd} (행이 없으면 수량차 무의미). "
                         f"입고수량차를 0으로 두세요.")
            if invmiss == "Y" and (pd != 0 or amt == "Y"):
                sys.exit(f"[중단] {cid}: 계산서누락=Y인데 청구단가차={pd}/금액오류={amt} "
                         f"(행이 없으면 단가·금액 무의미). 청구단가차=0·금액오류=N으로 두세요.")
            if namevar == "Y" and grmiss == "Y" and invmiss == "Y":
                sys.exit(f"[중단] {cid}: 품목명표기변형=Y인데 입고·계산서가 모두 누락 — "
                         f"표기를 변형할 장부가 없습니다.")

            rows.append(dict(
                case_id=cid, 발주번호=po, 품목코드=code,
                품목명=(r.get("품목명") or "").strip() or code,
                발주수량=po_qty, 발주단가=po_price,
                입고수량차=qd, 청구단가차=pd,
                금액오류=amt, 입고누락=grmiss, 계산서누락=invmiss, 품목명표기변형=namevar,
                _seq=len(rows),
            ))
    if not rows:
        sys.exit(f"[중단] {CSV_PATH} 에 거래 행이 없습니다.")
    return rows


# ----------------------------------------------------------------------------
# 3) 3장부 합성 — cases.csv → (발주/입고/계산서) 테이블 + 거래 객체
# ----------------------------------------------------------------------------
def vary_name(name):
    """품목명 표기 변형(공백·접두/접미 등). 조인은 품목코드로 하므로 위반이 아님(FP 함정)."""
    cands = [name.replace(" ", ""), name + "(주문)", "(구)" + name, name + " 외"]
    out = random.choice(cands)
    return out if out != name else name + " "


def amount_offset(c):
    """청구금액 주입 오차(원). 금액오류=Y → |오차|>1(경계 2 포함). N → 0 또는 +1(허용오차 내 반올림)."""
    if c["금액오류"] == "Y":
        return [2, 3, -50, 100, 250, -7][c["_seq"] % 6]   # 모두 |x|>1, 경계 2 포함
    return 1 if c["_seq"] % 6 == 0 else 0                  # 허용오차 내(<=1) → 정상(FP 함정)


def build_tables(cases):
    """3장부 행 리스트 + 거래(txn) 객체 리스트를 만든다."""
    po_rows, gr_rows, inv_rows = [], [], []
    txns = []
    for c in cases:
        po_qty, po_price = c["발주수량"], c["발주단가"]
        po_rows.append([c["발주번호"], c["품목코드"], c["품목명"], po_qty, po_price])

        gr_present = c["입고누락"] != "Y"
        gr_qty = None
        if gr_present:
            gr_qty = po_qty + c["입고수량차"]
            gr_name = vary_name(c["품목명"]) if c["품목명표기변형"] == "Y" else c["품목명"]
            gr_rows.append([c["발주번호"], c["품목코드"], gr_name, gr_qty])

        inv_present = c["계산서누락"] != "Y"
        inv_qty = inv_price = inv_amount = None
        if inv_present:
            inv_qty = po_qty                       # 청구수량 = 발주수량(공급사 청구 기준)
            inv_price = po_price + c["청구단가차"]
            inv_amount = inv_qty * inv_price + amount_offset(c)
            # 입고가 없을 때만 계산서 품목명을 변형(표기변형 함정을 둘 장부 확보)
            inv_name = c["품목명"]
            if c["품목명표기변형"] == "Y" and not gr_present:
                inv_name = vary_name(c["품목명"])
            inv_rows.append([c["발주번호"], c["품목코드"], inv_name,
                             inv_qty, inv_price, inv_amount])

        txns.append(dict(
            발주번호=c["발주번호"], 품목코드=c["품목코드"], case_id=c["case_id"],
            po_qty=po_qty, po_price=po_price,
            gr_present=gr_present, gr_qty=gr_qty,
            inv_present=inv_present, inv_qty=inv_qty,
            inv_price=inv_price, inv_amount=inv_amount,
            _flags=c,
        ))
    return po_rows, gr_rows, inv_rows, txns


# ----------------------------------------------------------------------------
# 4) 정답 도출 — 3장부에 허용오차 규칙을 *다시 적용*해 위반 계산(플래그 직접 읽지 않음)
# ----------------------------------------------------------------------------
def derive_violations(txn):
    """거래(txn)에서 위반유형 집합을 규칙으로 도출. 정상이면 빈 집합."""
    v = set()
    if not txn["gr_present"]:
        v.add("입고누락")
    if not txn["inv_present"]:
        v.add("계산서누락")
    # 수량: 존재하는 수량(발주·입고·청구)이 모두 같지 않으면 불일치
    qtys = [txn["po_qty"]]
    if txn["gr_present"]:
        qtys.append(txn["gr_qty"])
    if txn["inv_present"]:
        qtys.append(txn["inv_qty"])
    if len(set(qtys)) > 1:
        v.add("수량불일치")
    # 단가: |청구단가 - 발주단가| > 10
    if txn["inv_present"] and abs(txn["inv_price"] - txn["po_price"]) > PRICE_TOL:
        v.add("단가불일치")
    # 금액: |청구금액 - 청구수량×청구단가| > 1
    if txn["inv_present"] and abs(txn["inv_amount"] - txn["inv_qty"] * txn["inv_price"]) > AMOUNT_TOL:
        v.add("금액오류")
    return v


def expected_from_flags(c):
    """cases.csv 플래그가 의도한 위반 집합(자기검증/fail-loud 전용 — 정답 도출엔 미사용)."""
    exp = set()
    if c["입고누락"] == "Y":
        exp.add("입고누락")
    if c["계산서누락"] == "Y":
        exp.add("계산서누락")
    if c["입고누락"] != "Y" and c["입고수량차"] != 0:
        exp.add("수량불일치")
    if c["계산서누락"] != "Y" and abs(c["청구단가차"]) > PRICE_TOL:
        exp.add("단가불일치")
    if c["계산서누락"] != "Y" and c["금액오류"] == "Y":
        exp.add("금액오류")
    return exp


def cross_check(txns):
    """규칙도출(derive) == 플래그의도(expected) 를 거래마다 대조. 불일치=생성기 버그 → 즉시 중단.
       이것이 'fail-loud: 금액오류=N인데 금액 어긋남 모순' 등을 잡는 자기검증이다."""
    for t in txns:
        got = derive_violations(t)
        exp = expected_from_flags(t["_flags"])
        if got != exp:
            sys.exit(f"[중단] 자기검증 실패 {t['case_id']} ({t['발주번호']}/{t['품목코드']}): "
                     f"규칙도출={sorted(got)} ≠ 플래그의도={sorted(exp)}. 생성기/입력 모순.")


# ----------------------------------------------------------------------------
# 5) 엑셀 산출
# ----------------------------------------------------------------------------
HDR = Font(bold=True, color="FFFFFF")
HDRFILL = PatternFill("solid", fgColor="161616")
GUIDE = PatternFill("solid", fgColor="F4F4F4")

PO_COLS  = ["발주번호", "품목코드", "품목명", "발주수량", "발주단가"]
GR_COLS  = ["발주번호", "품목코드", "품목명", "입고수량"]
INV_COLS = ["발주번호", "품목코드", "품목명", "청구수량", "청구단가", "청구금액"]
RECON_COLS = ["발주번호", "품목코드", "발주수량", "발주단가", "입고수량",
              "청구수량", "청구단가", "청구금액"]
EXC_COLS = ["발주번호", "품목코드", "위반유형", "사유"]


def style_header(ws, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(1, c)
        cell.font = HDR
        cell.fill = HDRFILL
        cell.alignment = Alignment(vertical="center")


def write_table(path, title, cols, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    ws.append(cols)
    style_header(ws, len(cols))
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    wb.save(path)


def recon_row(txn, answer):
    """Part1 대사정리표 한 줄. answer=False면 식별칸(발주번호/품목코드)만, True면 조인 정답.
       누락측 칸은 빈칸."""
    if not answer:
        return [txn["발주번호"], txn["품목코드"], "", "", "", "", "", ""]
    return [
        txn["발주번호"], txn["품목코드"],
        txn["po_qty"], txn["po_price"],
        txn["gr_qty"] if txn["gr_present"] else "",
        txn["inv_qty"] if txn["inv_present"] else "",
        txn["inv_price"] if txn["inv_present"] else "",
        txn["inv_amount"] if txn["inv_present"] else "",
    ]


def build_reconcile(path, txns, answer):
    wb = openpyxl.Workbook()
    g = wb.active
    g.title = "안내"
    guide = [
        ["항목", "설명"],
        ["행 단위", "거래 1건 = (발주번호 + 품목코드). 발주서의 모든 거래가 한 줄씩(식별칸 채워져 있음)"],
        ["입력", "data/발주_table.xlsx · data/입고_table.xlsx · data/계산서_table.xlsx 3장부"],
        ["조인 키", "발주번호 + 품목코드로 세 장부를 맞댄다. 품목명 표기가 장부마다 달라도 품목코드로 조인"],
        ["채점 칸", "발주수량, 발주단가, 입고수량, 청구수량, 청구단가, 청구금액 (6칸, 셀 단위 Exact)"],
        ["발주수량/발주단가", "발주_table 그대로"],
        ["입고수량", "입고_table의 입고수량. 입고증에 그 거래 행이 없으면(입고누락) 빈칸"],
        ["청구수량/청구단가/청구금액", "계산서_table 그대로. 세금계산서에 행이 없으면(계산서누락) 세 칸 모두 빈칸"],
        ["★누락 = 빈칸", "장부에 행이 없으면 그 칸을 비운다. 0과 빈칸은 다르다 — 입고수량 0(0개 입고)과 입고누락(빈칸)은 구분"],
        ["★날짜 안 봄", "장부 일자는 보지 않는다 — 수량·단가·금액만 맞댄다"],
        ["행 순서", "채점에 영향 없음(거래 키로 대조)"],
    ]
    for r in guide:
        g.append(r)
    g.cell(1, 1).font = Font(bold=True)
    g.cell(1, 2).font = Font(bold=True)
    for row in g.iter_rows(min_row=1, max_row=g.max_row, max_col=2):
        row[0].fill = GUIDE
    g.column_dimensions["A"].width = 22
    g.column_dimensions["B"].width = 78

    s = wb.create_sheet("대사정리")
    s.append(RECON_COLS)
    style_header(s, len(RECON_COLS))
    for t in txns:
        s.append(recon_row(t, answer))
    s.freeze_panes = "A2"
    wb.save(path)


def build_exception(path, exc_rows, answer):
    wb = openpyxl.Workbook()
    g = wb.active
    g.title = "안내"
    guide = [
        ["항목", "설명"],
        ["행 단위", "위반 1건 = (발주번호, 품목코드, 위반유형). 정상 거래는 적지 않는다"],
        ["입력", "data/reconcile_table_given.xlsx (★제공된 정답 대사표). 1단계 조인을 틀려도 2단계는 영향 없음"],
        ["채점 칸", "발주번호, 품목코드, 위반유형 (3칸 집합 채점 = Set-F1). 사유는 참고용(채점 안 함)"],
        ["위반유형", "허용값: 수량불일치 · 단가불일치 · 금액오류 · 입고누락 · 계산서누락"],
        ["── 판정 규칙(허용오차) ──", "아래를 적용해 어긋남·누락만 예외로 올린다. 한 거래에 위반이 여러 개면 행도 여러 개"],
        ["수량불일치", "발주수량 = 입고수량 = 청구수량 이 정확히 일치하지 않으면 위반(허용오차 0). "
                     "입고/계산서가 누락이면 그 수량은 비교에서 빠짐(대신 누락 위반)"],
        ["단가불일치", "|청구단가 - 발주단가| 가 10원을 초과하면 위반. 딱 10원은 정상, 11원부터 위반"],
        ["금액오류",   "|청구금액 - (청구수량 × 청구단가)| 가 1원을 초과하면 위반. 딱 1원은 정상, 2원부터 위반"],
        ["입고누락",   "입고_table(입고증)에 그 거래 행이 아예 없으면 위반"],
        ["계산서누락", "계산서_table(세금계산서)에 그 거래 행이 아예 없으면 위반"],
        ["★정상은 올리지 않음", "허용오차 안이면(예: 단가차 10, 금액차 1, 수량 일치) 위반 아님 — 올리면 오탐(FP)으로 감점"],
        ["★표기변형은 위반 아님", "장부마다 품목명 표기가 달라도(품목코드 같음) 위반이 아니다 — 조인은 코드로"],
        ["행 순서", "채점에 영향 없음. 중복 행은 1개로 간주"],
        ["예시", "아래 '예외목록' 시트의 예시 행 참고(실제 답 아님, 채점 시 무시)"],
    ]
    for r in guide:
        g.append(r)
    g.cell(1, 1).font = Font(bold=True)
    g.cell(1, 2).font = Font(bold=True)
    for row in g.iter_rows(min_row=1, max_row=g.max_row, max_col=2):
        row[0].fill = GUIDE
    g.column_dimensions["A"].width = 22
    g.column_dimensions["B"].width = 78

    s = wb.create_sheet("예외목록")
    s.append(EXC_COLS)
    style_header(s, len(EXC_COLS))
    if answer:
        for row in exc_rows:
            s.append(row)
    else:
        s.append(["예시PO250001", "A1002", "단가불일치", "(예시 — 실제 답 아님)"])
    s.freeze_panes = "A2"
    wb.save(path)


# ----------------------------------------------------------------------------
# main
# ----------------------------------------------------------------------------
def main():
    cases = load_cases()
    print(f"[1/4] cases.csv 로드: 거래 {len(cases)}건")

    po_rows, gr_rows, inv_rows, txns = build_tables(cases)
    cross_check(txns)   # fail-loud: 규칙도출 == 플래그의도
    print(f"[2/4] 3장부 합성: 발주 {len(po_rows)} / 입고 {len(gr_rows)} / 계산서 {len(inv_rows)} 행"
          f" (자기검증 통과)")

    # 정답 도출: 각 거래에 규칙 재적용 → 예외목록
    exc_rows = []
    vio_count = {v: 0 for v in VIOLATIONS}
    normal = 0
    for t in txns:
        vs = derive_violations(t)
        if not vs:
            normal += 1
        for v in sorted(vs, key=VIOLATIONS.index):
            exc_rows.append([t["발주번호"], t["품목코드"], v, REASON[v]])
            vio_count[v] += 1
    print(f"[3/4] 정답 도출: 예외 {len(exc_rows)}건 / 정상 {normal}건  위반유형별={vio_count}")

    # 3장부 (학생 입력 원본)
    write_table(DATA / "발주_table.xlsx",   "발주서",     PO_COLS,  po_rows)
    write_table(DATA / "입고_table.xlsx",   "입고증",     GR_COLS,  gr_rows)
    write_table(DATA / "계산서_table.xlsx", "세금계산서", INV_COLS, inv_rows)
    # Part1 제출양식(빈) + 정답
    build_reconcile(DATA / "reconcile_table_template.xlsx", txns, answer=False)
    build_reconcile(ANS  / "reconcile_table_answer.xlsx",   txns, answer=True)
    # 분리형: Part2 입력용 정답 대사표(올바른 조인) = Part1 정답과 동일 내용
    build_reconcile(DATA / "reconcile_table_given.xlsx",    txns, answer=True)
    # Part2 제출양식(빈) + 정답
    build_exception(DATA / "exception_template.xlsx", exc_rows, answer=False)
    build_exception(ANS  / "exception_answer.xlsx",   exc_rows, answer=True)
    print("[4/4] 엑셀 산출 완료: 3장부 / 대사양식+given / 예외양식 / 정답키 2")

    print("=" * 60)
    print("생성 요약")
    print("  거래 수:", len(txns))
    print("  위반유형별 건수:", vio_count)
    print("  정상(위반없음) 거래:", normal)
    print("  예외 행 합계:", len(exc_rows))
    print("=" * 60)


if __name__ == "__main__":
    main()
