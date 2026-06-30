# -*- coding: utf-8 -*-
"""
VOC 3축 분류 과제 — 채점 스크립트
(분류 아키타입: 다클래스 macro-F1 + 혼동행렬 / 단일 산출)

사용:
  python3 eval/grade.py \
    --submission 데이터/1_분류표_제출용.xlsx \
    --answer-dir eval/answer_key

자기채점(정답키 자기검증, 100점 기대):
  python3 eval/grade.py \
    --submission eval/answer_key/1_분류표_정답.xlsx \
    --answer-dir eval/answer_key

배점: 문의유형 macro-F1 40 / 감성 macro-F1 30 / 위험 macro-F1 25 / 무결성 5.
"""
import argparse, json, re
from pathlib import Path
import openpyxl

# ---- 라벨 집합 -------------------------------------------------------------
TYPE_LABELS  = ["환불취소", "배송지연", "제품하자", "사용문의", "칭찬감사", "기타"]
SENTI_LABELS = ["긍정", "부정", "중립"]
RISK_LABELS  = ["위험", "일반"]
INVALID = "(무효)"   # 라벨집합 외 값/빈칸 센티넬


# ---- 정규화 ----------------------------------------------------------------
def norm_str(v):
    if v is None:
        return ""
    return re.sub(r"\s+", "", str(v)).strip()


TYPE_ALIAS = {
    "환불취소": "환불취소", "환불": "환불취소", "취소": "환불취소", "반품": "환불취소",
    "환급": "환불취소", "환불/취소": "환불취소", "환불·취소": "환불취소",
    "배송지연": "배송지연", "배송": "배송지연", "배송문제": "배송지연", "택배": "배송지연", "배송문의": "배송지연",
    "제품하자": "제품하자", "하자": "제품하자", "불량": "제품하자", "고장": "제품하자",
    "파손": "제품하자", "제품불량": "제품하자",
    "사용문의": "사용문의", "사용법": "사용문의", "사용방법": "사용문의", "사용법문의": "사용문의",
    "칭찬감사": "칭찬감사", "칭찬": "칭찬감사", "감사": "칭찬감사", "칭찬/감사": "칭찬감사", "칭찬·감사": "칭찬감사",
    "기타": "기타", "그외": "기타", "기타문의": "기타",
}
SENTI_ALIAS = {
    "긍정": "긍정", "positive": "긍정", "pos": "긍정", "+": "긍정", "긍": "긍정",
    "부정": "부정", "negative": "부정", "neg": "부정", "-": "부정", "부": "부정",
    "중립": "중립", "neutral": "중립", "neu": "중립", "중": "중립",
}
RISK_ALIAS = {
    "위험": "위험", "y": "위험", "긴급": "위험", "높음": "위험", "주의": "위험", "high": "위험",
    "일반": "일반", "n": "일반", "정상": "일반", "낮음": "일반", "low": "일반",
}


def norm_type(v):
    r = norm_str(v)
    if r == "":
        return INVALID
    if r in TYPE_ALIAS:
        return TYPE_ALIAS[r]
    return r if r in TYPE_LABELS else INVALID


def norm_senti(v):
    r = norm_str(v)
    if r == "":
        return INVALID
    low = r.lower()
    if r in SENTI_ALIAS:
        return SENTI_ALIAS[r]
    if low in SENTI_ALIAS:
        return SENTI_ALIAS[low]
    return r if r in SENTI_LABELS else INVALID


def norm_risk(v):
    r = norm_str(v)
    if r == "":
        return INVALID
    low = r.lower()
    if r in RISK_ALIAS:
        return RISK_ALIAS[r]
    if low in RISK_ALIAS:
        return RISK_ALIAS[low]
    return r if r in RISK_LABELS else INVALID


def load_sheet(path, sheet):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet in wb.sheetnames else wb[wb.sheetnames[-1]]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    head = [norm_str(h) for h in rows[0]]
    out = []
    for r in rows[1:]:
        if all(c is None or str(c).strip() == "" for c in r):
            continue
        out.append({head[i]: r[i] for i in range(len(head)) if i < len(r)})
    return out


# ---- 다클래스 macro-F1 + 혼동행렬 ------------------------------------------
def grade_axis(gold_map, sub_map, item_ids, labels, norm_fn):
    cols = labels + [INVALID]
    confusion = {g: {p: 0 for p in cols} for g in labels}
    tp = {c: 0 for c in labels}
    fp = {c: 0 for c in labels}
    fn = {c: 0 for c in labels}
    support = {c: 0 for c in labels}
    pred_count = {c: 0 for c in labels}
    invalid_count = 0

    for iid in item_ids:
        g = norm_fn(gold_map.get(iid))            # 정답은 항상 유효
        p = norm_fn(sub_map.get(iid))             # 제출(빈칸/오타 → INVALID)
        support[g] += 1
        confusion[g][p] += 1
        if p == INVALID:
            invalid_count += 1
        else:
            pred_count[p] += 1
        if p == g:
            tp[g] += 1
        else:
            fn[g] += 1
            if p in labels:
                fp[p] += 1

    per_class = {}
    f1s = []
    for c in labels:
        prec = tp[c] / (tp[c] + fp[c]) if (tp[c] + fp[c]) else 0.0
        rec  = tp[c] / (tp[c] + fn[c]) if (tp[c] + fn[c]) else 0.0
        f1   = (2 * prec * rec / (prec + rec)) if (prec + rec) else 0.0
        per_class[c] = {"정밀도": round(prec, 4), "재현율": round(rec, 4),
                        "F1": round(f1, 4), "support": support[c], "pred": pred_count[c]}
        if support[c] > 0 or pred_count[c] > 0:
            f1s.append(f1)
    macro = sum(f1s) / len(f1s) if f1s else 0.0
    return {"macro_f1": round(macro, 4), "per_class": per_class,
            "confusion": confusion, "invalid_count": invalid_count}


def fmt_confusion(conf, labels):
    cols = labels + [INVALID]
    short = {l: l[:4] for l in cols}
    head = "정답\\제출".ljust(10) + "".join(short[c].rjust(7) for c in cols)
    lines = [head]
    for g in labels:
        row = g[:9].ljust(10) + "".join(str(conf[g][c]).rjust(7) for c in cols)
        lines.append(row)
    return "\n".join(lines)


# ---- 배점 -----------------------------------------------------------------
# 유형 40 / 감성 30 / 위험 25 / 무결성 5.
def score(rt, rs, rr, ghost_n):
    s_type = rt["macro_f1"] * 40
    s_senti = rs["macro_f1"] * 30
    s_risk = rr["macro_f1"] * 25
    viol = rt["invalid_count"] + rs["invalid_count"] + rr["invalid_count"] + ghost_n
    s_int = max(0, 5 - viol * 0.5)
    return {
        "유형_macroF1(40)": round(s_type, 1),
        "감성_macroF1(30)": round(s_senti, 1),
        "위험_macroF1(25)": round(s_risk, 1),
        "무결성(5)": round(s_int, 1),
        "총점(100)": round(s_type + s_senti + s_risk + s_int, 1),
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--submission", required=True, help="분류표 제출용 xlsx")
    ap.add_argument("--answer-dir", required=True)
    args = ap.parse_args()

    adir = Path(args.answer_dir)
    sub = load_sheet(args.submission, "분류")
    ans = load_sheet(adir / "1_분류표_정답.xlsx", "분류")

    def key(r):
        return norm_str(r.get("item_id"))

    A_type = {key(r): r.get("문의유형") for r in ans}
    A_senti = {key(r): r.get("감성") for r in ans}
    A_risk = {key(r): r.get("위험") for r in ans}
    S_type = {key(r): r.get("문의유형") for r in sub if key(r)}
    S_senti = {key(r): r.get("감성") for r in sub if key(r)}
    S_risk = {key(r): r.get("위험") for r in sub if key(r)}
    item_ids = [key(r) for r in ans]
    valid_ids = set(item_ids)
    ghost = sorted({key(r) for r in sub if key(r) and key(r) not in valid_ids})

    rt = grade_axis(A_type, S_type, item_ids, TYPE_LABELS, norm_type)
    rs = grade_axis(A_senti, S_senti, item_ids, SENTI_LABELS, norm_senti)
    rr = grade_axis(A_risk, S_risk, item_ids, RISK_LABELS, norm_risk)

    sc = score(rt, rs, rr, len(ghost))

    print("=" * 64); print(" 채점 결과 (VOC 3축 분류)")
    print("=" * 64)
    for name, res, labs in [("문의유형", rt, TYPE_LABELS), ("감성", rs, SENTI_LABELS), ("위험", rr, RISK_LABELS)]:
        print(f"[{name}] macro-F1 {res['macro_f1']*100:.1f}%  (라벨 {len(labs)}클래스)")
        for c in labs:
            pc = res["per_class"][c]
            print(f"   - {c:<8} F1={pc['F1']*100:5.1f}  (P={pc['정밀도']*100:5.1f} R={pc['재현율']*100:5.1f} "
                  f"support={pc['support']:>3} pred={pc['pred']:>3})")
        if res["invalid_count"]:
            print(f"   ⚠ 라벨집합 외/빈칸: {res['invalid_count']}건")
        print("   [혼동행렬]")
        print("   " + fmt_confusion(res["confusion"], labs).replace("\n", "\n   "))
        print("-" * 64)
    if ghost:
        print(f"⚠ 유령 item_id(정답에 없음): {len(ghost)}건")
    print("-" * 64)
    for k, v in sc.items():
        print(f"  {k:<22} {v}")
    print("=" * 64)

    out = {"문의유형": rt, "감성": rs, "위험": rr, "유령ID": ghost, "score": sc}
    Path("result.json").write_text(json.dumps(out, ensure_ascii=False, indent=2))
    print("→ result.json 저장")


if __name__ == "__main__":
    main()
