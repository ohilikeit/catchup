# -*- coding: utf-8 -*-
"""
VOC(고객의 소리) 3축 분류 + 채널별 교차집계 과제 — 데이터셋/정답 생성기
(단일 입력 = items.csv)

[아키타입] 분류(다클래스 macro-F1 + 혼동행렬) + 분리형 교차집계(채널×유형 Exact).
  "텍스트 → 범주 라벨"을 3축(문의유형·감성·위험)으로 동시에 떨어뜨리고, 그 정답 분류표를
  별도 입력으로 제공해 채널×유형 건수표(교차집계)를 따로 세게 한다.
  판별 기준이 "이게 뭐냐(정답 모호)"라 결정성과 충돌 → **모호함을 규칙으로 결정화**한다:
  자유 판단이 아니라 **우선순위 규칙북**(키워드 → 라벨, 위에서부터 첫 매치)으로 모든 텍스트가
  유일 라벨로 떨어진다. 정답은 생성된 text에 규칙북을 **재적용**해 도출(시나리오 의도 라벨을
  직접 베끼지 않는다 — 규칙 재적용으로 도출해야 결정성/독립검증 성립).

[입력]
  eval/answer_key/items.csv   ★ 시나리오 명세(단일 진실 소스). 헤더(정확히):
    item_id,채널,템플릿,유형,감성,위험,충돌주입,충돌상대,경계케이스
    · item_id   : 고유키(중복 금지).
    · 채널       : 인입 채널 ∈ {앱리뷰,고객센터메일,전화상담,문의게시판,SNS}. (교차집계용)
    · 템플릿     : 문장 조립 힌트(auto). 채점·도출과 무관(가독성만).
    · 유형       : 최종 의도 문의유형 ∈ {환불취소,배송지연,제품하자,사용문의,칭찬감사,기타}.
    · 감성       : 최종 의도 감성 ∈ {긍정,부정,중립}. (충돌이면 부정>긍정이라 부정.)
    · 위험       : 최종 의도 위험 ∈ {위험,일반}. **감성과 독립** — 위험 키워드가 있어야 위험.
    · 충돌주입   : none | 유형충돌 | 감성충돌.
    · 충돌상대   : 충돌주입=유형충돌일 때 함께 주입할 하위 우선순위 유형. 그 외 '없음'.
    · 경계케이스 : none | 판별불가 | 약한신호.

[출력] (모두 한국어 네이밍)
  데이터/1_VOC_원문.xlsx           학생 입력(item_id, 채널, text). 라벨 없음.
  데이터/1_분류기준_위험규칙.pdf    규칙북 PDF (상수에서 렌더 — 안내시트와 동일 내용).
  데이터/1_분류표_제출용.xlsx       제출 양식(item_id·채널 채움 + 라벨 3칸 빈칸 + '안내').
  데이터/2_분류결과_참고용.xlsx     정답 분류표(교차집계 입력 — 학생이 보고 센다).
  데이터/2_채널별집계_제출용.xlsx   교차집계 양식(채널 채움 + 카운트 빈칸 + '안내').
  eval/answer_key/1_분류표_정답.xlsx       정답 분류표(채점용).
  eval/answer_key/2_채널별집계_정답.xlsx   정답 교차집계(채점용).

[규칙북 (rulebook — 안내시트/PDF/INPUT_GUIDE에 명시, 정답의 유일 근거)]
  · 문의유형(우선순위, 위에서부터 첫 매치): 1.환불취소 2.배송지연 3.제품하자 4.사용문의 5.칭찬감사 6.기타
  · 감성: 부정어 하나라도 있으면 부정, 없고 긍정어 있으면 긍정, 둘 다 없으면 중립. (부정 > 긍정)
  · 위험: 위험 키워드(소송·법적·소비자원 등)가 하나라도 있으면 위험, 없으면 일반. (감성과 독립)

[고정 기준일/시드]
  ASSIGN_DATE = 2026-06-28 (참고 상수, 채점 미사용). datetime.now() 금지.
  random.seed(20260628). items.csv만 바꿔 재실행하면 텍스트·양식·정답이 항상 정합 재생성.
"""
import csv
import sys
import random
import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle)
from reportlab.lib.styles import ParagraphStyle

# ----------------------------------------------------------------------------
# 경로 / 기준일 / 시드
# ----------------------------------------------------------------------------
ROOT = Path(__file__).resolve().parents[2]          # .../4.voc_classification
DATA = ROOT / "데이터"
ANS  = ROOT / "eval" / "answer_key"
CSV_PATH = ANS / "items.csv"
DATA.mkdir(parents=True, exist_ok=True)
ANS.mkdir(parents=True, exist_ok=True)

ASSIGN_DATE = datetime.date(2026, 6, 28)             # 고정 기준일(참고용)
random.seed(20260628)

# ── 규칙북(rulebook) ────────────────────────────────────────────────────────
# 문의유형: 우선순위 순(위에서부터 첫 매치). '기타'는 키워드 없음(fallback).
TYPE_ORDER = ["환불취소", "배송지연", "제품하자", "사용문의", "칭찬감사", "기타"]
TYPE_KEYWORDS = {
    "환불취소":  ["환불", "취소", "반품", "돈 돌려", "결제 취소", "환급"],
    "배송지연":  ["배송", "택배", "도착", "아직 안", "발송", "언제 와", "수령"],
    "제품하자":  ["불량", "고장", "깨져", "작동이 안", "하자", "파손", "흠집"],
    "사용문의":  ["어떻게 쓰", "사용법", "사용 방법", "설정", "연결하는 방법", "방법 문의", "방법 알려", "가입 방법"],
    "칭찬감사":  ["친절", "칭찬", "수고", "고맙", "고마웠", "덕분"],
    # 기타: 키워드 없음
}
# 감성: 부정 > 긍정 > 중립(없음).
SENTI_NEG = ["화가", "실망", "최악", "불편", "짜증", "별로", "엉망", "불쾌"]
SENTI_POS = ["좋아", "만족", "최고", "훌륭", "마음에", "좋네", "좋습니다", "감동", "흡족"]
SENTI_LABELS = ["긍정", "부정", "중립"]
# 위험: 키워드 있으면 위험(축 독립). 없으면 일반.
RISK_KEYWORDS = ["소송", "고소", "법적", "변호사", "소비자원", "공정위", "신고하겠",
                 "언론", "제보", "명예훼손", "불매", "유출", "해킹", "화상", "부상", "다쳤"]
RISK_LABELS = ["위험", "일반"]

# ── 채널(교차집계용) ────────────────────────────────────────────────────────
CHANNELS = ["앱리뷰", "고객센터메일", "전화상담", "문의게시판", "SNS"]

# ── 자연어 생성 어구(phrase) — 각 어구는 자기 축 키워드만 포함 ─────────────
TYPE_PHRASES = {
    "환불취소":  ["환불해 주세요", "주문을 취소하고 싶어요", "결제 취소 부탁드립니다",
                 "반품하고 환급받고 싶습니다", "돈 돌려주세요"],
    "배송지연":  ["배송이 언제 와요", "택배가 아직 안 왔어요", "발송이 됐는지 궁금해요",
                 "주문한 게 아직 도착을 안 했어요", "수령이 너무 늦네요"],
    "제품하자":  ["제품이 불량이에요", "받자마자 고장났어요", "화면이 깨져서 왔어요",
                 "작동이 안 됩니다", "하자가 있고 파손돼 있었어요"],
    "사용문의":  ["이거 어떻게 쓰나요", "사용법을 모르겠어요", "설정하는 방법 알려주세요",
                 "연결하는 방법이 궁금해요", "가입 방법 문의드려요"],
    "칭찬감사":  ["응대가 정말 친절하셨어요", "직원분을 칭찬하고 싶어요", "수고 많으셨습니다",
                 "고맙다는 말 전하고 싶어요", "덕분에 잘 해결됐어요"],
}
SENTI_PHRASES = {
    "긍정": ["정말 좋아요", "아주 만족합니다", "서비스 최고예요", "응대가 훌륭하네요", "마음에 쏙 들어요"],
    "부정": ["너무 화가 나요", "정말 실망했어요", "최악이네요", "너무 불편합니다", "짜증이 나요"],
}
RISK_PHRASES = ["소비자원에 신고하겠습니다", "법적으로 대응하겠습니다", "변호사를 통해 진행하겠습니다",
                "공정위에 제소하겠습니다", "언론에 제보하겠습니다", "불매 운동을 하겠습니다",
                "개인정보가 유출된 것 같습니다", "제품 때문에 화상을 입었습니다"]
# 판별불가/기타용 — 어떤 키워드도 포함하지 않는 중립 filler
NEUTRAL_FILLER = ["문의사항이 있어 연락드립니다", "확인 한번 부탁드려요", "안내 좀 받고 싶어요",
                  "담당자분과 통화하고 싶습니다", "궁금한 점이 있어서요"]
TEMPLATES = ["{body}.", "안녕하세요. {body}.", "{body} 확인 부탁드립니다.", "문의드립니다. {body}."]


def _assert_vocab_disjoint():
    """3축 pairwise 어휘 disjoint(유형↔감성↔위험). 한 어구가 두 축에 걸치면 derive 독립 도출 불가."""
    axes = {
        "유형": [k for ks in TYPE_KEYWORDS.values() for k in ks],
        "감성": SENTI_NEG + SENTI_POS,
        "위험": RISK_KEYWORDS,
    }
    names = list(axes)
    for i in range(len(names)):
        for j in range(i + 1, len(names)):
            for a in axes[names[i]]:
                for b in axes[names[j]]:
                    if a in b or b in a:
                        sys.exit(f"[중단] 어휘 충돌({names[i]}↔{names[j]}): '{a}' ↔ '{b}'. 키워드를 분리하세요.")
    # 합성 어구가 자기 축 외 키워드를 누수하지 않는지 추가 단언
    def axes_in(text, skip):
        hit = []
        for nm, kws in axes.items():
            if nm == skip:
                continue
            if any(k in text for k in kws):
                hit.append(nm)
        return hit
    for t, phrases in TYPE_PHRASES.items():
        for p in phrases:
            bad = axes_in(p, "유형")
            if bad:
                sys.exit(f"[중단] 유형 어구 누수 '{p}' → {bad} 키워드 포함.")
    for s_, phrases in SENTI_PHRASES.items():
        for p in phrases:
            bad = axes_in(p, "감성")
            if bad:
                sys.exit(f"[중단] 감성 어구 누수 '{p}' → {bad} 키워드 포함.")
    for p in RISK_PHRASES:
        bad = axes_in(p, "위험")
        if bad:
            sys.exit(f"[중단] 위험 어구 누수 '{p}' → {bad} 키워드 포함.")
    for p in NEUTRAL_FILLER:
        bad = axes_in(p, None)
        if bad:
            sys.exit(f"[중단] 중립필러 누수 '{p}' → {bad} 키워드 포함(판별불가 불가).")


# ----------------------------------------------------------------------------
# 1) 입력 파싱 헬퍼
# ----------------------------------------------------------------------------
def s(v):
    return (v or "").strip()


# ----------------------------------------------------------------------------
# 2) items.csv 로드 (fail-loud 검증)
# ----------------------------------------------------------------------------
def load_items():
    if not CSV_PATH.exists():
        sys.exit(
            f"[중단] 입력 파일이 없습니다: {CSV_PATH}\n"
            f"  INPUT_GUIDE.md를 참고해 items.csv를 먼저 채우세요(헤더 고정)."
        )
    rows = []
    seen = set()
    with CSV_PATH.open(encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        for i, r in enumerate(reader):
            iid = s(r.get("item_id"))
            if not iid:
                continue   # 빈 행 무시
            if iid in seen:
                sys.exit(f"[중단] 행 {i+2}: item_id 중복 — {iid}. 유일해야 합니다.")
            seen.add(iid)

            ch    = s(r.get("채널"))
            typ   = s(r.get("유형"))
            sen   = s(r.get("감성"))
            risk  = s(r.get("위험"))
            conf  = s(r.get("충돌주입")) or "none"
            rival = s(r.get("충돌상대")) or "없음"
            edge  = s(r.get("경계케이스")) or "none"

            if ch not in CHANNELS:
                sys.exit(f"[중단] {iid}: 채널 '{ch}' 허용 외 {CHANNELS}.")
            if typ not in TYPE_ORDER:
                sys.exit(f"[중단] {iid}: 유형 '{typ}' 허용 외 {TYPE_ORDER}.")
            if sen not in SENTI_LABELS:
                sys.exit(f"[중단] {iid}: 감성 '{sen}' 허용 외 {SENTI_LABELS}.")
            if risk not in RISK_LABELS:
                sys.exit(f"[중단] {iid}: 위험 '{risk}' 허용 외 {RISK_LABELS}.")
            if conf not in ("none", "유형충돌", "감성충돌"):
                sys.exit(f"[중단] {iid}: 충돌주입 '{conf}' 허용 외(none|유형충돌|감성충돌).")
            if edge not in ("none", "판별불가", "약한신호"):
                sys.exit(f"[중단] {iid}: 경계케이스 '{edge}' 허용 외(none|판별불가|약한신호).")

            # ── 입력 모순 fail-loud ─────────────────────────────────────────
            if conf == "유형충돌":
                if typ == "기타":
                    sys.exit(f"[중단] {iid}: 유형충돌인데 유형=기타 — 기타는 키워드가 없어 충돌 해소 불가.")
                if rival not in TYPE_ORDER or rival == "기타":
                    sys.exit(f"[중단] {iid}: 유형충돌의 충돌상대 '{rival}' 가 유효 유형이 아닙니다.")
                if TYPE_ORDER.index(rival) <= TYPE_ORDER.index(typ):
                    sys.exit(f"[중단] {iid}: 충돌상대 '{rival}'(우선순위 {TYPE_ORDER.index(rival)+1})가 "
                             f"유형 '{typ}'(우선순위 {TYPE_ORDER.index(typ)+1})보다 우선순위가 높거나 같음 — "
                             f"규칙이 유형을 못 고릅니다. 충돌상대는 유형보다 우선순위가 낮아야 합니다.")
            if conf == "감성충돌" and sen != "부정":
                sys.exit(f"[중단] {iid}: 감성충돌이면 규칙(부정>긍정)상 최종 감성은 '부정'이어야 합니다(감성='{sen}').")
            if edge == "판별불가":
                if typ != "기타" or sen != "중립" or risk != "일반":
                    sys.exit(f"[중단] {iid}: 판별불가는 유형=기타·감성=중립·위험=일반이어야 합니다"
                             f"(유형='{typ}', 감성='{sen}', 위험='{risk}').")
                if conf != "none":
                    sys.exit(f"[중단] {iid}: 판별불가인데 충돌주입='{conf}' — 키워드가 없어 충돌 불가.")
            if edge == "약한신호":
                if sen != "중립":
                    sys.exit(f"[중단] {iid}: 약한신호는 감성 신호가 없어 감성=중립이어야 합니다(감성='{sen}').")
                if typ == "기타":
                    sys.exit(f"[중단] {iid}: 약한신호는 유형 키워드가 있어야 합니다(유형=기타 불가).")
                if conf != "none":
                    sys.exit(f"[중단] {iid}: 약한신호인데 충돌주입='{conf}' — 단일 약한신호와 모순.")

            rows.append(dict(item_id=iid, 채널=ch, 템플릿=s(r.get("템플릿")) or "auto",
                             유형=typ, 감성=sen, 위험=risk, 충돌주입=conf, 충돌상대=rival,
                             경계케이스=edge, _seq=len(rows)))
    if not rows:
        sys.exit(f"[중단] {CSV_PATH} 에 항목 행이 없습니다.")
    return rows


# ----------------------------------------------------------------------------
# 3) 텍스트 합성 — items.csv → 자연어 문장(text)
# ----------------------------------------------------------------------------
def _pick(lst, seq):
    return lst[seq % len(lst)]


def build_text(c):
    """items.csv 한 행 → 자연어 VOC 문장. 의도 라벨이 규칙으로 재도출되도록 키워드를 심는다."""
    seq = c["_seq"]
    frags = []

    if c["경계케이스"] == "판별불가":
        # 어떤 키워드도 없는 중립 텍스트 (유형/감성/위험 키워드 0)
        frags.append(_pick(NEUTRAL_FILLER, seq))
    else:
        # ── 유형 어구 ───────────────────────────────────────────────
        if c["유형"] == "기타":
            frags.append(_pick(NEUTRAL_FILLER, seq + 2))
        else:
            frags.append(_pick(TYPE_PHRASES[c["유형"]], seq))
            if c["충돌주입"] == "유형충돌":
                frags.append(_pick(TYPE_PHRASES[c["충돌상대"]], seq + 1))

        # ── 감성 어구 ───────────────────────────────────────────────
        if c["감성"] != "중립":
            frags.append(_pick(SENTI_PHRASES[c["감성"]], seq))
        if c["충돌주입"] == "감성충돌":
            frags.append(_pick(SENTI_PHRASES["긍정"], seq + 1))

        # ── 위험 어구 (축 독립) ─────────────────────────────────────
        if c["위험"] == "위험":
            frags.append(_pick(RISK_PHRASES, seq))

    body = ". ".join(frags)
    return _pick(TEMPLATES, seq).format(body=body)


# ----------------------------------------------------------------------------
# 4) 정답 도출 — text에 규칙북을 *다시 적용*해 라벨 도출(의도 라벨 직접 안 읽음)
# ----------------------------------------------------------------------------
def derive_labels(text):
    """규칙북 재적용: (유형 우선순위 첫 매치, 감성 부정>긍정>중립, 위험 키워드 매치)."""
    t = text
    dtype = "기타"
    for label in TYPE_ORDER:                # 우선순위 순(위에서부터 첫 매치)
        if label == "기타":
            continue
        if any(kw in t for kw in TYPE_KEYWORDS[label]):
            dtype = label
            break
    if any(kw in t for kw in SENTI_NEG):
        dsenti = "부정"
    elif any(kw in t for kw in SENTI_POS):
        dsenti = "긍정"
    else:
        dsenti = "중립"
    drisk = "위험" if any(kw in t for kw in RISK_KEYWORDS) else "일반"
    return dtype, dsenti, drisk


def expected_from_flags(c):
    """items.csv 의도 라벨(자기검증/fail-loud 전용 — 정답 도출엔 미사용)."""
    return c["유형"], c["감성"], c["위험"]


def cross_check(items, texts):
    """규칙도출(derive) == 의도라벨(expected) 를 항목마다 대조. 불일치=생성기 버그 → 즉시 중단."""
    for c, text in zip(items, texts):
        got = derive_labels(text)
        exp = expected_from_flags(c)
        if got != exp:
            sys.exit(f"[중단] 자기검증 실패 {c['item_id']}: 규칙도출={got} ≠ 의도라벨={exp}\n"
                     f"       text='{text}'\n"
                     f"       (어휘 누수/우선순위/충돌 설정을 점검하세요.)")


# ----------------------------------------------------------------------------
# 5) 규칙북 콘텐츠 (안내시트 + PDF 공통 소스 — divergence 0)
# ----------------------------------------------------------------------------
def rulebook_sections():
    """규칙북을 비개발자 쉬운 말로. 안내시트(표)와 PDF(문서)가 같은 내용을 쓴다."""
    kw = TYPE_KEYWORDS
    return [
        ("이 표를 채우는 방법",
         ["고객 문장(text)을 한 줄씩 읽고, 아래 세 가지를 정해진 보기 중에서 골라 적습니다.",
          "느낌으로 찍지 말고, 문장에 어떤 말이 들어 있는지를 보고 아래 규칙대로 정합니다."]),
        ("① 무슨 문의인가 (문의유형) — 보기는 이 6가지뿐",
         ["환불취소 · 배송지연 · 제품하자 · 사용문의 · 칭찬감사 · 기타",
          "위에서부터 차례로 보고, 해당하는 말이 처음 걸리는 칸으로 정합니다.",
          "한 문장에 여러 개가 같이 있으면 더 위에 있는 것으로 정합니다.",
          f"1. 환불취소 — 다음 말이 있으면: {' · '.join(kw['환불취소'])}",
          f"2. 배송지연 — 다음 말이 있으면: {' · '.join(kw['배송지연'])}",
          f"3. 제품하자 — 다음 말이 있으면: {' · '.join(kw['제품하자'])}",
          f"4. 사용문의 — 다음 말이 있으면: {' · '.join(kw['사용문의'])}",
          f"5. 칭찬감사 — 다음 말이 있으면: {' · '.join(kw['칭찬감사'])}",
          "6. 기타 — 위 어디에도 걸리는 말이 없으면 기타."]),
        ("② 고객 기분은 어떤가 (감성) — 보기는 이 3가지뿐",
         ["긍정 · 부정 · 중립",
          f"나쁜 기분을 나타내는 말({' · '.join(SENTI_NEG)})이 하나라도 있으면 → 부정.",
          f"그런 말이 없고 좋은 기분을 나타내는 말({' · '.join(SENTI_POS)})이 있으면 → 긍정.",
          "둘 다 없으면 → 중립.",
          "좋은 말과 나쁜 말이 같이 있으면 → 부정(나쁜 말을 먼저 봅니다)."]),
        ("③ 위험한 문의인가 (위험) — 보기는 이 2가지뿐",
         ["위험 · 일반",
          f"다음과 같은 말이 하나라도 있으면 → 위험: {' · '.join(RISK_KEYWORDS)}",
          "그런 말이 하나도 없으면 → 일반.",
          "★ 위험은 기분과 따로 봅니다. 화가 났다고 무조건 위험이 아니고(부정인데 일반일 수 있어요),",
          "   차분한 말투라도 위 같은 말이 있으면 위험입니다(중립·긍정인데 위험일 수 있어요).",
          "   오직 위 목록의 말이 들어 있을 때만 위험으로 적습니다."]),
        ("★ 모르면 비웁니다 — 무리하게 짐작하지 마세요",
         ["어느 규칙에도 걸리는 말이 없으면, 억지로 끼워 맞추지 말고 문의유형=기타·감성=중립·위험=일반.",
          "문의유형과 기분과 위험은 서로 다른 칸입니다. 칭찬 문의라고 기분이 무조건 좋은 게 아니에요 —",
          "기분은 기분을 나타내는 말이 있을 때만, 위험은 위험을 나타내는 말이 있을 때만 적습니다.",
          "보기에 없는 이름을 새로 만들거나 빈칸으로 두면 틀린 답입니다.",
          "줄 순서는 점수와 상관없어요(번호 item_id로 맞춥니다)."]),
    ]


def crosstab_guide_sections():
    return [
        ("이 표를 채우는 방법",
         ["옆 파일 '2_분류결과_참고용.xlsx'의 '분류결과' 시트에는 모든 문의의 정답 분류가 이미 적혀 있습니다.",
          "그 표를 보고, 채널별로 각 문의유형이 몇 건인지와 위험 문의가 몇 건인지를 세어 적습니다.",
          "직접 분류하는 게 아니라, 이미 적힌 정답 분류표에서 건수만 세면 됩니다."]),
        ("어느 칸에 무엇을 적나",
         ["맨 왼쪽 줄에는 채널 5개가 미리 적혀 있습니다: " + " · ".join(CHANNELS),
          "그 오른쪽 칸에 문의유형 6개(" + " · ".join(TYPE_ORDER) + ") 각각의 건수를 적습니다.",
          "맨 오른쪽 '위험건수' 칸에는 그 채널의 위험 문의가 몇 건인지 적습니다.",
          "예: '앱리뷰' 줄의 '환불취소' 칸 = 참고용 표에서 채널이 앱리뷰이고 문의유형이 환불취소인 줄의 개수.",
          "숫자만 적습니다(0건이면 0). 빈칸으로 두지 마세요."]),
    ]


# ----------------------------------------------------------------------------
# 6) 엑셀 산출
# ----------------------------------------------------------------------------
HDR = Font(bold=True, color="FFFFFF")
HDRFILL = PatternFill("solid", fgColor="161616")
GUIDE = PatternFill("solid", fgColor="F4F4F4")

VOC_COLS    = ["item_id", "채널", "text"]
SUBMIT_COLS = ["item_id", "채널", "문의유형", "감성", "위험"]


def style_header(ws, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(1, c)
        cell.font = HDR
        cell.fill = HDRFILL
        cell.alignment = Alignment(vertical="center")


def _write_guide_sheet(ws, sections):
    ws.append(["항목", "설명"])
    ws.cell(1, 1).font = Font(bold=True)
    ws.cell(1, 2).font = Font(bold=True)
    for head, lines in sections:
        ws.append([head, lines[0] if lines else ""])
        for ln in lines[1:]:
            ws.append(["", ln])
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=2):
        row[0].fill = GUIDE
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 86


def write_voc(path, items, texts):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "VOC"
    ws.append(VOC_COLS)
    style_header(ws, len(VOC_COLS))
    for c, text in zip(items, texts):
        ws.append([c["item_id"], c["채널"], text])
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 80
    ws.freeze_panes = "A2"
    wb.save(path)


def build_classification(path, items, answers, answer, sheet_name, guide_sections):
    """분류표(빈 제출용/정답/참고용)를 찍는다. answer=False면 라벨 3칸 빈칸, True면 정답 채움."""
    wb = openpyxl.Workbook()
    g = wb.active
    g.title = "안내"
    _write_guide_sheet(g, guide_sections)

    sht = wb.create_sheet(sheet_name)
    sht.append(SUBMIT_COLS)
    style_header(sht, len(SUBMIT_COLS))
    for c in items:
        if answer:
            dt, ds, dr = answers[c["item_id"]]
            sht.append([c["item_id"], c["채널"], dt, ds, dr])
        else:
            sht.append([c["item_id"], c["채널"], "", "", ""])
    sht.column_dimensions["A"].width = 10
    sht.column_dimensions["B"].width = 14
    sht.freeze_panes = "A2"
    wb.save(path)


def _crosstab_counts(items, answers):
    """채널×유형 건수 + 채널별 위험건수."""
    cell = {ch: {t: 0 for t in TYPE_ORDER} for ch in CHANNELS}
    risk = {ch: 0 for ch in CHANNELS}
    for c in items:
        dt, ds, dr = answers[c["item_id"]]
        cell[c["채널"]][dt] += 1
        if dr == "위험":
            risk[c["채널"]] += 1
    return cell, risk


def build_crosstab(path, items, answers, answer, guide_sections):
    """교차집계(빈 제출용/정답). 행=채널, 열=유형6 + 위험건수."""
    cell, risk = _crosstab_counts(items, answers)
    wb = openpyxl.Workbook()
    g = wb.active
    g.title = "안내"
    _write_guide_sheet(g, guide_sections)

    sht = wb.create_sheet("교차집계")
    cols = ["채널"] + TYPE_ORDER + ["위험건수"]
    sht.append(cols)
    style_header(sht, len(cols))
    for ch in CHANNELS:
        if answer:
            sht.append([ch] + [cell[ch][t] for t in TYPE_ORDER] + [risk[ch]])
        else:
            sht.append([ch] + ["" for _ in TYPE_ORDER] + [""])
    sht.column_dimensions["A"].width = 14
    for i in range(len(cols) - 1):
        sht.column_dimensions[chr(ord("B") + i)].width = 11
    sht.freeze_panes = "B2"
    wb.save(path)


# ----------------------------------------------------------------------------
# 7) 규칙북 PDF (reportlab CID 폰트 — 한국어)
# ----------------------------------------------------------------------------
def _register_korean_font():
    """한글 글리프를 임베드한 TTF가 있으면 등록해 모든 뷰어에서 렌더되게 하고(강건),
       없으면 reportlab 내장 CID 폰트로 폴백한다(이식성). 등록한 폰트명을 반환."""
    candidates = [
        "/usr/share/fonts/truetype/nanum/NanumGothic.ttf",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/truetype/noto/NotoSansCJKkr-Regular.otf",
        "/mnt/c/Windows/Fonts/malgun.ttf",          # Windows 한글(맑은 고딕)
        "/Library/Fonts/AppleSDGothicNeo.ttc",       # macOS
    ]
    for p in candidates:
        if Path(p).exists():
            try:
                pdfmetrics.registerFont(TTFont("KoEmbed", p))
                return "KoEmbed"
            except Exception:
                continue
    pdfmetrics.registerFont(UnicodeCIDFont("HYSMyeongJo-Medium"))
    return "HYSMyeongJo-Medium"


def render_rulebook_pdf(path):
    """한 장짜리 규칙 요약 카드. 핵심(키워드·우선순위·세 규칙)만 표로 깔끔하게.
       키워드/규칙은 제출양식 '안내' 시트와 동일 상수에서 나온다(내용 불일치 0)."""
    FN = _register_korean_font()
    title = ParagraphStyle("t", fontName=FN, fontSize=15, leading=20, spaceAfter=2)
    sub = ParagraphStyle("s", fontName=FN, fontSize=9, leading=13, textColor=colors.HexColor("#525252"),
                         spaceAfter=8)
    h = ParagraphStyle("h", fontName=FN, fontSize=11, leading=15, spaceBefore=9, spaceAfter=3,
                       textColor=colors.HexColor("#161616"))
    note = ParagraphStyle("n", fontName=FN, fontSize=8.5, leading=12,
                          textColor=colors.HexColor("#525252"), spaceBefore=2)
    cell = ParagraphStyle("c", fontName=FN, fontSize=9, leading=12.5)

    def P(t, st=cell):
        return Paragraph(t, st)

    def mk(rows, widths):
        t = Table(rows, colWidths=widths, hAlign="LEFT")
        t.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, -1), FN),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#161616")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#C6C6C6")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F4F4")]),
            ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        return t

    doc = SimpleDocTemplate(str(path), pagesize=A4,
                            leftMargin=16 * mm, rightMargin=16 * mm,
                            topMargin=16 * mm, bottomMargin=16 * mm)
    W = A4[0] - 32 * mm
    flow = [P("고객의 소리(VOC) 분류 기준", title),
            P("고객 문장에 어떤 말이 들어 있는지를 보고 세 칸(문의유형·감성·위험)을 규칙대로 정합니다. "
              "느낌으로 정하지 않습니다.", sub)]

    # ① 문의유형
    flow.append(P("① 문의유형 — 위에서부터 첫 번째로 걸리는 칸 (보기 6개)", h))
    rows = [[P("유형"), P("이런 말이 있으면")]]
    for t in TYPE_ORDER[:-1]:
        rows.append([P(t), P(" · ".join(TYPE_KEYWORDS[t]))])
    rows.append([P("기타"), P("위 어디에도 걸리는 말이 없으면")])
    flow.append(mk(rows, [W * 0.20, W * 0.80]))
    flow.append(P("여러 개가 같이 있으면 더 위(번호 작은) 유형으로 정합니다.", note))

    # ② 감성
    flow.append(P("② 감성 — 나쁜 말이 우선 (보기 3개)", h))
    rows = [[P("감성"), P("기준")],
            [P("부정"), P(" · ".join(SENTI_NEG))],
            [P("긍정"), P("(부정어가 없을 때) " + " · ".join(SENTI_POS))],
            [P("중립"), P("위 말이 하나도 없으면")]]
    flow.append(mk(rows, [W * 0.20, W * 0.80]))
    flow.append(P("좋은 말과 나쁜 말이 같이 있으면 → 부정.", note))

    # ③ 위험
    flow.append(P("③ 위험 — 기분과 따로 봅니다 (보기 2개)", h))
    rows = [[P("위험"), P(" · ".join(RISK_KEYWORDS))],
            [P("일반"), P("위 말이 하나도 없으면")]]
    flow.append(mk(rows, [W * 0.20, W * 0.80]))
    flow.append(P("화가 나도 위 말이 없으면 일반, 차분해도 위 말이 있으면 위험. 위 목록의 말이 있을 때만 위험.", note))

    flow.append(Spacer(1, 6))
    flow.append(P("걸리는 말이 없으면 무리하게 짐작하지 말고 <b>문의유형=기타·감성=중립·위험=일반</b>. "
                  "보기에 없는 이름을 쓰거나 빈칸으로 두면 오답입니다.", sub))
    doc.build(flow)


# ----------------------------------------------------------------------------
# main
# ----------------------------------------------------------------------------
def main():
    _assert_vocab_disjoint()
    items = load_items()
    print(f"[1/5] items.csv 로드: 항목 {len(items)}건")

    texts = [build_text(c) for c in items]
    cross_check(items, texts)   # fail-loud: 규칙도출 == 의도라벨
    print(f"[2/5] 텍스트 합성 + 자기검증 통과: {len(texts)}문장")

    # 정답 도출: 각 텍스트에 규칙 재적용
    answers = {}
    type_count = {t: 0 for t in TYPE_ORDER}
    senti_count = {s_: 0 for s_ in SENTI_LABELS}
    risk_count = {r: 0 for r in RISK_LABELS}
    chan_count = {ch: 0 for ch in CHANNELS}
    for c, text in zip(items, texts):
        dt, ds, dr = derive_labels(text)
        answers[c["item_id"]] = (dt, ds, dr)
        type_count[dt] += 1
        senti_count[ds] += 1
        risk_count[dr] += 1
        chan_count[c["채널"]] += 1
    print(f"[3/5] 정답 도출(규칙 재적용): 유형={type_count}")
    print(f"      감성={senti_count} / 위험={risk_count}")

    # 산출
    write_voc(DATA / "1_VOC_원문.xlsx", items, texts)
    render_rulebook_pdf(DATA / "1_분류기준_위험규칙.pdf")
    build_classification(DATA / "1_분류표_제출용.xlsx", items, answers,
                         answer=False, sheet_name="분류", guide_sections=rulebook_sections())
    build_classification(DATA / "2_분류결과_참고용.xlsx", items, answers,
                         answer=True, sheet_name="분류결과",
                         guide_sections=[("이 표는 무엇인가요",
                                          ["모든 문의의 정답 분류가 이미 적혀 있는 표입니다(문의유형·감성·위험).",
                                           "여기서 직접 분류하지 않습니다. '2_채널별집계_제출용.xlsx'를 채울 때",
                                           "이 표를 보고 채널별 건수만 세면 됩니다."])])
    build_crosstab(DATA / "2_채널별집계_제출용.xlsx", items, answers,
                   answer=False, guide_sections=crosstab_guide_sections())
    # 정답키
    build_classification(ANS / "1_분류표_정답.xlsx", items, answers,
                         answer=True, sheet_name="분류", guide_sections=rulebook_sections())
    build_crosstab(ANS / "2_채널별집계_정답.xlsx", items, answers,
                   answer=True, guide_sections=crosstab_guide_sections())
    print("[4/5] 엑셀/PDF 산출 완료: 원문 / 규칙북PDF / 분류표(제출·정답) / 참고용 / 교차집계(제출·정답)")

    # 분포/함정 요약
    conf_t = sum(1 for c in items if c["충돌주입"] == "유형충돌")
    conf_s = sum(1 for c in items if c["충돌주입"] == "감성충돌")
    undet = sum(1 for c in items if c["경계케이스"] == "판별불가")
    weak = sum(1 for c in items if c["경계케이스"] == "약한신호")
    indep_neg = sum(1 for c in items if answers[c["item_id"]][1] == "부정" and answers[c["item_id"]][2] == "일반")
    indep_risk = sum(1 for c in items if answers[c["item_id"]][1] in ("중립", "긍정") and answers[c["item_id"]][2] == "위험")
    cell, risk = _crosstab_counts(items, answers)
    print("[5/5] 분포 요약")
    print("=" * 64)
    print("  항목 수:", len(items))
    print("  유형 분포:", type_count)
    print("  감성 분포:", senti_count)
    print("  위험 분포:", risk_count, f"(위험 비율 {risk_count['위험']/len(items)*100:.1f}%)")
    print("  채널 분포:", chan_count)
    print("  충돌: 유형충돌", conf_t, "/ 감성충돌", conf_s)
    print("  경계: 판별불가", undet, "/ 약한신호", weak)
    print("  위험 축독립: (부정&일반)", indep_neg, "/ (중립·긍정&위험)", indep_risk)
    print("  교차집계 채널×유형(0셀 개수):",
          sum(1 for ch in CHANNELS for t in TYPE_ORDER if cell[ch][t] == 0), "/ 30")
    print("=" * 64)


if __name__ == "__main__":
    main()
